"""Build the standard Excel entry templates in templates/ from config/templates.yaml.

Each template has:
  - an "Instructions" sheet,
  - one sheet per data table, with a single header row, dropdown lists and
    simple Excel checks (dates must be dates, counts must be whole numbers),
  - a hidden "Lists" sheet that holds the dropdown values.

HOW TO RUN
    python src/build_templates.py            (only rebuilds if something changed)
    python src/build_templates.py --force    (always rebuilds)
"""

from __future__ import annotations

import hashlib
import json
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from common import load_admin, load_settings, load_yaml, resolve_path

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
REQUIRED_FILL = PatternFill("solid", fgColor="FCE4D6")
ROWS_WITH_CHECKS = 2000   # how many rows below the header get dropdowns and checks


def resolve_list(spec) -> list[str]:
    """Turn a list entry from templates.yaml into the actual dropdown values."""
    if isinstance(spec, list):
        return [str(v) for v in spec]
    source = spec["from"]
    if source == "rwanda_admin":
        return sorted(d for ds in load_admin()["provinces"].values() for d in ds)
    if source == "donor_pipeline_stages":
        return load_settings()["donor_pipeline"]["stages"]
    synthetic = load_yaml("synthetic.yaml")
    if source == "synthetic_projects":
        return [p["name"] for p in synthetic["projects"]]
    if source == "synthetic_activities":
        return [a["activity_name"] for a in synthetic["activities"]]
    if source == "synthetic_officers":
        return list(synthetic["field_officers"]) + [synthetic["paper_data_entry_clerk"]]
    raise ValueError(f"Unknown list source: {source}")


def build_template(file_name: str, spec: dict, lists: dict) -> None:
    wb = Workbook()
    intro = wb.active
    intro.title = "Instructions"
    intro["A1"] = file_name.replace("_", " ").replace(".xlsx", "").title()
    intro["A1"].font = Font(bold=True, size=14)
    for i, line in enumerate(spec["instructions"], start=3):
        intro.cell(i, 1, f"- {line}")
    row = len(spec["instructions"]) + 4
    intro.cell(row, 1, "Orange headers are required. Do not rename, reorder or delete columns.").font = Font(italic=True)
    intro.column_dimensions["A"].width = 110

    # Hidden sheet with every dropdown list used in this template.
    list_ws = wb.create_sheet("Lists")
    list_ranges: dict[str, str] = {}
    used = {c["list"] for cols in spec["sheets"].values() for c in cols if "list" in c}
    for col_idx, name in enumerate(sorted(used), start=1):
        values = resolve_list(lists[name])
        letter = get_column_letter(col_idx)
        list_ws.cell(1, col_idx, name).font = Font(bold=True)
        for r, v in enumerate(values, start=2):
            list_ws.cell(r, col_idx, v)
        list_ranges[name] = f"Lists!${letter}$2:${letter}${len(values) + 1}"
    list_ws.sheet_state = "hidden"

    for sheet_name, columns in spec["sheets"].items():
        ws = wb.create_sheet(sheet_name, index=len(wb.sheetnames) - 1)
        for i, col in enumerate(columns, start=1):
            letter = get_column_letter(i)
            cell = ws.cell(1, i, col["header"])
            cell.font = Font(bold=True)
            cell.fill = REQUIRED_FILL if col.get("required") else HEADER_FILL
            if col.get("note"):
                cell.comment = Comment(col["note"], "Template")
            ws.column_dimensions[letter].width = max(12, len(col["header"]) + 4)
            cells = f"{letter}2:{letter}{ROWS_WITH_CHECKS + 1}"
            dv = None
            if "list" in col:
                dv = DataValidation(type="list", formula1=f"={list_ranges[col['list']]}", allow_blank=True)
                dv.error = "Please pick a value from the list."
            elif col.get("type") == "date":
                dv = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)",
                                    formula2="DATE(2100,12,31)", allow_blank=True)
                dv.error = "Please type a date, for example 12/03/2024."
                for r in range(2, ROWS_WITH_CHECKS + 2):
                    ws.cell(r, i).number_format = "dd/mm/yyyy"
            elif col.get("type") == "whole_number":
                dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
                dv.error = "Please type a whole number (0 or more)."
            elif col.get("type") == "number":
                dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
                dv.error = "Please type a number only (no RWF, no commas)."
            elif col.get("type") == "text":
                for r in range(2, ROWS_WITH_CHECKS + 2):
                    ws.cell(r, i).number_format = "@"   # keeps leading zeros in phone and ID numbers
            if dv is not None:
                dv.showErrorMessage = True
                ws.add_data_validation(dv)
                dv.add(cells)
        ws.freeze_panes = "A2"

    out = resolve_path("templates") / file_name
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Wrote {out}")


STAMP_FILE = ".build_stamp"


def content_stamp(cfg: dict) -> str:
    """A code that changes only when what the templates contain changes: templates.yaml itself
    and the values behind its dropdown lists (districts, projects, activities, officers, stages)."""
    resolved = {name: resolve_list(spec) for name, spec in cfg["lists"].items()}
    text = json.dumps({"templates": cfg["templates"], "lists": resolved}, sort_keys=True, default=str)
    return hashlib.sha256(text.encode()).hexdigest()


def main(force: bool = False) -> bool:
    """Rebuild the templates only when their content has changed. Returns True if rebuilt.

    Rebuilding on every run would give the .xlsx files a new save time each time, so git
    would always show them as changed.
    """
    cfg = load_yaml("templates.yaml")
    folder = resolve_path("templates")
    stamp = content_stamp(cfg)
    stamp_path = folder / STAMP_FILE
    all_present = all((folder / name).exists() for name in cfg["templates"])
    if not force and all_present and stamp_path.exists() and stamp_path.read_text().strip() == stamp:
        print("Templates are up to date (templates.yaml and its lists have not changed).")
        return False
    for file_name, spec in cfg["templates"].items():
        build_template(file_name, spec, cfg["lists"])
    stamp_path.write_text(stamp + "\n", encoding="utf-8")
    return True


if __name__ == "__main__":
    main(force="--force" in sys.argv)
