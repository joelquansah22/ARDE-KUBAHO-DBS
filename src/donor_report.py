"""Make the one-page quarterly donor report for a project, as Excel and PDF.

This is the reporting task timed in the evaluation (brief Section 10): producing a quarterly
progress summary for one project. The dashboard's "Generate donor report" button calls make_report();
every run records how long it took in evaluation/report_generation_log.csv.

HOW TO RUN (without the dashboard)
    python src/donor_report.py PRJ01 "FY2025/26 Q4"
"""

from __future__ import annotations

import csv
import datetime as dt
import sys
import time
from pathlib import Path

import duckdb
import matplotlib
import pandas as pd

matplotlib.use("Agg")                      # draw charts to files, no window
import matplotlib.pyplot as plt            # noqa: E402
from fpdf import FPDF                      # noqa: E402
from openpyxl import Workbook              # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

import measures                            # noqa: E402
from common import PROJECT_ROOT, load_settings, resolve_path  # noqa: E402

LOG = PROJECT_ROOT / "evaluation" / "report_generation_log.csv"
STATUS_WORDS = {"on track": "On track", "at risk": "At risk", "off track": "Off track",
                "no target set": "No target set", "target incomplete": "Target incomplete", "no data": "No data"}


def latin1(text) -> str:
    """The PDF's built-in fonts only know Western European characters."""
    return str(text if text is not None else "").replace("–", "-").replace("—", "-").encode(
        "latin-1", "replace").decode("latin-1")


def gather(db: Path, project_id: str, quarter: str) -> dict:
    """Everything the report shows, read from the warehouse."""
    con = duckdb.connect(str(db), read_only=True)
    q = lambda sql, *p: con.execute(sql, list(p)).df()
    try:
        project = q("""SELECT p.*, g.name AS programme FROM dim_project p JOIN dim_programme g USING (programme_id)
                       WHERE project_id = ?""", project_id).iloc[0]
        values = q("""SELECT v.*, i.name AS indicator FROM indicator_values v JOIN dim_indicator i USING (indicator_id)
                      WHERE v.period_type = 'quarter' AND v.period = ? AND v.project_id = ? AND v.district = 'All'
                        AND v.disaggregation = 'total' ORDER BY v.indicator_code, v.source DESC""", quarter, project_id)
        start = values["period_start"].min() if len(values) else None
        if start is None:
            any_value = q("SELECT MIN(period_start) AS s FROM indicator_values WHERE period_type='quarter' AND period=?",
                          quarter)
            start = any_value["s"].iloc[0]
        start = pd.Timestamp(start).date()
        end = (pd.Timestamp(start) + pd.DateOffset(months=3)).date()
        last_day = end - dt.timedelta(days=1)
        # The same definitions as the dashboard (src/measures.py), so the report and the pages agree.
        people = measures.people_reached(q, [project_id], start, last_day)
        totals = measures.activities_total(q, [project_id], start, last_day)
        acts = dict(done=totals["activities"], spend=totals["spend"], attendance=totals["attendance"])
        budget_check = measures.budget_status(q, [project_id], last_day, load_settings()).iloc[0]
        spend_to_date = q("""SELECT SUM(cost_rwf) AS s FROM fact_activity f JOIN dim_activity a USING (activity_id)
                             WHERE a.project_id = ? AND f.date < ?""", project_id, end)["s"].iloc[0]
        flags = measures.reconciliation_counts(
            measures.reconciliation(q, [project_id], start, last_day))["gap above tolerance"]
        warnings = q("""SELECT COUNT(*) AS n FROM fact_activity f JOIN dim_activity a USING (activity_id)
                        WHERE a.project_id = ? AND f.quality_flag = 'warning' AND f.date >= ? AND f.date < ?""",
                     project_id, start, end)["n"].iloc[0]
    finally:
        con.close()
    return dict(project=project, values=values, start=start, end=end, people=people, acts=acts,
                spend_to_date=spend_to_date or 0, flags=int(flags), warnings=int(warnings), budget_check=budget_check)


def indicator_rows(values: pd.DataFrame) -> list[list]:
    from common import load_yaml
    missing = {i["code"]: i["when_missing"] for i in load_yaml("indicator_catalogue.yaml")["indicators"]
               if i.get("when_missing")}
    rows = []
    for r in values.itertuples():
        if pd.isna(r.actual) and r.indicator_code in missing:      # e.g. cost per person: no reported reach
            rows.append([r.indicator_code, r.indicator, r.source, "-", missing[r.indicator_code], "-", "-", ""])
            continue
        pct = (f"{r.pct_achieved:.0f}%" if pd.notna(r.pct_achieved)
               else f"{r.partial_pct_achieved:.0f}% ({r.target_note})" if pd.notna(r.partial_pct_achieved) else "-")
        unit = "%" if r.unit == "percent" else (" RWF" if r.unit == "RWF" else "")
        fmt = lambda x: "-" if pd.isna(x) else (f"{x:,.1f}{unit}" if r.unit in ("percent", "hectares") else f"{x:,.0f}{unit}")
        cov = "" if pd.isna(r.coverage_pct) or r.coverage_pct >= 100 else f"{r.coverage_pct:.0f}%"
        rows.append([r.indicator_code, r.indicator, r.source, fmt(r.target), fmt(r.actual), pct,
                     STATUS_WORDS.get(r.status, r.status), cov])
    return rows


def chart(values: pd.DataFrame, path: Path) -> bool:
    """Bar chart of % achieved for reported indicators. A partial target (some months without one) is drawn
    hatched and labelled "partial", so it is never mistaken for a full result."""
    data = values[values["source"] == "reported"].copy()
    data["pct"] = data["pct_achieved"].where(data["pct_achieved"].notna(), data["partial_pct_achieved"])
    data = data[data["pct"].notna()]
    if not len(data):
        return False
    fig, ax = plt.subplots(figsize=(7.2, 0.5 + 0.38 * len(data)), dpi=150)
    for y, r in enumerate(data.itertuples()):
        partial = pd.isna(r.pct_achieved)
        colour = "#8c959f" if partial else "#1a7f37" if r.status == "on track" else (
            "#9a6700" if r.status == "at risk" else "#cf222e")
        ax.barh(y, r.pct, color=colour, hatch="///" if partial else None, edgecolor="white")
        label = f"{r.pct:.0f}%  " + (f"partial ({r.target_note.replace('partial: ', '')})" if partial
                                      else STATUS_WORDS.get(r.status, r.status))
        ax.text(r.pct + 1, y, label, va="center", fontsize=7)
    ax.set_yticks(range(len(data)), data["indicator_code"])
    ax.axvline(100, color="#555", linewidth=0.8, linestyle="--")
    ax.set_xlabel("% of quarterly target achieved", fontsize=7)
    ax.tick_params(labelsize=7)
    ax.set_xlim(0, max(140, data["pct"].max() + 45))
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    return True


def make_report(project_id: str, quarter: str, out_dir: Path | None = None, db: Path | None = None) -> dict:
    """Write the Excel and PDF reports. Returns their paths and how long it took (seconds)."""
    started = time.perf_counter()
    settings = load_settings()
    db = Path(db or resolve_path("warehouse"))
    out_dir = Path(out_dir or resolve_path("exports") / "donor_reports")
    out_dir.mkdir(parents=True, exist_ok=True)
    d = gather(db, project_id, quarter)
    p, people, acts = d["project"], d["people"], d["acts"]
    safe_q = quarter.replace("/", "-").replace(" ", "_")
    stem = f"{'SYNTHETIC_' if settings['synthetic_mode'] else ''}donor_report_{project_id}_{safe_q}"
    banner = settings["synthetic_banner"] if settings["synthetic_mode"] else ""
    title = f"Quarterly progress report: {p['name']} ({project_id})"
    period = f"{quarter} ({d['start']:%d %b %Y} to {(d['end'] - dt.timedelta(days=1)):%d %b %Y})"
    budget = float(p["budget_rwf"])
    kpis = [("People reached (different registered people)", f"{int(people['people']):,}"
             + (f" (includes {int(people['pending'])} pending duplicate review)" if people["pending"] else "")),
            ("Activities completed in the quarter", f"{int(acts['done'] or 0):,}"),
            ("Activity spending in the quarter", f"RWF {float(acts['spend'] or 0):,.0f}"),
            ("Spending to date / budget", f"RWF {float(d['spend_to_date']):,.0f} / RWF {budget:,.0f} "
             f"({100 * float(d['spend_to_date']) / budget:.0f}%)"
             + (f" - WARNING: {d['budget_check']['message']}" if d["budget_check"]["status"] == "over budget" else "")),
            ("Data checks", f"{d['flags']} reported-vs-records gaps flagged; {d['warnings']} activity rows with warnings")]
    headers = ["Code", "Indicator", "Source", "Target", "Actual", "% achieved", "Status", "Coverage"]
    rows = indicator_rows(d["values"])

    # ------------------------------------------------------------------ Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    if banner:
        ws.append([banner])
        ws["A1"].font = Font(bold=True, color="9A6700")
    ws.append([title])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=14)
    ws.append([f"Programme: {p['programme']}"])
    ws.append([f"Period: {period}"])
    ws.append([])
    for label, value in kpis:
        ws.append([label, value])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([])
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    for r in rows:
        ws.append(r)
    ws.append([])
    ws.append(["Status is shown in words. % achieved is against the quarterly target; 'partial' means some months "
               "have no target. Coverage is shown when a value uses only part of the data."])
    for col, width in zip("ABCDEFGH", (40, 46, 10, 12, 12, 30, 16, 10)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    xlsx = out_dir / f"{stem}.xlsx"
    wb.save(xlsx)

    # ------------------------------------------------------------------ PDF (A4, one page)
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(False)
    pdf.add_page()
    pdf.set_margins(12, 10, 12)
    if banner:
        pdf.set_fill_color(255, 243, 205)
        pdf.set_font("Helvetica", "B", 9)
        pdf.cell(0, 7, latin1(banner), border=0, align="C", fill=True, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)
    pdf.set_font("Helvetica", "B", 14)
    pdf.multi_cell(0, 7, latin1(title), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.cell(0, 5, latin1(f"Programme: {p['programme']}   |   Period: {period}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    for label, value in kpis:
        pdf.set_font("Helvetica", "B", 9)
        pdf.cell(78, 5.5, latin1(label))
        pdf.set_font("Helvetica", "", 9)
        pdf.cell(0, 5.5, latin1(value), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    widths = (16, 64, 15, 18, 18, 26, 22, 7)
    pdf.set_font("Helvetica", "B", 7.5)
    pdf.set_fill_color(221, 235, 247)
    for h, w in zip(headers, widths):
        pdf.cell(w, 6, latin1(h if h != "Coverage" else "Cov."), border=1, fill=True)
    pdf.ln()
    pdf.set_font("Helvetica", "", 7)
    for r in rows:
        name = r[1] if len(r[1]) <= 52 else r[1][:50] + "..."
        cells = [r[0], name, r[2], r[3], r[4], r[5] if len(r[5]) <= 22 else r[5][:21] + ".", r[6], r[7]]
        for value, w in zip(cells, widths):
            pdf.cell(w, 5.2, latin1(value), border=1)
        pdf.ln()
    png = out_dir / f"{stem}_chart.png"
    if chart(d["values"], png) and pdf.get_y() < 220:
        pdf.ln(3)
        pdf.image(str(png), x=12, w=186)
    pdf.set_y(-22)
    pdf.set_font("Helvetica", "I", 7)
    pdf.multi_cell(0, 3.6, latin1(
        "Status is given in words, not only colour. Targets are quarterly; 'partial' means some months have no "
        "target. People are counted once each (pseudonymised), never by attendance. Every figure can be traced to "
        f"its Excel source row in the data warehouse. Generated {dt.datetime.now():%d %B %Y %H:%M} by the "
        "ARDE/KUBAHO MEAL prototype."))
    pdf_path = out_dir / f"{stem}.pdf"
    pdf.output(str(pdf_path))
    seconds = round(time.perf_counter() - started, 2)
    log_timing(project_id, quarter, seconds)
    return dict(excel=xlsx, pdf=pdf_path, seconds=seconds, pages=pdf.pages_count)


def log_timing(project_id: str, quarter: str, seconds: float) -> None:
    LOG.parent.mkdir(parents=True, exist_ok=True)
    new = not LOG.exists()
    with open(LOG, "a", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        if new:
            w.writerow(["generated_at", "project_id", "quarter", "seconds"])
        w.writerow([dt.datetime.now().isoformat(timespec="seconds"), project_id, quarter, seconds])


def main():
    project_id, quarter = (sys.argv[1], sys.argv[2]) if len(sys.argv) > 2 else ("PRJ01", "FY2025/26 Q4")
    r = make_report(project_id, quarter)
    print(f"Donor report for {project_id} {quarter}: {r['excel'].name} and {r['pdf'].name} "
          f"({r['pages']} page) in {r['seconds']} s")


if __name__ == "__main__":
    main()
