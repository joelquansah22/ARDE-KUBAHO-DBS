"""Work out the MEAL indicators, reconcile reported and derived values, write the reference sheets.

WHAT THIS SCRIPT DOES
---------------------
For every indicator in config/indicator_catalogue.yaml that has a data source:
  1. works out its value for each project and district, for each project across all districts, and for
     all projects together, by month, fiscal quarter and fiscal year (settings: indicators.period_types);
  2. does this from up to two sources:
       reported  the values typed into the monthly indicator reports (fact_indicator_result)
       derived   values worked out from the activity and beneficiary records (fact_activity,
                 fact_beneficiary_service), e.g. different people trained, % women, cost per person;
  3. adds the target, % achieved and a traffic-light status (thresholds in settings: traffic_light).
     A missing target shows "no target set", never 0% or 100%;
  4. stores the COVERAGE next to each value: how much of the data the value could be worked out from
     (e.g. the share of activities whose timing can be checked);
  5. compares reported and derived values where both exist (the RECONCILIATION table), flagging gaps
     larger than indicators.reconciliation_tolerance_pct;
  6. records, for every value, the fact rows (source_file, source_sheet, source_row) that produced it.

Counting rule for people: always DIFFERENT pseudonymised people (beneficiary_key), never attendance,
so the same person is not counted twice in a period.

Outcome indicators whose survey does not exist yet ("data source not yet available") are skipped.

Results go to the warehouse (indicator_values, indicator_value_sources, indicator_reconciliation),
to data/exports/, and the reference sheets to docs/indicator_reference_sheets.md and .xlsx.

HOW TO RUN
    python src/indicators.py      (after model.py; pipeline.py runs everything in order)
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import duckdb
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import model
from common import fiscal_year_label, load_settings, load_yaml, resolve_path

ALL = "All"
LEVELS = [("project_id", "district"), ("project_id",), ()]      # by project and district, by project, overall
KEYS = ["project_id", "district", "period_type", "period", "period_start", "disaggregation"]
LINEAGE = ["source_table", "source_file", "source_sheet", "source_row", "record_key"]
OUTPUT_TABLES = ["indicator_values", "indicator_value_sources", "indicator_reconciliation"]


# ======================================================================================
# Periods
# ======================================================================================
def period_columns(dates: pd.Series, period_type: str, start_month: int) -> tuple[pd.Series, pd.Series]:
    """(label, first day) of the month, fiscal quarter or fiscal year each date falls in."""
    unique = pd.Series(pd.unique(dates.dropna()))
    labels, starts = {}, {}
    for d in unique:
        d = pd.Timestamp(d).date()
        fy_start_year = d.year if d.month >= start_month else d.year - 1
        if period_type == "month":
            labels[d], starts[d] = f"{d:%Y-%m}", d.replace(day=1)
        elif period_type == "quarter":
            q = ((d.month - start_month) % 12) // 3 + 1
            month0 = (start_month - 1 + 3 * (q - 1)) % 12 + 1
            year0 = fy_start_year + (1 if month0 < start_month else 0)
            labels[d], starts[d] = f"{fiscal_year_label(d, start_month)} Q{q}", dt.date(year0, month0, 1)
        else:
            labels[d], starts[d] = fiscal_year_label(d, start_month), dt.date(fy_start_year, start_month, 1)
    as_date = dates.map(lambda x: pd.Timestamp(x).date() if pd.notna(x) else None)
    return as_date.map(labels), as_date.map(starts)


def expand(df: pd.DataFrame, date_col: str, settings: dict) -> pd.DataFrame:
    """One copy of each row per aggregation level and period type, with its group keys filled in."""
    df = df[df[date_col].notna()]
    start_month = settings["fiscal_year"]["start_month"]
    frames = []
    for period_type in settings["indicators"]["period_types"]:
        label, start = period_columns(df[date_col], period_type, start_month)
        base = df.assign(period_type=period_type, period=label.values, period_start=start.values)
        for level in LEVELS:
            frames.append(base.assign(**{c: ALL for c in ("project_id", "district") if c not in level}))
    return pd.concat(frames, ignore_index=True) if frames else df.assign(**{k: None for k in KEYS})


def with_disaggregation(df: pd.DataFrame, by_sex: bool, sex_col: str = "sex") -> pd.DataFrame:
    """Every row counts in 'total'; for sex-disaggregated indicators it also counts in sex=female/male."""
    total = df.assign(disaggregation="total")
    if not by_sex:
        return total
    known = df[df[sex_col].isin(["female", "male"])]
    return pd.concat([total, known.assign(disaggregation="sex=" + known[sex_col])], ignore_index=True)


# ======================================================================================
# Loading the warehouse
# ======================================================================================
class Data:
    """The fact and dimension tables, joined with what the calculations need."""

    def __init__(self, db: Path):
        con = duckdb.connect(str(db), read_only=True)
        q = lambda sql: con.execute(sql).df()
        try:
            self.activities = q("""
                SELECT f.*, a.project_id, l.district FROM fact_activity f
                JOIN dim_activity a USING (activity_id) JOIN dim_location l USING (location_id)""")
            self.services = q("""
                SELECT s.*, a.project_id, b.sex, l.district FROM fact_beneficiary_service s
                JOIN dim_activity a USING (activity_id)
                JOIN dim_beneficiary b USING (beneficiary_key)
                LEFT JOIN dim_location l ON l.location_id = b.location_id""")
            self.results = q("""
                SELECT f.*, i.code AS indicator_code, l.district FROM fact_indicator_result f
                JOIN dim_indicator i USING (indicator_id) JOIN dim_location l USING (location_id)""")
            self.completeness = q("SELECT * FROM report_completeness")
            self.dim_indicator = q("SELECT * FROM dim_indicator")
        finally:
            con.close()
        for df, cols in ((self.activities, ["date", "planned_date"]), (self.services, ["date"]),
                         (self.results, ["period_date"]), (self.completeness, ["period"])):
            for c in cols:
                df[c] = pd.to_datetime(df[c]).dt.date
        self.services["district"] = self.services["district"].fillna("Unknown")


def fact_links(df: pd.DataFrame, table: str) -> pd.DataFrame:
    return df.assign(source_table=table, record_key=None)


# ======================================================================================
# Calculation methods: each returns (values, links). values has KEYS plus the measures.
# ======================================================================================
def reported(data: Data, spec: dict, settings: dict):
    code = spec["code"]
    rows = data.results[data.results["indicator_code"] == code].copy()
    by_sex = spec.get("disaggregation") == "sex"
    rows["sex"] = rows["disaggregation_key"].str.replace("sex=", "", regex=False)
    if not by_sex:
        rows = rows[rows["disaggregation_key"] == "total"]
    long = with_disaggregation(expand(rows, "period_date", settings), by_sex)
    if by_sex:                                    # the report's "total" rows are not double counted
        long = long[(long["disaggregation"] != "total") | (long["disaggregation_key"] != "total")]
    g = long.groupby(KEYS, dropna=False)
    values = pd.DataFrame(dict(
        actual=g["actual_value"].sum(min_count=1), target=g["target_value"].sum(min_count=1),
        n_target_missing=g["target_value"].apply(lambda s: int(s.isna().sum())), n_rows=g.size())).reset_index()
    # Coverage: reports received out of those expected, for the months this indicator was due.
    is_due = data.completeness["indicators_due"].fillna("").map(lambda codes: code in codes.split(", "))
    due = data.completeness[is_due.astype(bool)]
    exp = expand(due, "period", settings)
    disaggs = values["disaggregation"].unique() if len(values) else ["total"]
    exp = pd.concat([exp.assign(disaggregation=d) for d in disaggs], ignore_index=True)
    cov = exp.groupby(KEYS, dropna=False).agg(coverage_numerator=("submitted", "sum"),
                                              coverage_denominator=("submitted", "size")).reset_index()
    values = cov.merge(values, on=KEYS, how="outer")   # an expected period with no report still shows (no data)
    values["numerator"], values["denominator"] = values["actual"], np.nan
    return values, fact_links(long, "fact_indicator_result")


def unique_people(data: Data, spec: dict, settings: dict):
    method = spec["calculation"]["derived"]
    rows = data.services[data.services["activity_id"].isin(method["activity_ids"])]
    long = with_disaggregation(expand(rows, "date", settings), spec.get("disaggregation") == "sex")
    g = long.groupby(KEYS, dropna=False)
    values = pd.DataFrame(dict(actual=g["beneficiary_key"].nunique(), n_rows=g.size())).reset_index()
    values["numerator"], values["denominator"] = values["actual"], np.nan
    values["coverage_numerator"] = values["coverage_denominator"] = values["n_rows"]
    return values, fact_links(long, "fact_beneficiary_service")


def share_women(data: Data, spec: dict, settings: dict):
    long = with_disaggregation(expand(data.services, "date", settings), False)
    people = long.groupby(KEYS)["beneficiary_key"].nunique().rename("people")
    known = long[long["sex"].notna()].groupby(KEYS)["beneficiary_key"].nunique().rename("known")
    women = long[long["sex"] == "female"].groupby(KEYS)["beneficiary_key"].nunique().rename("women")
    values = pd.concat([people, known, women], axis=1).fillna(0).reset_index()
    values["numerator"], values["denominator"] = values["women"], values["known"]
    values["actual"] = np.where(values["known"] > 0, 100 * values["women"] / values["known"].replace(0, np.nan), np.nan)
    values["coverage_numerator"], values["coverage_denominator"] = values["known"], values["people"]
    values["n_rows"] = long.groupby(KEYS).size().reindex(pd.MultiIndex.from_frame(values[KEYS])).values
    return values, fact_links(long, "fact_beneficiary_service")


def on_time(data: Data, spec: dict, settings: dict):
    # Which activities count, and which can be checked, use the ONE definition of the data-quality
    # measure "on_time_checkable" (config/dq_metrics.yaml), applied to the rows kept in the warehouse.
    import validate
    measure = load_yaml("dq_metrics.yaml")["metrics"]["on_time_checkable"]
    acts = data.activities[validate.condition_mask(data.activities, measure["denominator"], settings)].copy()
    # The period is the month the activity was PLANNED for; records with no planned date (e.g. paper
    # forms) are placed by their activity date and count as "cannot be checked".
    acts["period_date"] = acts["planned_date"].where(acts["planned_date"].notna(), acts["date"])
    long = with_disaggregation(expand(acts, "period_date", settings), False)
    long["checkable"] = validate.condition_mask(long, measure["numerator"], settings)
    long["is_on_time"] = long["on_time"].fillna(False).astype(bool)
    g = long.groupby(KEYS)
    values = pd.DataFrame(dict(due=g.size(), checkable=g["checkable"].sum(), ok=g["is_on_time"].sum())).reset_index()
    values["numerator"], values["denominator"] = values["ok"], values["checkable"]
    values["actual"] = 100 * values["ok"] / values["checkable"].replace(0, np.nan)
    values["coverage_numerator"], values["coverage_denominator"] = values["checkable"], values["due"]
    values["n_rows"] = values["due"]
    return values, fact_links(long, "fact_activity")


def cost_per_person(data: Data, spec: dict, settings: dict):
    acts = data.activities[data.activities["status"].isin(["completed", "delayed"]) & data.activities["date"].notna()]
    cost_long = with_disaggregation(expand(acts, "date", settings), False)
    cost_long["has_cost"] = cost_long["cost_rwf"].notna()
    g = cost_long.groupby(KEYS)
    costs = pd.DataFrame(dict(cost=g["cost_rwf"].sum(min_count=1), acts=g.size(), with_cost=g["has_cost"].sum()))
    people_long = with_disaggregation(expand(data.services, "date", settings), False)
    people = people_long.groupby(KEYS)["beneficiary_key"].nunique().rename("people")
    values = costs.join(people, how="outer").reset_index()
    values["numerator"], values["denominator"] = values["cost"], values["people"]
    values["actual"] = values["cost"] / values["people"].replace(0, np.nan)
    values["coverage_numerator"], values["coverage_denominator"] = values["with_cost"], values["acts"]
    values["n_rows"] = values["acts"].fillna(0) + people_long.groupby(KEYS).size().reindex(
        pd.MultiIndex.from_frame(values[KEYS])).fillna(0).values
    links = pd.concat([fact_links(cost_long, "fact_activity"), fact_links(people_long, "fact_beneficiary_service")],
                      ignore_index=True)
    return values, links


def report_completeness(data: Data, spec: dict, settings: dict):
    rc = data.completeness.assign(
        record_key=lambda d: d["project_id"] + "|" + d["district"] + "|" + d["period"].astype(str))
    long = with_disaggregation(expand(rc, "period", settings), False)
    g = long.groupby(KEYS)
    values = pd.DataFrame(dict(received=g["submitted"].sum(), expected=g.size())).reset_index()
    values["numerator"], values["denominator"] = values["received"], values["expected"]
    values["actual"] = 100 * values["received"] / values["expected"]
    values["coverage_numerator"] = values["coverage_denominator"] = values["expected"]
    values["n_rows"] = values["expected"]
    return values, long.assign(source_table="report_completeness", source_file=None, source_sheet=None,
                               source_row=None)


DERIVED = {"unique_people": unique_people, "share_women": share_women, "on_time": on_time,
           "cost_per_person": cost_per_person, "report_completeness": report_completeness}


# ======================================================================================
# Targets, % achieved and status
# ======================================================================================
def add_status(values: pd.DataFrame, spec: dict, settings: dict) -> pd.DataFrame:
    tl, labels = settings["traffic_light"], settings["indicators"]
    v = values.copy()
    if "target" not in v:
        v["target"] = np.nan
    if "n_target_missing" not in v:
        v["n_target_missing"] = v["n_rows"].fillna(1)
    higher = bool(spec.get("higher_is_better", True))
    v["actual"] = v["actual"].round(2)
    no_data = v["actual"].isna()
    no_target = v["target"].isna()
    incomplete = ~no_target & (v["n_target_missing"].fillna(0) > 0)
    ratio = (v["actual"] / v["target"]) if higher else (v["target"] / v["actual"])
    v["pct_achieved"] = (100 * ratio).where(~no_data & ~no_target & ~incomplete).round(1)
    v["target"] = v["target"].where(~incomplete)          # never show a part-target as if it were complete
    status = np.select(
        [no_data, no_target, incomplete, v["pct_achieved"] >= tl["on_track_min"], v["pct_achieved"] >= tl["at_risk_min"]],
        ["no data", labels["no_target_label"], labels["target_incomplete_label"], "on track", "at risk"],
        default="off track")
    v["status"] = status
    v["coverage_pct"] = (100 * v["coverage_numerator"] / v["coverage_denominator"].replace(0, np.nan)).round(1)
    return v


# ======================================================================================
# Running everything
# ======================================================================================
class IndicatorBuilder:
    def __init__(self, db: Path, settings: dict | None = None):
        self.settings = settings or load_settings()
        self.catalogue = load_yaml("indicator_catalogue.yaml")["indicators"]
        self.data = Data(db)
        self.id_of = dict(zip(self.data.dim_indicator["code"], self.data.dim_indicator["indicator_id"]))

    def build(self) -> dict[str, pd.DataFrame]:
        all_values, all_links = [], []
        for spec in self.catalogue:
            if spec.get("availability", "available") != "available":
                continue                             # e.g. survey-based outcomes: no data source yet
            calc = spec.get("calculation", {})
            sources = []
            if "reported" in calc:
                sources.append(("reported", reported))
            if "derived" in calc:
                sources.append(("derived", DERIVED[calc["derived"]["method"]]))
            for source, fn in sources:
                values, links = fn(self.data, spec, self.settings)
                if not len(values):
                    continue
                values = add_status(values, spec, self.settings)
                values = values.assign(indicator_code=spec["code"], indicator_id=self.id_of[spec["code"]],
                                       source=source, unit=spec["unit"])
                all_values.append(values)
                all_links.append(links.assign(indicator_code=spec["code"], source=source))
        values = pd.concat(all_values, ignore_index=True)
        values = values.sort_values(["indicator_code", "source", "period_type", "period_start", "project_id",
                                     "district", "disaggregation"]).reset_index(drop=True)
        values.insert(0, "value_id", np.arange(1, len(values) + 1))
        ident = ["indicator_code", "source"] + KEYS
        links = pd.concat(all_links, ignore_index=True)
        links = links.merge(values[["value_id"] + ident], on=ident, how="inner")
        links = links[["value_id"] + LINEAGE].drop_duplicates().reset_index(drop=True)
        links["source_row"] = links["source_row"].astype("Int64")
        cols = ["value_id", "indicator_id", "indicator_code", "source", "project_id", "district", "period_type",
                "period", "period_start", "disaggregation", "unit", "target", "actual", "pct_achieved", "status",
                "numerator", "denominator", "coverage_pct", "n_rows"]
        values = values[cols]
        return dict(indicator_values=values, indicator_value_sources=links,
                    indicator_reconciliation=self.reconcile(values))

    def reconcile(self, values: pd.DataFrame) -> pd.DataFrame:
        """Reported versus derived, where an indicator has both."""
        tol = self.settings["indicators"]["reconciliation_tolerance_pct"]
        both = [s["code"] for s in self.catalogue if s.get("availability", "available") == "available"
                and {"reported", "derived"} <= set(s.get("calculation", {}))]
        ident = ["indicator_code", "indicator_id"] + KEYS
        v = values[values["indicator_code"].isin(both)]
        rep = v[v["source"] == "reported"][ident + ["value_id", "actual"]].rename(
            columns={"value_id": "reported_value_id", "actual": "reported"})
        der = v[v["source"] == "derived"][ident + ["value_id", "actual"]].rename(
            columns={"value_id": "derived_value_id", "actual": "derived"})
        rec = rep.merge(der, on=ident, how="outer")
        rec["difference"] = rec["derived"] - rec["reported"]
        rec["pct_difference"] = (100 * rec["difference"] / rec["reported"].replace(0, np.nan)).round(1)
        rec["tolerance_pct"] = tol
        rec["status"] = np.select(
            [rec["reported"].isna(), rec["derived"].isna(), rec["pct_difference"].abs() <= tol],
            ["derived only", "reported only", "within tolerance"], default="gap above tolerance")
        rec["flag"] = rec["status"] == "gap above tolerance"
        for c in ("reported_value_id", "derived_value_id"):
            rec[c] = rec[c].astype("Int64")
        return rec.sort_values(["indicator_code", "period_type", "period_start", "project_id", "district",
                                "disaggregation"]).reset_index(drop=True)[
            ["indicator_code", "indicator_id", "project_id", "district", "period_type", "period", "period_start",
             "disaggregation", "reported", "derived", "difference", "pct_difference", "tolerance_pct", "status",
             "flag", "reported_value_id", "derived_value_id"]]


def check(tables: dict) -> list[str]:
    problems = []
    v, links = tables["indicator_values"], tables["indicator_value_sources"]
    if v.duplicated(["indicator_code", "source"] + KEYS).any():
        problems.append("indicator_values: a value appears twice")
    unlinked = set(v.loc[v["actual"].notna() & (v["n_rows"].fillna(0) > 0), "value_id"]) - set(links["value_id"])
    if unlinked:
        problems.append(f"{len(unlinked)} indicator values have no source rows")
    shown_as_achieved = v[v["target"].isna() & v["pct_achieved"].notna()]
    if len(shown_as_achieved):
        problems.append("a value without a target has a % achieved")
    return problems


def save(tables: dict, db: Path) -> None:
    con = duckdb.connect(str(db))
    try:
        for name, df in tables.items():
            con.execute(f"DROP TABLE IF EXISTS {name}")
            con.register("frame", df)
            con.execute(f"CREATE TABLE {name} AS SELECT * FROM frame")
            con.unregister("frame")
    finally:
        con.close()


# ======================================================================================
# Reference sheets (Markdown and Excel) from the catalogue
# ======================================================================================
SHEET_FIELDS = [("code", "Code"), ("name", "Indicator"), ("status", "Status"), ("thematic_area", "Theme"),
                ("level", "Level"), ("unit", "Unit"), ("definition", "Definition"), ("numerator", "Numerator"),
                ("denominator", "Denominator"), ("calculation_rule", "Calculation rule"),
                ("counting_rule", "Counting rule"), ("coverage_rule", "Coverage"),
                ("disaggregation", "Disaggregation"), ("target_setting_rule", "Target-setting rule"),
                ("frequency", "Frequency"), ("data_source", "Data source"), ("availability", "Availability"),
                ("owner", "Owner"), ("higher_is_better", "Higher is better"), ("notes", "Notes")]


def text(value) -> str:
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return " ".join(str(value).split()) if value is not None else ""


def write_reference_sheets(docs: Path) -> None:
    cat = load_yaml("indicator_catalogue.yaml")
    settings = load_settings()
    lines = ["# Indicator reference sheets", "",
             f"> **{cat['catalogue_status']}.** Nothing here has been agreed with ARDE/KUBAHO. Use "
             "`docs/indicator_reference_sheets.xlsx` (with a Staff comments column) in the co-design session.", "",
             f"*Generated from `config/indicator_catalogue.yaml` by `src/indicators.py` on {dt.date.today():%d %B %Y}. "
             "Do not edit by hand: change the YAML file instead.*", "",
             "Traffic-light status (% of target achieved, from `config/settings.yaml`): on track "
             f"{settings['traffic_light']['on_track_min']}% or more; at risk {settings['traffic_light']['at_risk_min']}% "
             f"to {settings['traffic_light']['on_track_min'] - 1}%; off track below "
             f"{settings['traffic_light']['at_risk_min']}%. A missing target shows \"{settings['indicators']['no_target_label']}\".",
             "", "| Code | Indicator | Level | Unit | Frequency | Availability |", "|---|---|---|---|---|---|"]
    for ind in cat["indicators"]:
        lines.append(f"| {ind['code']} | {text(ind['name'])} | {ind['level']} | {ind['unit']} | {ind['frequency']} | "
                     f"{text(ind.get('availability'))} |")
    for ind in cat["indicators"]:
        lines += ["", f"## {ind['code']}: {text(ind['name'])}", "", "| Field | Value |", "|---|---|"]
        for key, label in SHEET_FIELDS[2:]:
            lines.append(f"| {label} | {text(ind.get(key)).replace('|', '/')} |")
    (docs / "indicator_reference_sheets.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "Indicators"
    headers = [label for _, label in SHEET_FIELDS] + ["Staff comments"]
    ws.append([f"{cat['catalogue_status']} - for the co-design session with ARDE/KUBAHO staff"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(headers)
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(2, len(headers)).fill = PatternFill("solid", fgColor="FFF2CC")
    for ind in cat["indicators"]:
        ws.append([text(ind.get(k)) for k, _ in SHEET_FIELDS] + [""])
    widths = {"Code": 10, "Indicator": 34, "Definition": 45, "Calculation rule": 45, "Counting rule": 40,
              "Notes": 45, "Staff comments": 45}
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 22)
        for row in ws.iter_rows(min_row=3, min_col=i, max_col=i):
            row[0].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C3"
    wb.save(docs / "indicator_reference_sheets.xlsx")


def run(warehouse=None, exports=None, docs=None) -> dict:
    db = Path(warehouse or resolve_path("warehouse"))
    tables = IndicatorBuilder(db).build()
    problems = check(tables)
    if problems:
        raise model.ModelError("Indicator checks failed, nothing saved:\n  - " + "\n  - ".join(problems))
    save(tables, db)
    dictionary = load_yaml("data_dictionary.yaml")
    for name, df in tables.items():                      # the dictionary must describe every column
        described = list(dictionary["tables"][name]["columns"])
        if list(df.columns) != described:
            raise model.ModelError(f"{name}: columns differ from config/data_dictionary.yaml")
    model.export(tables, Path(exports or resolve_path("exports")), dictionary["relationships"], clear=False)
    docs = Path(docs or resolve_path("docs"))
    model.write_dictionary(db, dictionary, docs)
    write_reference_sheets(docs)
    return tables


def main():
    t = run()
    v, rec = t["indicator_values"], t["indicator_reconciliation"]
    print(f"Indicators: {v['indicator_code'].nunique()} indicators, {len(v):,} values; reconciliation: "
          f"{int(rec['flag'].sum())} of {len(rec):,} comparisons above the tolerance. "
          "Reference sheets: docs/indicator_reference_sheets.md and .xlsx")


if __name__ == "__main__":
    main()
