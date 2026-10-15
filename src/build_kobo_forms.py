"""Build the KoboToolbox forms (XLSForm files) for field data collection: forms/*.xlsx.

WHAT THIS SCRIPT DOES
---------------------
Writes two XLSForm files, based on the Excel entry templates, that ARDE/KUBAHO can upload to KoboToolbox
(docs/kobo_guide.md):

  forms/activity_monitoring_kobo.xlsx   one submission per activity (the activity monitoring template)
  forms/indicator_report_kobo.xlsx      one submission per indicator, project, district and month

Each XLSForm has the three standard sheets: "survey" (the questions), "choices" (the dropdown lists) and "settings".
Dropdown lists come from the programme master the pipeline has read (data/staging/ref_*.parquet): projects,
activity types (filtered by project), officer codes, districts, sectors, cells and villages (each list filtered by
the one before), plus the activity statuses and the reported indicators (config). Checks on the phone: dates not in
the future, no negative numbers, and women + men = total.

The question NAMES are the pipeline's standard field names (date, district, female...), so an export with "XML values
and headers" is read without any new mapping; the LABELS are column names the pipeline already knows, so an export
with labels is read too.

HOW TO RUN
    python src/build_kobo_forms.py      (after python src/pipeline.py, which reads the programme master)
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from common import PROJECT_ROOT, load_settings, load_yaml, resolve_path

FORMS = PROJECT_ROOT / "forms"
SURVEY_COLUMNS = ["type", "name", "label", "hint", "required", "constraint", "constraint_message", "relevant",
                  "appearance", "choice_filter", "calculation"]
CHOICE_COLUMNS = ["list_name", "name", "label", "project", "district", "sector", "cell"]
DONE = "${status} = 'completed' or ${status} = 'delayed'"


def q(type_, name, label="", **kw) -> dict:
    return dict(type=type_, name=name, label=label, **kw)


def reference(staging: Path) -> dict[str, pd.DataFrame]:
    refs = {}
    for name in ("ref_project", "ref_activity", "ref_location", "ref_officer"):
        path = staging / f"{name}.parquet"
        if not path.exists():
            raise SystemExit(f"{path} not found: run python src/pipeline.py first (it reads the programme master).")
        refs[name] = pd.read_parquet(path)
    return refs


def common_choices(refs: dict) -> list[dict]:
    rows = [dict(list_name="officers", name=c, label=c) for c in refs["ref_officer"]["officer_code"]]
    rows += [dict(list_name="projects", name=r.project_id, label=r.name)      # labels the pipeline can read
             for r in refs["ref_project"].itertuples()]
    loc = refs["ref_location"]
    rows += [dict(list_name="districts", name=d, label=d) for d in sorted(loc["district"].unique())]
    return rows


def location_choices(refs: dict) -> list[dict]:
    loc = refs["ref_location"]
    rows = [dict(list_name="sectors", name=r.sector, label=r.sector, district=r.district)
            for r in loc[["district", "sector"]].drop_duplicates().itertuples()]
    rows += [dict(list_name="cells", name=r.cell, label=r.cell, district=r.district, sector=r.sector)
             for r in loc[["district", "sector", "cell"]].drop_duplicates().itertuples()]
    rows += [dict(list_name="villages", name=r.village, label=r.village, district=r.district, sector=r.sector,
                  cell=r.cell) for r in loc.itertuples()]
    return rows


def activity_form(refs: dict, settings: dict) -> tuple[list, list]:
    survey = [
        q("start", "start"), q("end", "end"),
        q("today", "entry_date"),                                   # the day the form was filled in (entry date)
        q("select_one officers", "officer", "Officer", required="yes"),
        q("select_one projects", "project", "Project", required="yes"),
        q("select_one activities", "activity", "Activity", required="yes", choice_filter="project=${project}"),
        q("select_one statuses", "status", "Status", required="yes"),
        q("date", "planned_date", "Planned date", hint="The date the activity was planned for, if known."),
        q("date", "date", "Activity date", required="yes", relevant="${status} != 'planned'",
          constraint=". <= today()", constraint_message="The activity date cannot be in the future."),
        q("select_one districts", "district", "District", required="yes"),
        q("select_one sectors", "sector", "Sector", required="yes", choice_filter="district=${district}"),
        q("select_one cells", "cell", "Cell", required="yes",
          choice_filter="district=${district} and sector=${sector}"),
        q("select_one villages", "village", "Village", required="yes",
          choice_filter="district=${district} and sector=${sector} and cell=${cell}"),
        q("integer", "female", "Female", required="yes", relevant=DONE, constraint=". >= 0",
          constraint_message="Cannot be negative."),
        q("integer", "male", "Male", required="yes", relevant=DONE, constraint=". >= 0",
          constraint_message="Cannot be negative."),
        q("integer", "total", "Total", required="yes", relevant=DONE, constraint=". >= 0 and . = ${female} + ${male}",
          constraint_message="The total must equal women + men."),
        q("decimal", "cost", "Cost (RWF)", relevant="${status} != 'planned'", constraint=". >= 0",
          constraint_message="Cannot be negative.", hint="Numbers only, in RWF."),
        q("text", "comment", "Comments", appearance="multiline"),
    ]
    choices = common_choices(refs) + location_choices(refs)
    choices += [dict(list_name="activities", name=r.activity_id, label=r.activity_name, project=r.project_id)
                for r in refs["ref_activity"].itertuples()]
    choices += [dict(list_name="statuses", name=s, label=s.capitalize()) for s in settings["activity_status_values"]]
    return survey, choices


def indicator_form(refs: dict, catalogue: list) -> tuple[list, list]:
    reported = [i for i in catalogue if i.get("availability") == "available" and "reported" in (i.get("calculation") or {})]
    by_sex = [i["code"] for i in reported if i.get("disaggregation") == "sex"]
    survey = [
        q("start", "start"), q("end", "end"), q("today", "entry_date"),
        q("select_one officers", "officer", "Officer", required="yes"),
        q("select_one projects", "project", "Project", required="yes"),
        q("select_one districts", "district", "District", required="yes"),
        q("date", "period", "Reporting period", required="yes", appearance="month-year",
          constraint=". <= today()", constraint_message="The month cannot be in the future."),
        q("select_one indicators", "indicator", "Indicator", required="yes"),
        q("select_one sexes", "sex", "Sex", required="yes",
          relevant=" or ".join(f"${{indicator}} = '{c}'" for c in by_sex) if by_sex else "false()",
          hint="This indicator is reported separately for women and men: one submission each."),
        q("decimal", "target", "Monthly target", constraint=". >= 0", constraint_message="Cannot be negative."),
        q("decimal", "actual", "Achieved", required="yes", constraint=". >= 0", constraint_message="Cannot be negative."),
        q("text", "comment", "Comments", appearance="multiline"),
    ]
    choices = common_choices(refs)
    choices += [dict(list_name="indicators", name=i["code"], label=" ".join(str(i["name"]).split()))
                for i in reported]
    choices += [dict(list_name="sexes", name="female", label="Female"), dict(list_name="sexes", name="male", label="Male")]
    return survey, choices


def write_form(path: Path, title: str, form_id: str, survey: list, choices: list) -> None:
    wb = Workbook()
    for sheet, columns, rows in (("survey", SURVEY_COLUMNS, survey), ("choices", CHOICE_COLUMNS, choices)):
        ws = wb.active if sheet == "survey" else wb.create_sheet(sheet)
        ws.title = sheet
        used = [c for c in columns if c in ("type", "name", "label", "list_name") or any(r.get(c) for r in rows)]
        ws.append(used)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
        for r in rows:
            ws.append([r.get(c, "") for c in used])
        for i, c in enumerate(used, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = 40 if c in ("label", "constraint",
                                                                                   "choice_filter", "relevant") else 18
    st = wb.create_sheet("settings")
    st.append(["form_title", "form_id", "version"])
    st.append([title, form_id, dt.date.today().strftime("%Y%m%d01")])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main(staging: Path | None = None, out: Path | None = None) -> list[Path]:
    settings = load_settings()
    refs = reference(Path(staging or resolve_path("staging")))
    catalogue = load_yaml("indicator_catalogue.yaml")["indicators"]
    out = Path(out or FORMS)
    banner = " (SYNTHETIC)" if settings.get("synthetic_mode") else ""
    paths = []
    for file_name, title, form_id, (survey, choices) in (
            ("activity_monitoring_kobo.xlsx", f"ARDE activity monitoring{banner}", "arde_activity_monitoring",
             activity_form(refs, settings)),
            ("indicator_report_kobo.xlsx", f"ARDE indicator report{banner}", "arde_indicator_report",
             indicator_form(refs, catalogue))):
        write_form(out / file_name, title, form_id, survey, choices)
        paths.append(out / file_name)
        print(f"Wrote {out / file_name}: {len(survey)} questions, {len(choices)} choices")
    return paths


if __name__ == "__main__":
    main()
