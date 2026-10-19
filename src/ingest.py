"""Read the Excel files, pseudonymise personal data and stage the records.

WHAT THIS SCRIPT DOES
---------------------
1. Opens every file in data/raw/ listed in config/source_mappings.yaml.
2. On each sheet, finds the real header row by itself (it may sit below a title
   or a merged group header), and matches each column to a standard field name.
   Blank rows and totals rows are removed and logged.
3. Finds duplicate beneficiaries BEFORE anything is hashed: names are compared
   with fuzzy matching (rapidfuzz), within the same district and sex, and kept
   apart if phone, ID or age clearly disagree. Threshold: config/settings.yaml.
4. Replaces people with salted hashes (HMAC-SHA256):
     beneficiary_key   the same for every registration of one person
     registration_key  one per registration row
   Names, phone numbers, ID numbers and dates of birth are then removed.
   Individual donors become "Individual donor #NNN"; donor contact details are removed.
5. Writes data/staging/ingested_*.parquet, ingest_issues.parquet and
   sheet_inventory.csv (which lists any column the pipeline did not recognise).

The salt (the secret that makes the hashes impossible to reverse) lives outside the
project folder; see privacy.salt_file in config/settings.yaml.

HOW TO RUN
    python src/ingest.py
"""

from __future__ import annotations

import datetime as dt
import fnmatch
import hashlib
import hmac
import os
import re
import secrets
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from rapidfuzz import fuzz, process

from clean import (FIXED, FLAGGED, REMOVED, IssueLog, is_blank, norm_key, parse_date, recode,
                   standardise_district)
from common import PROJECT_ROOT, load_settings, load_yaml, resolve_path

HEADER_SEARCH_ROWS = 15


# ======================================================================================
# Salt and hashing
# ======================================================================================
def salt_path() -> Path:
    """Where the salt is kept. The ARDE_SALT_FILE environment variable overrides settings (used by tests)."""
    configured = os.environ.get("ARDE_SALT_FILE") or load_settings()["privacy"]["salt_file"]
    return Path(os.path.expanduser(configured)).resolve()


def load_salt() -> bytes:
    """Read the salt, creating a new random one the first time. Refuses a salt inside the project."""
    path = salt_path()
    if PROJECT_ROOT.resolve() in path.parents:
        raise RuntimeError(f"The salt file must be outside the project folder, not {path}")
    if not path.exists():
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(secrets.token_hex(32), encoding="utf-8")
        print(f"Created a new pseudonymisation salt at {path}. Back it up somewhere safe: "
              "without it, new hashes will not match old ones.")
    return path.read_text(encoding="utf-8").strip().encode()


def pseudonym(salt: bytes, prefix: str, text: str) -> str:
    """A salted hash: the same input always gives the same key, but the key cannot be reversed."""
    return f"{prefix}-{hmac.new(salt, text.encode(), hashlib.sha256).hexdigest()[:12]}"


# ======================================================================================
# Reading sheets
# ======================================================================================
def to_text(value):
    """Store every raw cell as text so nothing is lost before cleaning."""
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def load_mappings(overlays: tuple[str, ...] = ()) -> dict:
    """config/source_mappings.yaml, plus any extra mapping files (their column names are added)."""
    mappings = load_yaml("source_mappings.yaml")
    for name in overlays:
        extra = load_yaml(name)
        mappings["files"] = extra.get("files", []) + mappings["files"]
        for entity, fields in (extra.get("fields") or {}).items():
            for fld, aliases in fields.items():
                mappings["fields"][entity].setdefault(fld, [])
                mappings["fields"][entity][fld] = mappings["fields"][entity][fld] + list(aliases)
    return mappings


def alias_index(entity: str, mappings: dict | None = None) -> dict[str, list[str]]:
    """{normalised header text: [fields it can mean]}"""
    index: dict[str, list[str]] = {}
    for fld, aliases in (mappings or load_mappings())["fields"][entity].items():
        for alias in aliases:
            index.setdefault(norm_key(alias), []).append(fld)
    return index


def find_header(rows: list[tuple], index: dict) -> tuple[int, int]:
    """Return (position of the header row, number of known column names in it)."""
    best, best_score = 0, -1
    for i, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        score = sum(1 for v in row if v is not None and norm_key(v) in index)
        if score > best_score:
            best, best_score = i, score
    return best, best_score


def assign_fields(headers: list, index: dict) -> tuple[dict[int, str], list[str]]:
    """Match each column to a field. A header that could mean two things (e.g. 'ID')
    takes whichever field is not already claimed by another column."""
    candidates = {i: index.get(norm_key(h), []) for i, h in enumerate(headers) if h is not None}
    assigned: dict[int, str] = {}
    taken: set[str] = set()
    for i, options in candidates.items():
        if len(options) == 1 and options[0] not in taken:
            assigned[i] = options[0]
            taken.add(options[0])
    for i, options in candidates.items():
        if i not in assigned:
            free = [o for o in options if o not in taken]
            if free:
                assigned[i] = free[0]
                taken.add(free[0])
    unmapped = [str(headers[i]) for i in candidates if i not in assigned]
    return assigned, unmapped


def sheet_entity(file_name: str, sheet: str, mappings: dict | None = None) -> str | None:
    for spec in (mappings or load_mappings())["files"]:
        if fnmatch.fnmatch(file_name.lower(), spec["pattern"].lower()):
            sheets = spec["sheets"]
            return sheets.get(sheet) or sheets.get("*")
    return None


class Reader:
    """Reads every mapped sheet into lists of raw records, with lineage."""

    def __init__(self, raw_dir: Path, settings: dict, mappings: dict | None = None):
        self.raw_dir, self.s = raw_dir, settings
        self.mappings = mappings or load_mappings()
        self.skipped: list[dict] = []
        self.records: dict[str, list[dict]] = {}
        self.refs: dict[str, pd.DataFrame] = {}
        self.inventory: list[dict] = []
        self.log = IssueLog()
        self.template_sheets = {t.lower() for t in self.mappings.get("template_sheets", [])}
        self.totals = {" ".join(re.sub(r"[^a-z ]", " ", norm_key(t)).split())
                       for t in settings["cleaning"]["totals_labels"]}

    def read_all(self):
        for path in sorted(self.raw_dir.glob("*.xlsx")):
            if path.name.startswith("~$"):
                continue
            wb = load_workbook(path, read_only=True, data_only=False)
            for ws in wb.worksheets:
                entity = sheet_entity(path.name, ws.title, self.mappings)
                rows = list(ws.iter_rows(values_only=True))
                if ws.title.strip().lower() in self.template_sheets:
                    self.inventory.append(dict(source_file=path.name, source_sheet=ws.title, entity=None,
                                               note="template help sheet (source_mappings.yaml: template_sheets): ignored"))
                elif entity is None:
                    self.inventory.append(dict(source_file=path.name, source_sheet=ws.title, entity=None,
                                               note="sheet not listed in source_mappings.yaml: ignored"))
                elif entity.startswith("ref_"):
                    self.read_reference(entity, rows)
                else:
                    self.read_sheet(path.name, ws.title, entity, rows)
            wb.close()

    def missing_required(self, entity: str, found: set) -> list[str]:
        """Required fields (source_mappings.yaml: required_fields) that this sheet does not have."""
        missing = []
        for item in self.mappings.get("required_fields", {}).get(entity, []):
            options = item if isinstance(item, list) else [item]
            if not set(options) & found:
                missing.append(" or ".join(options))
        return missing

    def skip(self, entity, base, code, reason, unrecognised, known, missing):
        """Skip a whole sheet: print a clear message, list it in skipped_sheets.csv and log it."""
        message = (f"SKIPPED {base['source_file']} / {base['source_sheet']}: {reason}. "
                   f"Unrecognised column headers: {', '.join(unrecognised) or '(none)'}. "
                   "Add them to config/source_mappings.yaml and run the pipeline again.")
        print(message)
        self.skipped.append(dict(source_file=base["source_file"], source_sheet=base["source_sheet"], entity=entity,
                                 reason=reason, known_columns=known, missing_required=", ".join(missing),
                                 headers_found=", ".join(unrecognised), message=message))
        self.inventory.append(dict(source_file=base["source_file"], source_sheet=base["source_sheet"], entity=entity,
                                   header_row=None, fields_found="", unmapped_columns=", ".join(unrecognised),
                                   note=f"skipped: {reason}"))
        self.log.add(entity, {**base, "source_row": None}, None, code, FLAGGED, detail=message)

    def is_totals_row(self, row: tuple, first) -> bool:
        """True if the first filled cell is a totals label ("TOTAL", "Total (RWF)", "Sub-total:")
        and every other filled cell is a number. A row with any other text is never dropped."""
        if not isinstance(first, str):
            return False
        words = re.sub(r"[^a-z ]", " ", norm_key(first)).split()
        label = " ".join(words)
        if not any(label == t or label.startswith(t + " ") for t in self.totals):
            return False
        others = [v for v in row if not is_blank(v) and v is not first]
        return all(isinstance(v, (int, float)) or re.fullmatch(r"[\d\s,.\-]+", str(v)) for v in others)

    def read_reference(self, entity: str, rows: list[tuple]):
        header = [h for h in rows[0] if h is not None]
        data = [[to_text(v) for v in r[: len(header)]] for r in rows[1:] if any(v is not None for v in r)]
        self.refs[entity] = pd.DataFrame(data, columns=header)

    def read_sheet(self, file_name: str, sheet: str, entity: str, rows: list[tuple]):
        index = alias_index(entity, self.mappings)
        h, score = find_header(rows, index)
        needed = self.s["cleaning"]["header_min_matches"]
        base = dict(source_file=file_name, source_sheet=sheet)
        if score < needed:
            # Do not guess: skip the sheet, say exactly why, and carry on with the other files.
            first = next((r for r in rows[:HEADER_SEARCH_ROWS] if any(not is_blank(v) for v in r)), ())
            headers = [str(v).strip() for v in first if not is_blank(v)]
            self.skip(entity, base, "sheet_skipped_unrecognised_layout",
                      f"layout not recognised (best row has {max(score, 0)} known column name(s); {needed} needed)",
                      headers, max(score, 0), [])
            return
        assigned, unmapped = assign_fields(list(rows[h]) if rows else [], index)
        missing = self.missing_required(entity, set(assigned.values()))
        if missing:
            # No silent partial reads: a sheet without all its required fields is skipped.
            self.skip(entity, base, "sheet_skipped_required_fields_missing",
                      f"required field(s) not found: {', '.join(missing)}", unmapped, len(assigned), missing)
            return
        self.inventory.append(dict(source_file=file_name, source_sheet=sheet, entity=entity, header_row=h + 1,
                                   fields_found=", ".join(assigned.values()), unmapped_columns=", ".join(unmapped),
                                   note=""))
        for i in range(h):
            if any(not is_blank(v) for v in rows[i]):
                self.log.add(entity, {**base, "source_row": i + 1}, None, "header_preamble_row_removed", REMOVED,
                             detail="title or group-header row above the column names")
        out = self.records.setdefault(entity, [])
        # Rows after the last filled one are not records (a template is formatted far below its data): not logged.
        last = max((i for i in range(h + 1, len(rows)) if any(not is_blank(v) for v in rows[i])), default=h)
        for i in range(h + 1, last + 1):
            row, excel_row = rows[i], i + 1
            rec = {**base, "source_row": excel_row}
            if all(is_blank(v) for v in row):
                self.log.add(entity, rec, None, "blank_row_removed", REMOVED)
                continue
            first = next((v for v in row if not is_blank(v)), None)
            if any(isinstance(v, str) and v.startswith("=") for v in row) or self.is_totals_row(row, first):
                self.log.add(entity, rec, None, "totals_row_removed", REMOVED, detail=f"'{first}' row")
                continue
            for col, fld in assigned.items():
                rec[fld] = row[col] if col < len(row) else None
            out.append(rec)


# ======================================================================================
# Identities, duplicate matching and pseudonymisation
# ======================================================================================
def name_key(text) -> str:
    """Normalised name for matching: lower case, letters only, words sorted elsewhere by the scorer."""
    return re.sub(r"[^a-z ]", "", norm_key(text or "")).strip()


def identity_frame(records: list[dict], settings: dict, log: IssueLog | None = None) -> pd.DataFrame:
    """Build the (in-memory only) identity details used for matching. Never written to disk."""
    log = log or IssueLog()
    rows = []
    for rec in records:
        name = rec.get("name")
        if is_blank(name):
            name = " ".join(str(rec[f]) for f in ("surname", "first_name") if not is_blank(rec.get(f)))
        # Phone: digits only, 07XXXXXXXX
        raw_phone = rec.get("phone")
        phone = None
        if not is_blank(raw_phone):
            digits = re.sub(r"\D", "", str(int(raw_phone)) if isinstance(raw_phone, (int, float)) else str(raw_phone))
            if digits.startswith("250") and len(digits) == 12:
                digits = "0" + digits[3:]
                log.add("registration", rec, "phone", "phone_standardised", FIXED)
            elif len(digits) == 9 and digits.startswith("7"):
                digits = "0" + digits
                log.add("registration", rec, "phone", "phone_leading_zero_restored", FIXED)
            if len(digits) == 10 and digits.startswith("07"):
                phone = digits
            else:
                log.add("registration", rec, "phone", "phone_invalid", FLAGGED, detail=f"{len(digits)} digits")
        # National ID: 16 digits
        raw_id = rec.get("national_id")
        national_id = None
        if not is_blank(raw_id):
            if isinstance(raw_id, (int, float)):
                log.add("registration", rec, "national_id", "id_stored_as_number", FLAGGED,
                        detail="stored as a number: Excel keeps only 15 digits, so the last may be wrong")
                raw_id = str(int(raw_id))
            digits = re.sub(r"\D", "", str(raw_id))
            if len(digits) == 16:
                national_id = digits
            else:
                log.add("registration", rec, "national_id", "id_invalid", FLAGGED, detail=f"{len(digits)} digits")
        # Age, possibly from date of birth
        age, age_source = None, "age"
        if not is_blank(rec.get("age")):
            try:
                age = float(str(rec["age"]).strip())
            except ValueError:
                age = None
        elif not is_blank(rec.get("dob")):
            age_source = "dob"
            dob = parse_date(to_text(rec["dob"]), settings).value
            ref = parse_date(to_text(rec.get("reg_date")), settings).value
            if dob and ref:
                age = ref.year - dob.year - ((ref.month, ref.day) < (dob.month, dob.day))
                log.add("registration", rec, "dob", "age_from_date_of_birth", FIXED, clean=age)
            else:
                # Age at registration needs BOTH dates. Never guess with another date.
                log.add("registration", rec, "dob", "age_not_computable", FLAGGED,
                        detail="date of birth could not be read" if not dob
                        else "registration date could not be read, so age at registration is unknown")
        rows.append(dict(
            source_file=rec["source_file"], source_sheet=rec["source_sheet"], source_row=rec["source_row"],
            reg_no=None if is_blank(rec.get("reg_no")) else str(rec["reg_no"]).strip(),
            name=name_key(name), phone=phone, national_id=national_id,
            sex=recode(rec.get("sex"), settings["sex_mapping"], "sex").value,
            district=standardise_district(rec.get("district"), settings).value,
            age=age, age_source=age_source, reg_date=to_text(rec.get("reg_date"))))
    return pd.DataFrame(rows)


def _digits_apart(a: str, b: str) -> int:
    if len(a) != len(b):
        return 99
    return sum(x != y for x, y in zip(a, b))


def cluster_identities(ids: pd.DataFrame, threshold: float, settings: dict,
                       use_id_rule: bool = True, use_vetoes: bool = True,
                       review_pairs: list | None = None) -> np.ndarray:
    """Group registrations that belong to the same person. Returns a cluster number per row.

    Two rows are joined when they have the same national ID, or when their names score
    at least `threshold` (0 to 100, word order ignored), they are in the same district
    and of the same sex (when known), their phone, ID and age do not clearly disagree,
    AND a phone number or ID confirms the match. Similar names with nothing to confirm them
    are not joined; they are added to review_pairs for a person to check.

    use_id_rule / use_vetoes can be switched off only to measure what each rule adds
    (score_cleaning.py does this); the pipeline always uses both.
    """
    rules = settings["duplicate_matching"]
    n = len(ids)
    parent = list(range(n))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    col = lambda c: [None if is_blank(v) else v for v in ids[c].tolist()]   # pandas may give NaN for empty
    names, phones, nids = col("name"), col("phone"), col("national_id")
    ages, sexes, districts = col("age"), col("sex"), col("district")
    id_digits = rules["compare_id_digits"]

    # Rule 1: the same national ID means the same person, however the name is spelt.
    # (Phones are NOT used this way: family members often share one phone.)
    first_with_id: dict[str, int] = {}
    for i, nid in enumerate(nids):
        if nid and use_id_rule:
            key = nid[:id_digits]
            if key in first_with_id:
                parent[find(i)] = find(first_with_id[key])
            else:
                first_with_id[key] = i

    # Rule 2: similar names, with nothing that clearly disagrees.
    if rules["require_same_district"]:
        blocks = {}
        for i, d in enumerate(districts):
            blocks.setdefault(d, []).append(i)
        unknown = blocks.pop(None, [])
        groups = [members + unknown for members in blocks.values()] + ([unknown] if unknown else [])
    else:
        groups = [list(range(n))]

    for members in groups:
        if len(members) < 2:
            continue
        scores = process.cdist([names[i] for i in members], [names[i] for i in members],
                               scorer=fuzz.token_sort_ratio, score_cutoff=threshold, workers=-1)
        for a, b in zip(*np.nonzero(np.triu(scores, k=1))):
            i, j = members[a], members[b]
            if rules["require_same_sex"] and sexes[i] and sexes[j] and sexes[i] != sexes[j]:
                continue
            if not use_vetoes:
                parent[find(i)] = find(j)
                continue
            if phones[i] and phones[j] and _digits_apart(phones[i], phones[j]) > rules["max_phone_digit_difference"]:
                continue
            if nids[i] and nids[j] and nids[i][:id_digits] != nids[j][:id_digits]:
                continue
            if ages[i] is not None and ages[j] is not None and abs(ages[i] - ages[j]) > rules["max_age_difference_years"]:
                continue
            confirmed = ((phones[i] and phones[j]
                          and _digits_apart(phones[i], phones[j]) <= rules["max_phone_digit_difference"])
                         or (nids[i] and nids[j] and nids[i][:id_digits] == nids[j][:id_digits]))
            if rules.get("merge_requires_phone_or_id", True) and not confirmed:
                # Similar names alone: not merged. Sent to the review queue instead.
                if review_pairs is not None:
                    review_pairs.append((i, j))
                continue
            parent[find(i)] = find(j)
    return np.array([find(i) for i in range(n)])


def initials_review_pairs(ids: pd.DataFrame, settings: dict) -> list[tuple[int, int]]:
    """Find registrations where one name uses initials for the given names ("J. C. Uwimana")
    and another writes them in full ("Jean Claude Uwimana"). REVIEW ONLY: never merged.

    A pair is listed when the full words of the shorter name (normally the surname) each match a
    word of the other name (score at least initials_surname_min_score), the initials match the
    first letters of the other name's remaining words, both are in the same district, both ages
    are known and close, and sex, phone and ID do not disagree.
    """
    rules = settings["duplicate_matching"]
    if not rules.get("initials_rule", True):
        return []
    col = lambda c: [None if is_blank(v) else v for v in ids[c].tolist()]
    names, phones, nids = col("name"), col("phone"), col("national_id")
    ages, sexes, districts = col("age"), col("sex"), col("district")
    id_digits, min_score = rules["compare_id_digits"], rules["initials_surname_min_score"]
    words = [(n or "").split() for n in names]
    pairs = []
    for i, wi in enumerate(words):
        initials = [w for w in wi if len(w) == 1]
        full = [w for w in wi if len(w) > 1]
        if not initials or not full or districts[i] is None or ages[i] is None:
            continue
        for j, wj in enumerate(words):
            if j == i or districts[j] != districts[i] or ages[j] is None or any(len(w) == 1 for w in wj):
                continue
            if abs(ages[i] - ages[j]) > rules["max_age_difference_years"]:
                continue
            if sexes[i] and sexes[j] and sexes[i] != sexes[j]:
                continue
            if phones[i] and phones[j] and _digits_apart(phones[i], phones[j]) > rules["max_phone_digit_difference"]:
                continue
            if nids[i] and nids[j] and nids[i][:id_digits] != nids[j][:id_digits]:
                continue
            remaining = list(wj)
            for w in full:                          # every full word (the surname) must be found
                best = max(remaining, key=lambda r: fuzz.ratio(w, r), default=None)
                if best is None or fuzz.ratio(w, best) < min_score:
                    break
                remaining.remove(best)
            else:
                if sorted(w[0] for w in remaining) == sorted(initials):
                    pairs.append((min(i, j), max(i, j)))
    return sorted(set(pairs))


def assign_keys(ids: pd.DataFrame, clusters: np.ndarray, salt: bytes, id_digits: int) -> pd.DataFrame:
    """Give each person one stable beneficiary_key and each row a registration_key."""
    ids = ids.copy()
    ids["cluster"] = clusters
    keys = {}
    for cluster, group in ids.sort_values(["reg_date", "source_file", "source_row"], na_position="last").groupby(
            "cluster", sort=False):
        first = group.iloc[0]
        identity = (f"id:{first['national_id'][:id_digits]}" if not is_blank(first["national_id"])
                    else f"name:{first['name']}|{first['phone']}|{first['district']}")
        keys[cluster] = pseudonym(salt, "BEN", identity)
    ids["beneficiary_key"] = ids["cluster"].map(keys)
    ids["registration_key"] = [pseudonym(salt, "REG", f"{f}|{r}") for f, r in zip(ids["source_file"], ids["reg_no"])]
    ids["duplicate_group_size"] = ids.groupby("cluster")["cluster"].transform("size")
    return ids


def duplicate_certainty(ids: pd.DataFrame, settings: dict) -> dict:
    """For each matched group, say whether the match is 'certain' or only 'suspected'.

    Certain: every record in the group is linked to another one by the same national ID,
    or by the same phone number with names scoring at least duplicate_certain_name_score.
    Certain matches are merged without review; suspected ones go to the review queue.
    """
    id_digits = settings["duplicate_matching"]["compare_id_digits"]
    min_score = settings["duplicate_certain_name_score"]
    result = {}
    for cluster, group in ids[ids["duplicate_group_size"] > 1].groupby("cluster"):
        rows = group.to_dict("records")
        linked = set()
        for i, a in enumerate(rows):
            for j, b in enumerate(rows):
                if i >= j:
                    continue
                same_id = not is_blank(a["national_id"]) and not is_blank(b["national_id"])                     and a["national_id"][:id_digits] == b["national_id"][:id_digits]
                same_phone = not is_blank(a["phone"]) and a["phone"] == b["phone"]                     and fuzz.token_sort_ratio(a["name"], b["name"]) >= min_score
                if same_id or same_phone:
                    linked |= {i, j}
        result[cluster] = "certain" if len(linked) == len(rows) else "suspected"
    return result


def load_registration_identities(raw_dir: Path | None = None, settings: dict | None = None,
                                 mapping_overlays: tuple[str, ...] = ()) -> pd.DataFrame:
    """Read only the registration identities (used by the scoring script to test thresholds)."""
    settings = settings or load_settings()
    reader = Reader(Path(raw_dir or resolve_path("raw")), settings, load_mappings(tuple(mapping_overlays)))
    reader.read_all()
    return identity_frame(reader.records.get("registration", []), settings)


# ======================================================================================
# Main
# ======================================================================================
APP_SOURCE = "entries.db"


def read_app_programme_entries(path: Path | None = None) -> tuple[list[dict], list[dict]]:
    """Activity records and monthly indicator reports entered on the dashboard (Data entry page), in the same form as
    rows read from Excel, with lineage (source_file entries.db, the table, the row number)."""
    import app_store
    import entry_checks
    path = Path(path or app_store.db_path())
    if not path.exists():
        return [], []
    acts = [entry_checks.raw_record("activity", r, r["entry_date"], "activities", r["row_id"])
            for r in app_store.table("activities", path)]
    reps = [dict(entry_checks.raw_record("indicator_report", r, None, "indicator_reports", r["row_id"]),
                 entered_by=r["created_by"]) for r in app_store.table("indicator_reports", path)]
    return acts, reps


def read_app_entries(path: Path | None = None) -> tuple[list[dict], list[dict]]:
    """Donors and opportunities entered on the dashboard's Donor entry page (data/app/entries.db), as records in the
    same form as rows read from the Excel donor tracker, with lineage (source_file entries.db, the table, the row)."""
    import app_store
    path = Path(path or app_store.db_path())
    if not path.exists():
        return [], []
    donors = [dict(source_file=APP_SOURCE, source_sheet="donors", source_row=r["row_id"], donor_id=r["donor_id"],
                   donor_name=r["donor_name"], donor_type=r["donor_type"], country=r["country"],
                   thematic_interests=r["thematic_interests"]) for r in app_store.table("donors", path)]
    opps = [dict(source_file=APP_SOURCE, source_sheet="opportunities", source_row=r["row_id"],
                 opportunity_id=r["opportunity_id"], title=r["title"], thematic_fit=r["thematic_fit"], stage=r["stage"],
                 deadline=r["deadline"], requested=r["requested_amount_rwf"], awarded=r["awarded_amount_rwf"],
                 probability=r["probability"], owner=r["owner"], next_action=r["next_action"],
                 next_action_date=r["next_action_date"], app_donor_id=r["donor_id"], app_donor_key=r["donor_key"])
            for r in app_store.table("opportunities", path)]
    return donors, opps


def run(raw_dir: Path | None = None, staging: Path | None = None, settings: dict | None = None,
        mapping_overlays: tuple[str, ...] = (), app_entries: Path | None = None) -> dict:
    settings = settings or load_settings()
    # Dashboard entries belong to the production data: read them in the production run (no raw_dir given), or
    # when a path is passed explicitly (tests); never in a scoring run on a test set.
    use_app_entries = app_entries is not None or raw_dir is None
    raw_dir = Path(raw_dir or resolve_path("raw"))
    staging = Path(staging or resolve_path("staging"))
    staging.mkdir(parents=True, exist_ok=True)
    salt = load_salt()
    mappings = load_yaml("source_mappings.yaml")
    personal = mappings["personal_fields"]

    reader = Reader(raw_dir, settings, load_mappings(tuple(mapping_overlays)))
    reader.read_all()
    log = reader.log
    # Donor tracking entered on the dashboard: read next to the Excel donor tracker (both sources work together).
    app_donors, app_opps = read_app_entries(app_entries) if use_app_entries else ([], [])
    app_acts, app_reps = read_app_programme_entries(app_entries) if use_app_entries else ([], [])
    if app_acts or app_reps:
        reader.records.setdefault("activity", []).extend(app_acts)
        reader.records.setdefault("indicator_result", []).extend(app_reps)
        for sheet, entity, n in (("activities", "activity", len(app_acts)),
                                 ("indicator_reports", "indicator_result", len(app_reps))):
            reader.inventory.append(dict(source_file=APP_SOURCE, source_sheet=sheet, entity=entity, rows=n))
    if use_app_entries:
        # One home for donor data: rows imported into the dashboard are no longer read from the Excel tracker.
        import app_store
        entries = Path(app_entries) if app_entries else app_store.db_path()
        moved_d, moved_o = app_store.imported_ids(entries) if entries.exists() else (set(), set())
        for entity, moved, id_field in (("donor", moved_d, "donor_id"), ("opportunity", moved_o, "opportunity_id")):
            before = reader.records.get(entity, [])
            kept = [r for r in before if r.get("source_file") == APP_SOURCE or str(r.get(id_field)) not in moved]
            if len(kept) < len(before):
                reader.records[entity] = kept
                for (f, sh), n in pd.Series([(r["source_file"], r["source_sheet"]) for r in before if r not in kept]
                                            ).value_counts().items():
                    # A note on the sheet's own inventory row (one row per sheet: the model checks this).
                    note = f"{n} rows already moved to the dashboard (Donor entry): not read"
                    for row in reader.inventory:
                        if row.get("source_file") == f and row.get("source_sheet") == sh:
                            row["note"] = "; ".join(x for x in (row.get("note"), note) if isinstance(x, str) and x)
    if app_donors or app_opps:
        reader.records.setdefault("donor", []).extend(app_donors)
        reader.records.setdefault("opportunity", []).extend(app_opps)
        for sheet, entity, n in (("donors", "donor", len(app_donors)), ("opportunities", "opportunity", len(app_opps))):
            reader.inventory.append(dict(source_file=APP_SOURCE, source_sheet=sheet, entity=entity, rows=n))

    # ---- beneficiaries: match duplicates, then pseudonymise
    regs = reader.records.get("registration", [])
    ids = identity_frame(regs, settings, log)
    review_pairs: list = []
    if not regs:                                       # e.g. checking one uploaded activity file on its own
        ids = ids.assign(beneficiary_key=None, registration_key=None, cluster=None, duplicate_group_size=0,
                         age=None, age_source=None)
    clusters = None if not regs else cluster_identities(ids, settings["duplicate_matching"]["name_similarity_threshold"], settings,
                                  review_pairs=review_pairs)
    if regs:
        ids = assign_keys(ids, clusters, salt, settings["duplicate_matching"]["compare_id_digits"])
    certainty = duplicate_certainty(ids, settings) if regs else {}
    review_partners: dict[int, int] = {}
    for i, j in review_pairs:
        if clusters[i] != clusters[j]:
            review_partners[i] = review_partners.get(i, 0) + 1
            review_partners[j] = review_partners.get(j, 0) + 1
    for rec, row in zip(regs, ids.itertuples()):
        rec["beneficiary_key"], rec["registration_key"] = row.beneficiary_key, row.registration_key
        rec["age"] = None if row.age is None or pd.isna(row.age) else str(row.age)
        rec["age_source"] = row.age_source
        # Not personal: only whether an ID was given and how it was typed (for data-quality measures).
        rec["has_national_id"] = not is_blank(rec.get("national_id"))
        rec["id_stored_as_text"] = rec["has_national_id"] and not isinstance(rec.get("national_id"), (int, float))
        if row.duplicate_group_size > 1 and certainty[row.cluster] == "certain":
            log.add("registration", rec, "name", "duplicate_merged_certain", FIXED,
                    detail=f"same person as {row.duplicate_group_size - 1} other registration(s) "
                           "(same ID, or same phone and name); same beneficiary_key given")
        elif row.duplicate_group_size > 1:
            log.add("registration", rec, "name", "possible_duplicate_person", FLAGGED,
                    detail=f"probably the same person as {row.duplicate_group_size - 1} other registration(s); "
                           "same beneficiary_key given - please confirm")
    for idx, count in review_partners.items():
        log.add("registration", regs[idx], "name", "possible_duplicate_person", FLAGGED,
                detail=f"similar name to {count} other registration(s) in the same district, but no phone or ID "
                       "to confirm it; NOT merged - please check")
    initials_partners: dict[int, int] = {}
    for i, j in (initials_review_pairs(ids, settings) if regs else []):
        if clusters[i] != clusters[j]:
            initials_partners[i] = initials_partners.get(i, 0) + 1
            initials_partners[j] = initials_partners.get(j, 0) + 1
    for idx, count in initials_partners.items():
        log.add("registration", regs[idx], "name", "possible_duplicate_initials", FLAGGED,
                detail=f"initials match the full name of {count} other registration(s) (same surname, district "
                       "and similar age); NOT merged - please check")
    reg_lookup = {(r["source_file"], norm_key(r.get("reg_no") or "")): (r["registration_key"], r["beneficiary_key"])
                  for r in regs}
    for rec in reader.records.get("service", []):
        keys = reg_lookup.get((rec["source_file"], norm_key(rec.get("reg_no") or "")))
        rec["registration_key"], rec["beneficiary_key"] = keys if keys else (None, None)
        if not keys:
            log.add("service", rec, "reg_no", "service_unknown_registration", FLAGGED,
                    detail="registration number not found in this file's register")

    # ---- donors: pseudonymise individuals, drop contact details
    donors = reader.records.get("donor", [])
    individuals = sorted((d for d in donors
                          if recode(d.get("donor_type"), settings["donor_pipeline"]["donor_type_mapping"],
                                    "donor_type").value == "individual"),
                         key=lambda d: (d.get("source_file") == APP_SOURCE,
                                        d["source_row"] if d.get("source_file") == APP_SOURCE else str(d.get("donor_id"))))
    # Donors entered on the dashboard are numbered after the Excel tracker's, in the order they were added (an import
    # adds the tracker's donors in donor-ID order), so existing "Individual donor #NNN" labels never change.
    label = settings["privacy"]["individual_donor_label"]
    name_to_key = {}
    for d in donors:
        d["donor_key"] = pseudonym(salt, "DON", str(d.get("donor_id") or d.get("donor_name")))
        name_to_key[norm_key(d.get("donor_name") or "")] = d["donor_key"]
        if any(not is_blank(d.get(f)) for f in personal["donor"]):
            log.add("donor", d, "contact_person", "personal_data_removed", FIXED,
                    detail="contact name, email and phone removed")
    for number, d in enumerate(individuals, 1):
        d["donor_name"] = label.format(number=number)
        log.add("donor", d, "donor_name", "individual_donor_pseudonymised", FIXED, clean=d["donor_name"])
    for o in reader.records.get("opportunity", []):
        if o.get("source_file") == APP_SOURCE:       # linked by the donor's ID (added on the dashboard) or key
            app_id, app_key = o.pop("app_donor_id", None), o.pop("app_donor_key", None)
            o["donor_key"] = pseudonym(salt, "DON", str(app_id)) if app_id else app_key
        else:
            o["donor_key"] = name_to_key.get(norm_key(o.get("donor") or ""))

    # ---- write staging: text columns only, no personal fields
    written = {}
    for entity, recs in reader.records.items():
        drop = set(personal.get(entity, [])) | ({"donor", "reg_no"} if entity in ("opportunity", "registration",
                                                                               "service") else set())
        df = pd.DataFrame([{k: (to_text(v) if k not in ("source_row",) else v) for k, v in r.items() if k not in drop}
                           for r in recs])
        df.to_parquet(staging / f"ingested_{entity}.parquet", index=False)
        written[entity] = df
    # Every kind of record gets a table, empty if no file held it (e.g. checking one uploaded file on its own).
    for entity, fields in reader.mappings["fields"].items():
        if entity not in written:
            cols = ["source_file", "source_sheet", "source_row"] + [f for f in fields
                                                                   if f not in personal.get(entity, [])]
            empty = pd.DataFrame({c: pd.Series(dtype="object") for c in cols})
            empty.to_parquet(staging / f"ingested_{entity}.parquet", index=False)
            written[entity] = empty
    for entity, df in reader.refs.items():
        df.to_parquet(staging / f"{entity}.parquet", index=False)
    log.frame().to_parquet(staging / "ingest_issues.parquet", index=False)
    pd.DataFrame(reader.skipped, columns=["source_file", "source_sheet", "entity", "reason", "known_columns",
                                          "missing_required", "headers_found", "message"]).to_csv(
        staging / "skipped_sheets.csv", index=False)
    pd.DataFrame(reader.inventory).to_csv(staging / "sheet_inventory.csv", index=False)
    return dict(tables=written, inventory=pd.DataFrame(reader.inventory), issues=log.frame(),
                skipped=pd.DataFrame(reader.skipped))


def main():
    result = run()
    counts = ", ".join(f"{len(df)} {name}" for name, df in result["tables"].items())
    print(f"Ingested {counts}")
    unmapped = result["inventory"].query("unmapped_columns != '' and unmapped_columns.notna()", engine="python") \
        if "unmapped_columns" in result["inventory"] else []
    for row in getattr(unmapped, "itertuples", lambda: [])():
        print(f"  Note: unrecognised columns in {row.source_file} / {row.source_sheet}: {row.unmapped_columns}")


if __name__ == "__main__":
    main()
