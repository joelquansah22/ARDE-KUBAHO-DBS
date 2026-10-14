"""Create realistic, deliberately messy SYNTHETIC data for the prototype.

WHAT THIS SCRIPT DOES
---------------------
1. Builds a fictional "true" world: programmes, projects, villages, activities,
   beneficiaries, services, indicator reports, donors and opportunities.
2. Writes that world into Excel files the way real field staff might: each field
   officer has their own layout, column names and habits, and mistakes are
   injected at the rates set in config/synthetic.yaml.
3. Writes an ANSWER KEY to data/ground_truth/: the true values, where each record
   ended up (file, sheet, Excel row) and a log of every injected problem. The
   pipeline must never read this folder; it is only used to score the pipeline.

HOW TO RUN
----------
    python src/generate_synthetic.py

The same seed always gives exactly the same files.

Everything produced is invented. File names start with SYNTHETIC_ so the data can
never be mistaken for real ARDE/KUBAHO records.
"""

from __future__ import annotations

import argparse
import calendar
import datetime as dt
import hashlib
import math
import sys
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import (
    district_to_province,
    fiscal_year_label,
    load_admin,
    load_settings,
    load_yaml,
    resolve_path,
    to_excel_serial,
)

PREFIX = "SYNTHETIC_"

# --------------------------------------------------------------------------------------
# Fictional name lists. These are common Rwandan given names and surnames combined at
# random, so any resemblance to a real person is coincidental.
# --------------------------------------------------------------------------------------
FEMALE_GIVEN = [
    "Marie", "Claudine", "Diane", "Aline", "Esperance", "Solange", "Josiane", "Vestine",
    "Chantal", "Grace", "Clementine", "Jeanne", "Immaculee", "Beatha", "Odette",
    "Alphonsine", "Francine", "Agnes", "Speciose", "Donatha", "Jacqueline", "Liliane",
    "Sandrine", "Yvonne", "Consolee", "Gaudence", "Seraphine", "Claire", "Angelique", "Pascasie",
]
MALE_GIVEN = [
    "Jean Claude", "Emmanuel", "Eric", "Jean de Dieu", "Patrick", "Innocent", "Theoneste",
    "Olivier", "Fabrice", "Alexis", "Didier", "Celestin", "Felix", "Samuel", "Pacifique",
    "Jean Bosco", "Evariste", "Damascene", "Faustin", "Augustin", "Gaspard", "Venant",
    "Callixte", "Silas", "Valens", "Anastase", "Jean Pierre", "Martin", "Joseph", "Aimable",
]
SECOND_GIVEN_FEMALE = ["Claire", "Rose", "Josee", "Louise", "Therese", "Goretti"]
SECOND_GIVEN_MALE = ["Paul", "Marie", "Baptiste", "Bosco", "Damascene", "Pierre"]
SURNAMES_ANY = [
    "Uwimana", "Niyonzima", "Habimana", "Nshimiyimana", "Uwase", "Mugisha", "Ishimwe",
    "Nsengiyumva", "Ingabire", "Hakizimana", "Tuyisenge", "Niyonsaba", "Bizimana",
    "Uwamahoro", "Munyaneza", "Iradukunda", "Nkurunziza", "Byiringiro", "Uwera", "Gasana",
    "Rukundo", "Kamanzi", "Ndayisaba", "Uwizeye", "Ntirenganya", "Sibomana", "Dusabimana",
    "Niyitegeka", "Tuyishime", "Manirakiza", "Nsabimana", "Twizeyimana", "Harerimana",
    "Mutabazi", "Nzabonimpa", "Ndagijimana", "Habiyaremye", "Uwamariya", "Mukeshimana",
]
SURNAMES_FEMALE = [
    "Mukamana", "Nyirahabimana", "Mukandayisenga", "Murekatete", "Kayitesi", "Mutoni",
    "Nyiraneza", "Mukarukundo", "Nyiransabimana",
]
SURNAMES_MALE = ["Habiyambere", "Nsengimana", "Ntakirutimana", "Hategekimana"]

# Syllables for invented sector, cell and village names (all synthetic).
PLACE_START = ["Ka", "Ki", "Ru", "Nya", "Gi", "Mu", "Bu", "Ri", "Ga", "Ma", "Rwa", "Bi", "Ngo", "Kabu", "Nyaka"]
PLACE_END = ["bare", "honga", "ganza", "rambi", "simbi", "tovu", "sange", "rehe", "hinga", "bona",
             "yaga", "shenyi", "bugu", "tare", "rama", "gera", "kiro", "zinga", "biro", "sasa"]

# Fictional donors. Organisation names are invented; countries are only for realism.
DONORS = [
    ("DNR01", "Bluewater Futures Foundation", "foundation", "United States", "WASH"),
    ("DNR02", "Greenfield Resilience Facility", "multilateral", "International", "Agriculture; Climate"),
    ("DNR03", "Halvorsen Development Partnership", "bilateral", "Norway", "WASH; Livelihoods"),
    ("DNR04", "Calloway Family Foundation", "foundation", "United Kingdom", "Agriculture"),
    ("DNR05", "Ostrander Water Trust", "foundation", "Canada", "WASH"),
    ("DNR06", "Rhinefeld Cooperation Fund", "bilateral", "Germany", "Agriculture; Environment"),
    ("DNR07", "Meridian Climate Adaptation Window", "multilateral", "International", "Climate; Environment"),
    ("DNR08", "Bellamy Rural Futures Fund", "foundation", "United States", "Livelihoods; Agriculture"),
    ("DNR09", "Lakeside Microfinance Community Fund", "private", "Rwanda", "Livelihoods"),
    ("DNR10", "Sunrise Agro-Processing Ltd", "private", "Rwanda", "Agriculture"),
    ("DNR11", "Riverside Municipal Twinning Grant", "government", "Belgium", "WASH"),
    ("DNR12", "Fenwick Charitable Trust", "foundation", "United Kingdom", "WASH; Environment"),
    ("DNR13", "Northstar Solar Cooperative Fund", "private", "Netherlands", "Agriculture; Environment"),
    ("DNR14", "Eastbrook Faith Alliance", "foundation", "United States", "WASH; Livelihoods"),
    ("DNR15", "Aldmere Bilateral Aid Programme", "bilateral", "Sweden", "Environment"),
    ("DNR16", "Kestrel Impact Partners", "private", "Kenya", "Livelihoods"),
    ("DNR17", "Corvane Global Health Fund", "multilateral", "International", "WASH"),
    ("DNR18", "Community Infrastructure Matching Grant", "government", "Rwanda", "WASH; Agriculture"),
    ("DNR19", "Wrenfield Women's Fund", "foundation", "Ireland", "Livelihoods"),
    ("DNR20", "Mr Thomas Whitcombe", "individual", "United Kingdom", "WASH"),
    ("DNR21", "Ms Ingrid Solberg", "individual", "Norway", "Environment"),
    ("DNR22", "Dr Amani Okafor", "individual", "United States", "Agriculture"),
    ("DNR23", "Mrs Claudine Mukasine", "individual", "Rwanda", "WASH; Livelihoods"),
    ("DNR24", "Mr Daniel Ferreira", "individual", "Portugal", "Agriculture"),
    ("DNR25", "Ms Grace Uwitonze", "individual", "Canada", "Livelihoods"),
]
CONTACT_NAMES = ["A. Lindqvist", "R. Patel", "S. Moreau", "J. Okonkwo", "L. Brennan", "K. Haugen",
                 "M. Duarte", "P. Kowalski", "T. Nakamura", "E. Van Dijk"]
DONOR_TYPE_VARIANTS = {
    "bilateral": ["Bilateral donor", "Bi-lateral"],
    "multilateral": ["Multi-lateral", "UN / multilateral"],
    "foundation": ["Private foundation", "Trust/Foundation"],
    "government": ["Govt", "Local government"],
    "private": ["Private sector", "Company"],
    "individual": ["Individual (diaspora)", "Private individual"],
}
STAGE_VARIANTS = {
    "Identified": ["Idea", "Prospect"],
    "Qualified": ["Eligible", "Go decision"],
    "Concept note": ["CN", "CN submitted", "concept note"],
    "Proposal submitted": ["Proposal sent", "Submitted", "Full proposal"],
    "Awarded": ["Won", "Grant signed"],
    "Declined": ["Lost", "Rejected", "Unsuccessful"],
}
STAGE_PROBABILITY = {"Identified": 0.10, "Qualified": 0.25, "Concept note": 0.40,
                     "Proposal submitted": 0.60, "Awarded": 1.0, "Declined": 0.0}
NEXT_ACTIONS = ["Draft concept note", "Call programme officer", "Collect budget inputs",
                "Submit due diligence documents", "Follow up on decision", "Arrange site visit",
                "Prepare budget narrative", "Request letter of support"]
DIAL_CODES = ["1", "44", "47", "49", "31", "46", "32", "351", "353", "254", "250"]
STAFF_OWNERS = ["ED01", "PM01", "PM02", "MEAL01", "FIN01"]
REPORT_COMMENTS = ["late submission", "figures to be verified", "includes catch-up from last month",
                   "awaiting sector data", ""]

# Held-out set only: nicknames, and old (pre-2006) names still used for some districts.
NICKNAMES = {"Jean Claude": "JC", "Emmanuel": "Manu", "Jean de Dieu": "Jado", "Innocent": "Inno",
             "Jean Bosco": "Bosco", "Pacifique": "Paci", "Olivier": "Oli", "Josiane": "Josy",
             "Esperance": "Espe", "Clementine": "Clemy", "Immaculee": "Imma", "Jacqueline": "Jacky",
             "Theoneste": "Theo", "Alexis": "Alex", "Samuel": "Sam", "Patrick": "Pato", "Francine": "Fanny",
             "Liliane": "Lili", "Fabrice": "Fab", "Jean Pierre": "JP", "Marie": "Mary", "Aline": "Alina"}
OLD_DISTRICT_NAMES = {"Huye": "Butare", "Nyamagabe": "Gikongoro", "Karongi": "Kibuye"}

DATE_ISSUE = {
    "d_mon_yy": ("date_new_format", "fix"),        # held-out set only: 12 Mar 24
    "d_month_yyyy": ("date_new_format", "fix"),    # held-out set only: 12 March 2024
    "dotted": ("date_new_format", "fix"),          # held-out set only: 12.03.2024
    "dmy": ("date_dmy_text", "fix"),
    "month_year": ("date_month_year_text", "fix_or_flag"),
    "serial": ("date_excel_serial", "fix"),
}


# ======================================================================================
# Small building blocks
# ======================================================================================
class Rand:
    """A thin wrapper round numpy's random generator with friendlier method names."""

    def __init__(self, seed: int):
        self.g = np.random.default_rng(seed)

    def chance(self, p: float) -> bool:
        return bool(self.g.random() < p)

    def pick(self, seq, weights=None):
        seq = list(seq)
        if weights is not None:
            w = np.asarray(weights, dtype=float)
            return seq[int(self.g.choice(len(seq), p=w / w.sum()))]
        return seq[int(self.g.integers(len(seq)))]

    def sample(self, seq, k: int) -> set:
        """Pick exactly k different items (used where coin-flips on a small group could give zero)."""
        seq = list(seq)
        k = min(k, len(seq))
        return {seq[i] for i in self.g.choice(len(seq), size=k, replace=False)}

    def pick_style(self, styles: dict[str, float]) -> str:
        return self.pick(list(styles), list(styles.values()))

    def int(self, lo: int, hi: int) -> int:
        """Whole number between lo and hi, both included."""
        return int(self.g.integers(lo, hi + 1))

    def uniform(self, lo: float, hi: float) -> float:
        return float(self.g.uniform(lo, hi))

    def normal(self, mean: float, sd: float) -> float:
        return float(self.g.normal(mean, sd))

    def poisson(self, lam: float) -> int:
        return int(self.g.poisson(lam))

    def binomial(self, n: int, p: float) -> int:
        return int(self.g.binomial(n, p))


def month_starts(start: dt.date, end: dt.date):
    """Yield the first day of every month from start to end."""
    current = dt.date(start.year, start.month, 1)
    while current <= end:
        yield current
        current = dt.date(current.year + (current.month == 12), current.month % 12 + 1, 1)


def days_in_month(d: dt.date) -> int:
    return calendar.monthrange(d.year, d.month)[1]


def render_date(d: dt.date | None, style: str):
    """Write a date the way a particular officer would type it into Excel."""
    if d is None:
        return None
    if style == "datetime":
        return dt.datetime(d.year, d.month, d.day)   # a real Excel date cell
    if style == "iso":
        return d.isoformat()                          # text 2024-03-12
    if style == "dmy":
        return d.strftime("%d/%m/%Y")                 # text 12/03/2024
    if style == "month_year":
        return f"{calendar.month_name[d.month]} {d.year}"  # text March 2024 (day lost)
    if style == "serial":
        return to_excel_serial(d)                     # plain number 45363
    if style == "d_mon_yy":
        return f"{d.day} {calendar.month_abbr[d.month]} {d:%y}"     # 12 Mar 24
    if style == "d_month_yyyy":
        return f"{d.day} {calendar.month_name[d.month]} {d.year}"   # 12 March 2024
    if style == "dotted":
        return d.strftime("%d.%m.%Y")                               # 12.03.2024
    raise ValueError(style)


def year_typo(d: dt.date) -> dt.date:
    """Simulate a mistyped year, e.g. 2024 -> 2042 or 2025 -> 2052."""
    y = str(d.year)
    new_year = int(y[:2] + y[3] + y[2]) if y[2] != y[3] else d.year + 10
    return dt.date(new_year, d.month, min(d.day, calendar.monthrange(new_year, d.month)[1]))


def letter_slip(rand: Rand, word: str) -> str:
    """Introduce one realistic spelling slip (swap, drop, double or vowel change)."""
    for _ in range(20):
        i = rand.int(1, max(1, len(word) - 2))
        op = rand.pick(["swap", "drop", "double", "vowel"])
        if op == "swap" and i < len(word) - 1:
            new = word[:i] + word[i + 1] + word[i] + word[i + 2:]
        elif op == "drop":
            new = word[:i] + word[i + 1:]
        elif op == "double":
            new = word[:i] + word[i] + word[i:]
        else:
            vowels = "aeiou"
            if word[i].lower() not in vowels:
                continue
            new = word[:i] + rand.pick([v for v in vowels if v != word[i].lower()]) + word[i + 1:]
        if new.lower() != word.lower():
            return new
    return word + word[-1]


@dataclass
class Cells:
    """The cell values for one Excel row, plus the problems deliberately put into it."""

    values: dict
    issues: list = field(default_factory=list)

    def note(self, fld, error_type, true_value, messy_value, expected, detail=""):
        self.issues.append(dict(field=fld, error_type=error_type, true_value=true_value,
                                messy_value=messy_value, expected_handling=expected, detail=detail))

    def blank(self, fld, true_value):
        """Empty a cell and record it as a missing required value."""
        self.issues = [i for i in self.issues if i["field"] != fld]
        self.values[fld] = None
        self.note(fld, "missing_required", true_value, None, "flag")


class InjectionLog:
    """Every injected problem, with the exact file, sheet, row and column it is in."""

    def __init__(self):
        self.rows: list[dict] = []

    def add(self, **kw):
        row = dict(source_file=None, source_sheet=None, source_row=None, source_column=None,
                   column_letter=None, field=None, entity=None, record_id=None, error_type=None,
                   true_value=None, messy_value=None, expected_handling=None, detail="")
        row.update(kw)
        self.rows.append(row)


class SheetWriter:
    """Writes rows top to bottom into one worksheet and logs where every problem landed."""

    HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")

    def __init__(self, ws, file_name: str, entity: str, log: InjectionLog):
        self.ws, self.file_name, self.entity, self.log = ws, file_name, entity, log
        self.row = 0
        self.fields: list[str] = []
        self.headers: list[str] = []
        self.formats: dict[str, str] = {}
        self.header_row = None

    def _log(self, **kw):
        self.log.add(source_file=self.file_name, source_sheet=self.ws.title, entity=self.entity, **kw)

    def title(self, text: str, width: int):
        self.row += 1
        cell = self.ws.cell(self.row, 1, text)
        cell.font = Font(bold=True, size=13)
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=width)
        self._log(source_row=self.row, error_type="title_row", expected_handling="fix",
                  detail="title above the header row")

    def blank_row(self, logged: bool = True):
        self.row += 1
        if logged:
            self._log(source_row=self.row, error_type="blank_row", expected_handling="fix")

    def group_header(self, groups: list[tuple[str, int]]):
        self.row += 1
        col = 1
        for label, span in groups:
            cell = self.ws.cell(self.row, col, label or None)
            cell.font, cell.fill = Font(bold=True), self.HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
            if span > 1:
                self.ws.merge_cells(start_row=self.row, start_column=col, end_row=self.row,
                                    end_column=col + span - 1)
            col += span
        self._log(source_row=self.row, error_type="merged_header_row", expected_handling="fix",
                  detail="two-row header with merged group labels")

    def header(self, columns: list[tuple[str, str]], formats: dict[str, str] | None = None):
        self.row += 1
        self.fields = [f for f, _ in columns]
        self.headers = [h for _, h in columns]
        self.formats = formats or {}
        self.header_row = self.row
        for i, h in enumerate(self.headers, 1):
            cell = self.ws.cell(self.row, i, h)
            cell.font, cell.fill = Font(bold=True), self.HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    def record(self, cells: Cells, record_id: str, only_issue: dict | None = None) -> int:
        """Write one data row. Returns the Excel row number it was written to."""
        self.row += 1
        for i, fld in enumerate(self.fields, 1):
            value = cells.values.get(fld)
            cell = self.ws.cell(self.row, i, value)
            if isinstance(value, dt.datetime):
                cell.number_format = self.formats.get(fld, "dd/mm/yyyy")
            elif fld in self.formats and isinstance(value, (int, float)):
                cell.number_format = self.formats[fld]
        issues = [only_issue] if only_issue else cells.issues
        for issue in issues:
            fld = issue["field"]
            col = self.fields.index(fld) + 1 if fld in self.fields else None
            details = {k: v for k, v in issue.items() if k != "field"}
            if col and not only_issue:
                # A cell can carry two problems (e.g. negative AND typed as text): always log
                # what is finally in the cell.
                details["messy_value"] = cells.values.get(fld)
            self._log(source_row=self.row,
                      source_column=self.headers[col - 1] if col else None,
                      column_letter=get_column_letter(col) if col else None,
                      field=fld if col else None, record_id=record_id, **details)
        return self.row

    def totals(self, label: str, sum_fields: list[str]):
        """Add a totals row at the bottom, as staff often do (with live SUM formulas)."""
        if self.row == self.header_row:
            return
        first, last = self.header_row + 1, self.row
        self.row += 1
        self.ws.cell(self.row, 1, label).font = Font(bold=True)
        for fld in sum_fields:
            if fld in self.fields:
                letter = get_column_letter(self.fields.index(fld) + 1)
                cell = self.ws.cell(self.row, self.fields.index(fld) + 1, f"=SUM({letter}{first}:{letter}{last})")
                cell.font = Font(bold=True)
        self._log(source_row=self.row, error_type="totals_row", expected_handling="fix",
                  detail=f"'{label}' row with SUM formulas")

    def finish(self):
        """Set readable column widths so the file looks like a real working sheet."""
        for i, h in enumerate(self.headers, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = min(32, max(10, len(str(h)) + 4))


# ======================================================================================
# 1. Build the "true" world
# ======================================================================================
class World:
    """Holds the true (clean) version of every synthetic record."""

    def __init__(self, cfg: dict, settings: dict, rand: Rand):
        self.cfg, self.settings, self.rand = cfg, settings, rand
        self.window_start = settings["reporting_window"]["start"]
        self.window_end = settings["reporting_window"]["end"]
        self.as_of = settings["as_of_date"] or dt.date.today()
        self.d2p = district_to_province()
        self.programmes = cfg["programmes"]
        self.projects = {p["project_id"]: p for p in cfg["projects"]}
        self.activities = {a["activity_id"]: a for a in cfg["activities"]}
        self.officer_of = {d: fo for fo, spec in cfg["field_officers"].items() for d in spec["districts"]}
        self.districts = list(dict.fromkeys(d for p in cfg["projects"] for d in p["districts"]))
        self.indicators = {i["code"]: i for i in load_yaml("indicator_catalogue.yaml")["indicators"]}
        self.build_locations()
        self.build_activity_records()
        self.build_beneficiaries()
        self.build_services()
        self.build_indicator_reports()
        self.build_donors()

    # ---------------------------------------------------------------- locations
    def build_locations(self):
        used: set[str] = set()

        def place():
            while True:
                name = self.rand.pick(PLACE_START) + self.rand.pick(PLACE_END)
                if name not in used:
                    used.add(name)
                    return name

        self.locations, n = [], 0
        for district in self.districts:
            for _ in range(3):
                sector = place()
                for _ in range(2):
                    cell = place()
                    for _ in range(3):
                        n += 1
                        self.locations.append(dict(
                            location_id=f"LOC{n:04d}", province=self.d2p[district], district=district,
                            sector=sector, cell=cell, village=place()))
        self.locs_by_district = {d: [l for l in self.locations if l["district"] == d] for d in self.districts}

    # ---------------------------------------------------------------- activity records
    def _project_months(self, project):
        start = max(project["start_date"], self.window_start)
        end = min(project["end_date"], self.window_end)
        return list(month_starts(start, end))

    def build_activity_records(self):
        rates, vol = self.cfg["injection_rates"], self.cfg["volumes"]
        tol = self.settings["data_quality"]["on_time_tolerance_days"]
        combos = [(p, d, m) for p in self.projects.values() for d in p["districts"] for m in self._project_months(p)]
        lam = vol["activity_records"] / len(combos)
        records = []
        for project, district, month in combos:
            acts = [a for a in self.activities.values() if a["project_id"] == project["project_id"]]
            weights = [0.6 if a["activity_type"] == "construction" else 1.0 for a in acts]
            for _ in range(self.rand.poisson(lam)):
                act = self.rand.pick(acts, weights)
                loc = self.rand.pick(self.locs_by_district[district])
                planned = month + dt.timedelta(days=self.rand.int(0, days_in_month(month) - 1))
                roll = self.rand.uniform(0, 1)
                if roll < rates["cancelled_activity"]:
                    status, actual = "cancelled", None
                elif roll < rates["cancelled_activity"] + rates["delayed_activity"]:
                    status, actual = "delayed", planned + dt.timedelta(days=self.rand.int(tol + 1, 45))
                else:
                    status, actual = "completed", planned + dt.timedelta(days=self.rand.int(-3, 5))
                if actual and actual > self.window_end:
                    # Completed late in the month: keep it inside the window. Delayed past the
                    # window: still not done when reporting closed (no date, no participants).
                    actual = None if status == "delayed" else self.window_end
                on_paper = self.rand.chance(vol["paper_share_activities"])
                if on_paper:
                    lag = self.rand.int(7, 45)          # paper is typed up in batches, often late
                elif self.rand.chance(rates["late_entry"]):
                    lag = self.rand.int(31, 120)
                else:
                    lag = self.rand.int(0, 14)
                entry = (actual or planned) + dt.timedelta(days=lag)
                if status == "cancelled":
                    female = male = total = cost = 0
                elif actual is None:
                    female = male = total = cost = None
                else:
                    total = self.rand.int(*act["participants"])
                    female = self.rand.binomial(total, act["female_share"])
                    male = total - female
                    cost = int(round(self.rand.uniform(*act["cost_rwf"]), -3))
                records.append(dict(
                    project_id=project["project_id"], activity_id=act["activity_id"],
                    location_id=loc["location_id"], district=district, planned_date=planned, date=actual,
                    entry_date=entry, status=status, participants_female=female,
                    participants_male=male, participants_total=total, cost_rwf=cost,
                    entered_by=self.cfg["paper_data_entry_clerk"] if on_paper else self.officer_of[district],
                    field_officer=self.officer_of[district], on_paper=on_paper))
        # A few activities planned after the reporting window (status "planned").
        for project in self.projects.values():
            if project["end_date"] <= self.window_end:
                continue
            acts = [a for a in self.activities.values() if a["project_id"] == project["project_id"]]
            for district in project["districts"]:
                for month in list(month_starts(self.window_end + dt.timedelta(days=1),
                                               self.window_end + dt.timedelta(days=60))):
                    for _ in range(self.rand.int(1, 2)):
                        act = self.rand.pick(acts)
                        loc = self.rand.pick(self.locs_by_district[district])
                        records.append(dict(
                            project_id=project["project_id"], activity_id=act["activity_id"],
                            location_id=loc["location_id"], district=district,
                            planned_date=month + dt.timedelta(days=self.rand.int(0, 27)), date=None,
                            entry_date=self.window_end - dt.timedelta(days=self.rand.int(0, 20)),
                            status="planned", participants_female=None, participants_male=None,
                            participants_total=None, cost_rwf=None, entered_by=self.officer_of[district],
                            field_officer=self.officer_of[district], on_paper=False))
        records.sort(key=lambda r: (r["planned_date"], r["district"]))
        late_days = self.settings["data_quality"]["late_entry_days"]
        for i, r in enumerate(records, 1):
            r["record_id"] = f"ACR{i:05d}"
            r["entry_lag_days"] = (r["entry_date"] - (r["date"] or r["planned_date"])).days
            # Late entry is measured from the activity date, so only activities that took place count.
            r["is_late_entry"] = r["entry_lag_days"] > late_days and r["date"] is not None
            r["on_time"] = (r["status"] == "completed") if r["status"] != "planned" else None
        self.activity_records = records

    # ---------------------------------------------------------------- beneficiaries
    def _person_name(self, sex):
        given = self.rand.pick(FEMALE_GIVEN if sex == "female" else MALE_GIVEN)
        if self.rand.chance(0.3):
            second = self.rand.pick(SECOND_GIVEN_FEMALE if sex == "female" else SECOND_GIVEN_MALE)
            if second not in given:
                given = f"{given} {second}"
        pool = SURNAMES_ANY + (SURNAMES_FEMALE if sex == "female" else SURNAMES_MALE)
        return self.rand.pick(pool), given

    def build_beneficiaries(self):
        vol, rates = self.cfg["volumes"], self.cfg["injection_rates"]
        projects = list(self.projects.values())
        phones, people = set(), []
        for i in range(1, vol["beneficiaries"] + 1):
            project = self.rand.pick(projects, [len(p["districts"]) for p in projects])
            district = self.rand.pick(project["districts"])
            loc = self.rand.pick(self.locs_by_district[district])
            female_p = 0.5 if project["programme_id"] == "PRG02" else 0.56
            sex = "female" if self.rand.chance(female_p) else "male"
            surname, given = self._person_name(sex)
            age = int(min(80, max(18, round(self.rand.normal(38, 12)))))
            first_month = max(project["start_date"], self.window_start)
            last_day = min(project["end_date"], self.window_end) - dt.timedelta(days=90)
            reg_date = first_month + dt.timedelta(days=self.rand.int(0, (last_day - first_month).days))
            dob = dt.date(reg_date.year - age, 1, 1) + dt.timedelta(days=self.rand.int(0, 364))
            while True:
                phone = "07" + self.rand.pick("8923") + "".join(str(self.rand.int(0, 9)) for _ in range(7))
                if phone not in phones:
                    phones.add(phone)
                    break
            national_id = ("1" + str(dob.year) + ("7" if sex == "female" else "8")
                           + "".join(str(self.rand.int(0, 9)) for _ in range(7)) + "0"
                           + f"{self.rand.int(0, 99):02d}")
            people.append(dict(
                person_id=f"PER{i:05d}", project_id=project["project_id"], district=district,
                location_id=loc["location_id"], sex=sex, surname=surname, given_name=given,
                age=age, date_of_birth=dob, phone=phone, national_id=national_id,
                household_size=int(min(15, max(1, self.rand.poisson(4.5)))),
                disability=self.rand.chance(0.06), registration_date=reg_date,
                on_paper=self.rand.chance(vol["paper_share_beneficiaries"])))
        hm = self.cfg.get("holdout_mess") or {}
        if hm:
            people += self.namesakes(people, hm, phones)
        self.people = people

        # Registration rows: one per person, plus a second row for duplicated people.
        regs = []
        for p in people:
            home = "PAPER" if p["on_paper"] else self.officer_of[p["district"]]
            regs.append(dict(person_id=p["person_id"], file_key=home, is_duplicate=False,
                             surname=p["surname"], given_name=p["given_name"],
                             registration_date=p["registration_date"], phone=p["phone"],
                             national_id=p["national_id"], age=p["age"], dup_change=""))
            if self.rand.chance(rates["duplicate_beneficiary"]):
                other = "PAPER" if home != "PAPER" else self.officer_of[p["district"]]
                where = home if self.rand.chance(0.6) else other
                surname, given, change = p["surname"], p["given_name"], ""
                if hm:
                    surname, given, change = self.harder_variant(p)
                    phone_roll = self.rand.uniform(0, 1)
                    phone = (None if (phone_roll < hm["dup_phone_missing"] or not p["phone"]) else p["phone"]
                             if phone_roll < 0.85 else p["phone"][:-1] + str((int(p["phone"][-1]) + 1) % 10))
                else:
                    op = self.rand.pick(["surname_slip", "given_slip", "order_swap", "surname_slip"])
                    if op == "surname_slip":
                        surname, change = letter_slip(self.rand, surname), "surname spelling"
                    elif op == "given_slip":
                        given, change = letter_slip(self.rand, given), "given-name spelling"
                    else:
                        change = "name order swapped"
                    phone_roll = self.rand.uniform(0, 1)
                    phone = (p["phone"] if phone_roll < 0.6 else None if phone_roll < 0.8
                             else p["phone"][:-1] + str((int(p["phone"][-1]) + 1) % 10))
                room = max(1, min(300, (self.window_end - p["registration_date"]).days))
                later = p["registration_date"] + dt.timedelta(days=self.rand.int(min(20, room), room))
                regs.append(dict(person_id=p["person_id"], file_key=where, is_duplicate=True,
                                 surname=surname, given_name=given,
                                 registration_date=min(later, self.window_end), phone=phone,
                                 national_id=p["national_id"] if self.rand.chance(
                                     hm["dup_id_kept"] if hm else 0.5) else None,
                                 age=p["age"] + self.rand.pick([0, 0, 1, -1]), dup_change=change))
        for i, r in enumerate(regs, 1):
            r["registration_id"] = f"REG{i:05d}"
        self.registrations = regs

    def harder_variant(self, p: dict) -> tuple[str, str, str]:
        """Held-out set only: harder ways the same person's name is written twice."""
        surname, given = p["surname"], p["given_name"]
        op = self.rand.pick(["swap_first_last", "missing_letters", "nickname", "nickname", "surname_slip"])
        first = given.split(" ")[0] if given not in NICKNAMES else given
        if op == "nickname" and (given in NICKNAMES or first in NICKNAMES):
            nick = NICKNAMES.get(given) or NICKNAMES[first]
            return surname, nick, f"nickname ({given} written as {nick})"
        if op == "swap_first_last":
            return given, surname, "first and last names swapped between columns"
        if op in ("missing_letters", "nickname"):
            for _ in range(2):
                i = self.rand.int(1, len(surname) - 2)
                surname = surname[:i] + surname[i + 1:]
            return surname, given, "two letters missing from the surname"
        return letter_slip(self.rand, surname), given, "surname spelling"

    def namesakes(self, people: list[dict], hm: dict, phones: set) -> list[dict]:
        """Held-out set only: DIFFERENT people with exactly the same name, district and sex.

        Some have a similar age and no phone or ID, so only careful matching keeps them apart.
        """
        extra = []
        n = len(people)
        for p in list(people):
            if not self.rand.chance(hm["namesake_rate"]):
                continue
            n += 1
            close_age = self.rand.chance(hm["namesake_close_age"])
            age = (p["age"] + self.rand.pick([-2, -1, 1, 2])) if close_age else self.rand.int(18, 80)
            age = min(80, max(18, age))
            loc = p["location_id"] if self.rand.chance(0.5) else self.rand.pick(
                self.locs_by_district[p["district"]])["location_id"]
            phone = None
            if not self.rand.chance(hm["namesake_phone_missing"]):
                while True:
                    phone = "07" + self.rand.pick("8923") + "".join(str(self.rand.int(0, 9)) for _ in range(7))
                    if phone not in phones:
                        phones.add(phone)
                        break
            reg_date = p["registration_date"] + dt.timedelta(days=self.rand.int(-60, 60))
            reg_date = min(max(reg_date, self.window_start), self.window_end)
            dob = dt.date(reg_date.year - age, 1, 1) + dt.timedelta(days=self.rand.int(0, 364))
            national_id = None if self.rand.chance(hm["namesake_id_missing"]) else (
                "1" + str(dob.year) + ("7" if p["sex"] == "female" else "8")
                + "".join(str(self.rand.int(0, 9)) for _ in range(7)) + "0" + f"{self.rand.int(0, 99):02d}")
            extra.append(dict(p, person_id=f"PER{n:05d}", location_id=loc, age=age, date_of_birth=dob,
                              phone=phone, national_id=national_id, registration_date=reg_date,
                              household_size=int(min(15, max(1, self.rand.poisson(4.5)))),
                              disability=self.rand.chance(0.06), namesake_of=p["person_id"],
                              on_paper=self.rand.chance(self.cfg["volumes"]["paper_share_beneficiaries"])))
        return extra

    def build_services(self):
        lo, hi = self.cfg["volumes"]["services_per_beneficiary"]
        regs_by_person: dict[str, list] = {}
        for r in self.registrations:
            regs_by_person.setdefault(r["person_id"], []).append(r)
        services, n = [], 0
        for p in self.people:
            acts = [a for a in self.activities.values() if a["project_id"] == p["project_id"]]
            end = min(self.projects[p["project_id"]]["end_date"], self.window_end)
            for _ in range(self.rand.int(lo, hi)):
                n += 1
                act = self.rand.pick(acts)
                reg = regs_by_person[p["person_id"]][-1] if (
                    len(regs_by_person[p["person_id"]]) > 1 and self.rand.chance(0.3)) \
                    else regs_by_person[p["person_id"]][0]
                start = reg["registration_date"]
                date = start + dt.timedelta(days=self.rand.int(0, max(0, (end - start).days)))
                services.append(dict(service_id=f"SRV{n:05d}", person_id=p["person_id"],
                                     registration_id=reg["registration_id"], file_key=reg["file_key"],
                                     activity_id=act["activity_id"], date=date,
                                     service_type=act["service_type"]))
        self.services = services

    # ---------------------------------------------------------------- indicator reports
    def build_indicator_reports(self):
        targets = self.cfg["indicator_targets"]
        rows, expected = [], []
        for project in self.projects.values():
            for district in project["districts"]:
                performance = self.rand.uniform(0.62, 1.12)   # some districts do better than others
                fy_targets: dict[tuple, float] = {}
                for month in self._project_months(project):
                    due = []
                    for code in project["indicators"]:
                        freq = self.indicators[code]["frequency"]
                        if (freq == "monthly" or (freq == "quarterly" and month.month in (3, 6, 9, 12))
                                or (freq == "semi-annual" and month.month in (6, 12))):
                            due.append(code)
                    if not due:
                        continue
                    expected.append(dict(project_id=project["project_id"], district=district,
                                         period=month, indicators_due=", ".join(due),
                                         reporter=self.officer_of[district]))
                    fy = fiscal_year_label(month)
                    for code in due:
                        spec = targets[code]
                        key = (code, fy)
                        if key not in fy_targets:   # targets are set once per fiscal year
                            if "percent" in spec:
                                fy_targets[key] = round(self.rand.uniform(*spec["percent"]))
                            elif "per_month" in spec:
                                fy_targets[key] = self.rand.int(*spec["per_month"])
                            else:
                                fy_targets[key] = self.rand.int(*spec["per_quarter"])
                        target = fy_targets[key]
                        for sex in (["Female", "Male"] if self.indicators[code]["disaggregation"] == "sex" else [None]):
                            factor = max(0.0, self.rand.normal(performance, 0.18))
                            if self.indicators[code]["unit"] == "percent":
                                actual = round(min(100.0, target * factor), 1)
                            elif self.indicators[code]["unit"] == "hectares":
                                actual = round(target * factor, 1)
                            else:
                                actual = int(round(target * factor))
                            rows.append(dict(project_id=project["project_id"],
                                             programme_id=project["programme_id"], district=district,
                                             period=month, fiscal_year=fy, indicator_code=code,
                                             sex=sex, target_value=target, actual_value=actual,
                                             entered_by=self.officer_of[district]))
        for i, r in enumerate(rows, 1):
            r["result_id"] = f"IRR{i:05d}"
        self.indicator_rows = rows
        self.expected_reports = expected

    # ---------------------------------------------------------------- donors
    def build_donors(self):
        rate = self.settings["currency"]["usd_to_rwf"]
        self.donors = [dict(donor_id=d[0], donor_name=d[1], donor_type=d[2], country=d[3],
                            thematic_interests=d[4]) for d in DONORS[: self.cfg["volumes"]["donors"]]]
        stages = list(STAGE_PROBABILITY)
        stage_weights = [6, 5, 6, 6, 4, 3]
        themes = {"WASH": ["Rural water supply scale-up", "WASH in schools", "Sanitation marketing",
                           "Water user committee strengthening", "Hygiene behaviour change"],
                  "Agriculture": ["Climate-smart agriculture", "Solar irrigation expansion",
                                  "Cooperative market access", "Seed system resilience",
                                  "Women farmers' livelihoods"]}
        opps = []
        for i in range(1, self.cfg["volumes"]["opportunities"] + 1):
            donor = self.rand.pick(self.donors)
            fit = ("WASH" if "WASH" in donor["thematic_interests"] else "Agriculture"
                   if any(t in donor["thematic_interests"] for t in ("Agriculture", "Climate", "Livelihoods"))
                   else self.rand.pick(["WASH", "Agriculture"]))
            if self.rand.chance(0.15):
                fit = "WASH; Agriculture"
            stage = self.rand.pick(stages, stage_weights)
            title = f"{self.rand.pick(themes[fit.split(';')[0]])} - {self.rand.pick(self.districts)}"
            individual = donor["donor_type"] == "individual"
            requested = (self.rand.int(2, 20) * 1_000_000 if individual
                         else self.rand.int(30, 600) * 1_000_000)
            if stage in ("Proposal submitted", "Awarded", "Declined"):
                deadline = self.as_of - dt.timedelta(days=self.rand.int(10, 300))
            else:
                deadline = self.as_of + dt.timedelta(days=self.rand.int(5, 150))
            prob = STAGE_PROBABILITY[stage]
            if 0 < prob < 1:
                prob = round(min(0.95, max(0.05, prob + self.rand.pick([-0.05, 0, 0, 0.05]))), 2)
            awarded = (int(round(requested * self.rand.uniform(0.6, 1.0), -6)) if stage == "Awarded" else None)
            if stage == "Declined":
                next_action, next_date = None, None
            else:
                next_action = self.rand.pick(NEXT_ACTIONS) if stage != "Awarded" else "Prepare inception report"
                next_date = self.as_of + dt.timedelta(days=self.rand.int(-35, 45))
            opps.append(dict(opportunity_id=f"OPP{i:03d}", donor_id=donor["donor_id"],
                             donor_name=donor["donor_name"], title=title, thematic_fit=fit,
                             stage=stage, deadline=deadline, requested_amount_rwf=requested,
                             awarded_amount_rwf=awarded, probability=prob,
                             owner=self.rand.pick(STAFF_OWNERS), next_action=next_action,
                             next_action_date=next_date,
                             is_overdue=bool(next_date and next_date < self.as_of),
                             usd_equivalent=round(requested / rate)))
        self.opportunities = opps


# ======================================================================================
# 2. Write the messy Excel files
# ======================================================================================
ACTIVITY_LAYOUTS = {
    "A": dict(
        title=None, groups=None, per_fy_sheets=False, project_label="name", status_style="title",
        date_styles={"datetime": 0.85, "dmy": 0.15}, entry_styles={"datetime": 0.85, "dmy": 0.15},
        cost_always_text=False, totals_label="TOTAL", kinyarwanda=False,
        columns=[("date", "Date"), ("planned_date", "Planned date"), ("entry_date", "Date entered"),
                 ("project", "Project"), ("activity", "Activity"), ("district", "District"),
                 ("sector", "Sector"), ("cell", "Cell"), ("village", "Village"), ("status", "Status"),
                 ("female", "Female"), ("male", "Male"), ("total", "Total"), ("cost", "Cost (RWF)"),
                 ("officer", "Officer")]),
    "B": dict(
        title="Activity monitoring log - {officer} - {districts} (SYNTHETIC)",
        groups=[("Activity details", 5), ("Location", 5), ("", 1), ("Participants", 3), ("Finance", 1), ("", 1)],
        per_fy_sheets=False, project_label="code", status_style="lower",
        date_styles={"iso": 0.85, "month_year": 0.10, "dmy": 0.05}, entry_styles={"iso": 0.9, "dmy": 0.1},
        cost_always_text=False, totals_label="Grand total", kinyarwanda=False,
        columns=[("date", "Activity date"), ("planned_date", "Planned for"), ("entry_date", "Entry date"),
                 ("project", "Project code"), ("activity", "Activity"), ("province", "Province"),
                 ("district", "District"), ("sector", "Sector"), ("cell", "Cell"), ("village", "Village"),
                 ("status", "Status"), ("female", "F"), ("male", "M"), ("total", "Total"),
                 ("cost", "Cost"), ("officer", "Entered by")]),
    "C": dict(
        title=None, groups=None, per_fy_sheets=False, project_label="name", status_style="kinyarwanda",
        date_styles={"serial": 0.9, "datetime": 0.1}, entry_styles={"serial": 0.9, "datetime": 0.1},
        cost_always_text=True, totals_label="IGITERANYO", kinyarwanda=True,
        columns=[("date", "Itariki y'igikorwa"), ("planned_date", "Planned date"), ("entry_date", "Entry date"),
                 ("project", "Umushinga"), ("activity", "Igikorwa"), ("district", "Akarere"),
                 ("sector", "Umurenge"), ("cell", "Akagari"), ("village", "Umudugudu"), ("status", "Status"),
                 ("female", "Abagore"), ("male", "Abagabo"), ("total", "Bose"),
                 ("cost", "Amafaranga (RWF)"), ("officer", "Officer")]),
    "D": dict(
        title="Activity report {fy} - {districts} (SYNTHETIC)", groups=None, per_fy_sheets=True,
        project_label="name", status_style="mixed",
        date_styles={"dmy": 0.70, "month_year": 0.25, "iso": 0.05}, entry_styles={"dmy": 0.95, "iso": 0.05},
        cost_always_text=False, totals_label="Total", kinyarwanda=False,
        columns=[("date", "Activity Date"), ("planned_date", "Planned Date"), ("entry_date", "Reported On"),
                 ("project", "Project"), ("activity", "Activity Name"), ("district", "District"),
                 ("sector", "Sector"), ("village", "Village"), ("status", "Status"),
                 ("total", "Total participants"), ("female", "Of which women"), ("male", "Of which men"),
                 ("cost", "Cost RWF"), ("officer", "FO")]),
    # Held-out set only: a fifth officer whose column names the pipeline has never seen.
    "E": dict(
        title=None, groups=None, per_fy_sheets=False, project_label="name", status_style="title",
        date_styles={"d_mon_yy": 0.8, "dmy": 0.2}, entry_styles={"d_mon_yy": 0.8, "dmy": 0.2},
        cost_always_text=False, totals_label="Sum", kinyarwanda=False,
        columns=[("date", "Date of event"), ("planned_date", "Scheduled for"), ("entry_date", "Captured on"),
                 ("project", "Programme/Project"), ("activity", "Activity type"),
                 ("district", "Location (District)"), ("sector", "Location (Sector)"), ("cell", "Location (Cell)"),
                 ("village", "Village name"), ("status", "Progress"), ("female", "No. of women"),
                 ("male", "No. of men"), ("total", "Attendance"), ("cost", "Amount spent"), ("officer", "Staff")]),
    "PAPER": dict(
        title="Digitised paper attendance sheets (typed from paper forms) - SYNTHETIC", groups=None,
        per_fy_sheets=False, project_label="name", status_style="mixed",
        date_styles={"dmy": 0.40, "iso": 0.15, "month_year": 0.20, "serial": 0.10, "datetime": 0.15},
        entry_styles={"dmy": 0.5, "datetime": 0.5},
        cost_always_text=False, totals_label=None, kinyarwanda=False,
        columns=[("row_no", "No."), ("date", "Date"), ("activity", "Activity"), ("project", "Project"),
                 ("district", "District"), ("sector", "Sector"), ("village", "Village"),
                 ("status", "Status"), ("female", "Women"), ("male", "Men"), ("total", "Total"),
                 ("cost", "Cost"), ("officer", "Field officer"), ("entry_date", "Digitised on")]),
}

REGISTER_LAYOUTS = {
    "A": dict(title=None, name_style="surname_first", sex_style={"F": 1.0}, uses_dob=False,
              phone_style="plain", id_style="text", disability_style=("Yes", "No"),
              date_styles={"datetime": 0.85, "dmy": 0.15}, project_label="name", allow_code_1=False,
              columns=[("reg_no", "Reg No"), ("name", "Full name"), ("sex", "Sex"), ("age", "Age"),
                       ("phone", "Phone"), ("national_id", "National ID"), ("household_size", "Household size"),
                       ("disability", "Disability"), ("district", "District"), ("sector", "Sector"),
                       ("cell", "Cell"), ("village", "Village"), ("project", "Project"),
                       ("reg_date", "Registration date")]),
    "B": dict(title="Beneficiary register - {officer} (SYNTHETIC)", name_style="split",
              sex_style={"Female": 0.9, "lower": 0.1}, uses_dob=True, phone_style="intl", id_style="text",
              disability_style=("Y", "N"), date_styles={"iso": 0.9, "dmy": 0.1}, project_label="code",
              allow_code_1=False,
              columns=[("reg_no", "Reg. No."), ("surname", "Surname"), ("first_name", "First name"),
                       ("sex", "Gender"), ("dob", "Date of birth"), ("phone", "Tel"),
                       ("national_id", "ID number"), ("household_size", "HH size"), ("disability", "PWD"),
                       ("province", "Province"), ("district", "District"), ("sector", "Sector"),
                       ("cell", "Cell"), ("village", "Village"), ("project", "Project code"),
                       ("reg_date", "Date registered")]),
    "C": dict(title=None, name_style="upper_surname", sex_style={"Gore": 1.0}, uses_dob=False, kinyarwanda=True,
              phone_style="spaced", id_style="text", disability_style=("Yego", "Oya"),
              date_styles={"serial": 0.9, "datetime": 0.1}, project_label="name", allow_code_1=False,
              columns=[("reg_no", "No"), ("name", "Amazina"), ("sex", "Igitsina"), ("age", "Imyaka"),
                       ("phone", "Telefone"), ("national_id", "Indangamuntu"),
                       ("household_size", "Abagize urugo"), ("disability", "Ubumuga"),
                       ("district", "Akarere"), ("sector", "Umurenge"), ("cell", "Akagari"),
                       ("village", "Umudugudu"), ("project", "Umushinga"), ("reg_date", "Itariki")]),
    "D": dict(title=None, name_style="given_first", sex_style={"mixed": 1.0}, uses_dob=False,
              phone_style="number", id_style="number", disability_style=("Yes", "No"),
              date_styles={"dmy": 0.9, "iso": 0.1}, project_label="name", allow_code_1=True,
              columns=[("reg_no", "ID"), ("name", "Name"), ("sex", "Sex"), ("age", "Age"),
                       ("phone", "Phone number"), ("national_id", "National ID"),
                       ("household_size", "Household members"), ("disability", "Disability"),
                       ("district", "District"), ("sector", "Sector"), ("village", "Village"),
                       ("project", "Project"), ("reg_date", "Date")]),
    # Held-out set only: a fifth officer whose column names the pipeline has never seen.
    "E": dict(title=None, name_style="given_first", sex_style={"F": 1.0}, uses_dob=False, phone_style="intl",
              id_style="text", disability_style=("Yes", "No"), date_styles={"d_mon_yy": 0.8, "dmy": 0.2},
              project_label="name", allow_code_1=False,
              columns=[("reg_no", "Beneficiary code"), ("name", "Names of beneficiary"), ("sex", "M/F"),
                       ("age", "Age (years)"), ("phone", "Mobile"), ("national_id", "NID number"),
                       ("household_size", "Family size"), ("disability", "Living with disability"),
                       ("district", "Location (District)"), ("sector", "Location (Sector)"),
                       ("cell", "Location (Cell)"), ("village", "Village name"), ("project", "Programme/Project"),
                       ("reg_date", "Date enrolled")],
              service_columns=[("reg_no", "Beneficiary code"), ("activity", "Activity type"),
                               ("date", "Date of event"), ("service_type", "Support given")]),
    "PAPER": dict(title="Digitised paper registration forms - SYNTHETIC", name_style="paper",
                  sex_style={"mixed": 1.0}, uses_dob=False, phone_style="plain", id_style="text",
                  disability_style=("Y", "N"),
                  date_styles={"dmy": 0.5, "month_year": 0.1, "iso": 0.2, "datetime": 0.2},
                  project_label="name", allow_code_1=True,
                  columns=[("reg_no", "Form No."), ("name", "Name"), ("sex", "Sex"), ("age", "Age"),
                           ("phone", "Phone"), ("national_id", "ID"), ("household_size", "HH"),
                           ("disability", "Disability"), ("village", "Village"), ("sector", "Sector"),
                           ("district", "District"), ("project", "Project"), ("reg_date", "Date")]),
}


class MessyWriter:
    """Turns the true World into messy Excel files and an injection log."""

    def __init__(self, world: World, raw_dir: Path):
        self.w, self.raw_dir = world, raw_dir
        self.hm = world.cfg.get("holdout_mess") or {}   # extra, unseen mess: held-out set only
        self.rand = world.rand
        self.rates = world.cfg["injection_rates"]
        self.log = InjectionLog()
        self.loc = {l["location_id"]: l for l in world.locations}
        self.all_districts_cf = {d.casefold() for d in district_to_province()}
        self.provinces = list(load_admin()["provinces"])
        self.activity_positions: dict[str, tuple] = {}
        self.registration_positions: dict[str, tuple] = {}
        self.service_positions: dict[str, tuple] = {}
        self.indicator_positions: dict[str, tuple] = {}
        self.opportunity_positions: dict[str, tuple] = {}
        self.donor_positions: dict[str, tuple] = {}
        self.missing_reports: list[dict] = []
        self.books: dict[str, Workbook] = {}

    def book(self, file_name: str) -> Workbook:
        """Return the workbook for a file, creating it the first time it is needed."""
        if file_name not in self.books:
            wb = Workbook()
            wb.remove(wb.active)
            self.books[file_name] = wb
        return self.books[file_name]

    # ---------------------------------------------------------------- shared messing helpers
    def date_cell(self, cells: Cells, fld: str, d: dt.date | None, styles: dict, fields: list):
        if fld not in fields or d is None:
            cells.values[fld] = None
            return
        style = self.rand.pick_style(styles)
        if self.hm and self.rand.chance(self.hm["new_date_format_rate"]):
            style = self.rand.pick(["d_mon_yy", "d_month_yyyy", "dotted"])
        cells.values[fld] = render_date(d, style)
        if style in DATE_ISSUE:
            etype, expected = DATE_ISSUE[style]
            cells.note(fld, etype, d.isoformat(), cells.values[fld], expected)

    def new_district_variant(self, district: str) -> tuple[str, str, str]:
        """Held-out set only: district spellings the cleaner has never seen."""
        options = ["r_l_swap", "two_slips", "abbreviation", "dist_prefix", "k_g_swap"]
        if district in OLD_DISTRICT_NAMES:
            options.append("old_name")
        kind = self.rand.pick(options)
        if kind == "old_name":
            return OLD_DISTRICT_NAMES[district], "district_old_name", "fix_or_flag"
        if kind == "abbreviation":
            return f"{district[:4]}.", "district_abbreviation", "fix_or_flag"
        if kind == "dist_prefix":
            return f"Dist. {district}", "district_variant", "fix"
        for _ in range(20):
            if kind == "r_l_swap" and ("r" in district.lower() or "l" in district.lower()):
                i = next(i for i, ch in enumerate(district.lower()) if ch in "rl")
                ch = district[i]
                swap = {"r": "l", "l": "r", "R": "L", "L": "R"}[ch]
                value = district[:i] + swap + district[i + 1:]
            elif kind == "k_g_swap" and any(c in district.lower() for c in "kg"):
                i = next(i for i, ch in enumerate(district.lower()) if ch in "kg")
                swap = {"k": "g", "g": "k", "K": "G", "G": "K"}[district[i]]
                value = district[:i] + swap + district[i + 1:]
            else:
                value = letter_slip(self.rand, letter_slip(self.rand, district))
            if value.casefold() not in self.all_districts_cf:
                return value, "district_typo", "fix_or_flag"
            kind = "two_slips"
        return district.upper(), "district_variant", "fix"

    def district_cell(self, cells: Cells, district: str, kinyarwanda: bool, fields: list):
        if self.hm and self.rand.chance(self.hm["district_new_variant_rate"]):
            value, etype, expected = self.new_district_variant(district)
            cells.note("district", etype, district, value, expected, detail="spelling not seen in the main set")
            roll = 2.0          # skip the usual variants below
        else:
            roll = self.rand.uniform(0, 1)
        typo, variant = self.rates["district_typo"], self.rates["district_variant"]
        if roll > 1.0:
            pass
        elif roll < typo:
            value = district
            while value.casefold() in self.all_districts_cf:
                value = letter_slip(self.rand, district)
            cells.note("district", "district_typo", district, value, "fix_or_flag")
        elif roll < typo + variant:
            forms = [district.upper(), district.lower(), f"{district} District", f" {district} "]
            if kinyarwanda:
                forms.append(f"Akarere ka {district}")
            value = self.rand.pick(forms)
            cells.note("district", "district_variant", district, value, "fix")
        elif roll <= 1.0:
            value = district
        cells.values["district"] = value
        if "province" in fields:
            province = self.w.d2p[district]
            if self.rand.chance(self.rates["district_wrong_province"]):
                wrong = self.rand.pick([p for p in self.provinces if p != province])
                cells.values["province"] = f"{wrong} Province" if wrong != "Kigali City" else wrong
                cells.note("province", "district_wrong_province", province, cells.values["province"], "flag",
                           detail=f"{district} is in {province}")
            else:
                cells.values["province"] = f"{province} Province" if province != "Kigali City" else province
                if province != "Kigali City":
                    cells.note("province", "province_variant", province, cells.values["province"], "fix")

    def number_as_text(self, cells: Cells, fld: str, kind: str, force: bool = False):
        value = cells.values.get(fld)
        if value is None or not isinstance(value, (int, float)):
            return
        if not force and not self.rand.chance(self.rates["number_as_text"]):
            return
        v = int(value)
        if kind == "money":
            text = self.rand.pick([f"{v:,}", f"{v:,} RWF", f"RWF {v:,}", f"{v:,}".replace(",", " ") + " Frw"])
        elif kind == "count":
            text = self.rand.pick([str(v), f"{v} ", f"{v} people"])
        elif kind == "hectares":
            text = self.rand.pick([f"{value} ha", str(value), f"{value} Ha"])
        elif kind == "percent":
            text = self.rand.pick([f"{value}%", f"{value} %"])
        else:
            text = f"{v:,}"
        cells.values[fld] = text
        cells.note(fld, "number_as_text", value, text, "fix")

    def project_label(self, project_id: str, style: str) -> str:
        return project_id if style == "code" else self.w.projects[project_id]["name"]

    def status_text(self, status: str, style: str) -> str:
        if style == "title":
            return status.title()
        if style == "lower":
            return status
        if style == "kinyarwanda":
            return {"completed": "Done", "delayed": "Late"}.get(status, status.title())
        options = {"completed": ["Completed", "Completed", "Complete", "DONE"],
                   "delayed": ["Delayed", "Late", "delayed"],
                   "cancelled": ["Cancelled", "Canceled", "CANCELLED"],
                   "planned": ["Planned", "Not yet"]}[status]
        return self.rand.pick(options)

    def write_status(self, cells: Cells, status: str, style: str, fields: list):
        if status == "delayed" and self.rand.chance(self.rates["status_says_completed_but_late"]):
            text = self.status_text("completed", style)
            cells.values["status"] = text
            if "planned_date" in fields:
                cells.note("status", "status_completed_but_late", status, text, "flag",
                           detail="actual date is well after the planned date")
            else:
                cells.note("status", "status_completed_but_late", status, text, "undetectable",
                           detail="this file has no planned date, so lateness cannot be seen")
            return
        text = self.status_text(status, style)
        cells.values["status"] = text
        if text != status.title() and text.casefold() != status:
            cells.note("status", "status_variant", status, text, "fix")

    # ---------------------------------------------------------------- activity files
    def activity_cells(self, rec: dict, layout: dict, fields: list, is_paper: bool) -> Cells:
        loc = self.loc[rec["location_id"]]
        act = self.w.activities[rec["activity_id"]]
        cells = Cells(values=dict(
            project=self.project_label(rec["project_id"], layout["project_label"]),
            activity=act["activity_name"], sector=loc["sector"], cell=loc["cell"], village=loc["village"],
            female=rec["participants_female"], male=rec["participants_male"],
            total=rec["participants_total"], cost=rec["cost_rwf"], officer=rec["field_officer"]))
        # Dates
        true_date = rec["date"]
        if true_date and self.rand.chance(self.rates["date_out_of_window"]):
            typo = year_typo(true_date)
            self.date_cell(cells, "date", typo, layout["date_styles"], fields)
            cells.issues = [i for i in cells.issues if i["field"] != "date"]
            cells.note("date", "date_out_of_window", true_date.isoformat(), cells.values["date"], "flag",
                       detail=f"year typed as {typo.year}")
        else:
            self.date_cell(cells, "date", true_date, layout["date_styles"], fields)
        self.date_cell(cells, "planned_date", rec["planned_date"], layout["date_styles"], fields)
        self.date_cell(cells, "entry_date", rec["entry_date"], layout["entry_styles"], fields)
        if rec["is_late_entry"] and "entry_date" in fields:
            cells.note("entry_date", "late_entry", rec["entry_lag_days"], cells.values["entry_date"], "flag",
                       detail=f"entered {rec['entry_lag_days']} days after the activity")
        # Location and status
        self.district_cell(cells, rec["district"], layout["kinyarwanda"], fields)
        self.write_status(cells, rec["status"], layout["status_style"], fields)
        # Participant numbers: at most one numeric problem per row
        if rec["status"] in ("completed", "delayed") and rec["participants_total"] is not None:
            roll = self.rand.uniform(0, 1)
            neg = self.rates["negative_value"]
            over = neg + self.rates["female_exceeds_total"]
            mismatch = over + self.rates["sum_mismatch"]
            if roll < neg:
                fld = self.rand.pick(["male", "cost"])
                cells.values[fld] = -abs(cells.values[fld]) if cells.values[fld] else -self.rand.int(1, 9)
                cells.note(fld, "negative_value", rec["cost_rwf"] if fld == "cost" else rec["participants_male"],
                           cells.values[fld], "flag")
            elif roll < over:
                total = rec["participants_total"]
                cells.values["female"] = math.ceil(total * self.rand.uniform(1.1, 1.4))
                cells.note("female", "female_exceeds_total", rec["participants_female"], cells.values["female"],
                           "flag", detail=f"{cells.values['female'] / total:.0%} female participation")
            elif roll < mismatch:
                shift = self.rand.pick([-8, -5, -3, -2, -1, 1, 2, 3, 5, 10])
                cells.values["total"] = max(0, rec["participants_total"] + shift)
                if cells.values["total"] != rec["participants_total"]:
                    cells.note("total", "sum_mismatch", rec["participants_total"], cells.values["total"], "flag",
                               detail="female + male does not equal total")
        for fld in ("female", "male", "total"):
            self.number_as_text(cells, fld, "count")
        self.number_as_text(cells, "cost", "money", force=layout["cost_always_text"] and rec["cost_rwf"] is not None)
        # Missing required values
        rate = self.rates["missing_required_paper" if is_paper else "missing_required"]
        if rec["status"] != "planned" and self.rand.chance(rate):
            # Do not blank a value another injected problem depends on (that would hide it),
            # and a cancelled activity legitimately has no date or participants.
            protect = set()
            if rec["status"] == "cancelled":
                protect |= {"date", "total"}
            if rec["is_late_entry"] or any(i["error_type"] == "status_completed_but_late" for i in cells.issues):
                protect.add("date")
            options = [f for f in ("date", "district", "activity", "total", "village")
                       if f in fields and f not in protect and cells.values.get(f) not in (None, "")]
            if options:
                fld = self.rand.pick(options)
                true = {"date": rec["date"].isoformat() if rec["date"] else None, "district": rec["district"],
                        "activity": act["activity_name"], "total": rec["participants_total"],
                        "village": loc["village"]}[fld]
                cells.blank(fld, true)
        # Typing errors when paper forms were digitised
        if is_paper and self.rand.chance(self.rates["paper_transcription_error"]):
            self.paper_slip(cells, loc["village"])
        return cells

    def paper_slip(self, cells: Cells, true_village: str):
        kind = self.rand.pick(["O_for_0", "l_for_1", "swapped_digits", "village_spelling"])
        if kind in ("O_for_0", "l_for_1"):
            digit, letter = ("0", "O") if kind == "O_for_0" else ("1", "l")
            for fld in ("cost", "total", "female", "male"):
                value = cells.values.get(fld)
                if value is not None and digit in str(value):
                    new = str(value).replace(digit, letter, 1)
                    cells.issues = [i for i in cells.issues if i["field"] != fld]
                    cells.values[fld] = new
                    cells.note(fld, "transcription_error", value, new, "fix_or_flag", detail=kind)
                    return
        if kind == "swapped_digits":
            value = cells.values.get("total")
            if isinstance(value, int) and value >= 10 and str(value)[-1] != str(value)[-2]:
                s = str(value)
                new = int(s[:-2] + s[-1] + s[-2])
                cells.issues = [i for i in cells.issues if i["field"] != "total"]
                cells.values["total"] = new
                cells.note("total", "transcription_error", value, new, "flag",
                           detail="swapped digits (shows up as female + male not equal to total)")
                return
        if cells.values.get("village"):
            new = letter_slip(self.rand, true_village)
            cells.values["village"] = new
            cells.note("village", "transcription_error", true_village, new, "fix_or_flag",
                       detail="village spelling")

    def write_activity_file(self, key: str, records: list[dict]):
        is_paper = key == "PAPER"
        spec = self.w.cfg["field_officers"].get(key, {})
        layout = ACTIVITY_LAYOUTS["PAPER" if is_paper else spec["layout"]]
        file_name = (f"{PREFIX}digitised_paper_records.xlsx" if is_paper
                     else f"{PREFIX}activity_monitoring_{key}.xlsx")
        wb = self.book(file_name)
        fields = [f for f, _ in layout["columns"]]
        if layout["per_fy_sheets"]:
            groups = {}
            for r in records:
                groups.setdefault(fiscal_year_label(r["date"] or r["planned_date"]), []).append(r)
            sheets = [(fy.replace("/", "-"), fy, recs) for fy, recs in sorted(groups.items())]
        else:
            sheets = [("Activities", None, records)]
        districts = " & ".join(spec.get("districts", []))
        for sheet_name, fy, recs in sheets:
            ws = wb.create_sheet(sheet_name)
            writer = SheetWriter(ws, file_name, "activity", self.log)
            if layout["title"]:
                writer.title(layout["title"].format(officer=key, districts=districts, fy=fy), len(fields))
                if layout["groups"] is None:
                    writer.blank_row(logged=False)
            if layout["groups"]:
                writer.group_header(layout["groups"])
            writer.header(layout["columns"])
            pending_duplicates: list[tuple[int, Cells, str, int]] = []
            for n, rec in enumerate(recs, 1):
                if self.rand.chance(self.rates["blank_row"]) and not is_paper:
                    writer.blank_row()
                cells = self.activity_cells(rec, layout, fields, is_paper)
                cells.values["row_no"] = n
                row = writer.record(cells, rec["record_id"])
                self.activity_positions[rec["record_id"]] = (file_name, ws.title, row)
                if rec["status"] != "planned" and self.rand.chance(self.rates["duplicate_activity_row"]):
                    pending_duplicates.append((n + self.rand.int(0, 3), cells, rec["record_id"], row))
                for item in [p for p in pending_duplicates if p[0] <= n]:
                    pending_duplicates.remove(item)
                    _, dup_cells, rid, original_row = item
                    writer.record(dup_cells, rid, only_issue=dict(
                        field="date", error_type="duplicate_activity_row", true_value=None,
                        messy_value=None, expected_handling="fix", detail=f"exact copy of row {original_row}"))
            for _, dup_cells, rid, original_row in pending_duplicates:
                writer.record(dup_cells, rid, only_issue=dict(
                    field="date", error_type="duplicate_activity_row", true_value=None, messy_value=None,
                    expected_handling="fix", detail=f"exact copy of row {original_row}"))
            if layout["totals_label"]:
                writer.totals(layout["totals_label"], ["female", "male", "total", "cost"])
            writer.finish()

    # ---------------------------------------------------------------- beneficiary registers
    def name_values(self, reg: dict, style: str) -> dict:
        surname, given = reg["surname"], reg["given_name"]
        if style == "paper":
            style = self.rand.pick(["surname_first", "given_first", "upper_all"])
        swapped = reg["dup_change"] == "name order swapped"
        if style == "split":
            if swapped:
                surname, given = given, surname
            return dict(surname=surname, first_name=given)
        if style == "surname_first":
            text = f"{surname} {given}"
        elif style == "upper_surname":
            text = f"{surname.upper()} {given}"
        elif style == "upper_all":
            text = f"{surname} {given}".upper()
        else:
            text = f"{given} {surname}"
        if swapped:
            parts_first = text.split(" ")[0]
            text = " ".join(text.split(" ")[1:] + [parts_first])
        return dict(name=text)

    def sex_text(self, sex: str, style: dict) -> str:
        choice = self.rand.pick_style(style)
        if choice == "F":
            return "F" if sex == "female" else "M"
        if choice == "Female":
            return sex.title()
        if choice == "lower":
            return sex
        if choice == "Gore":
            return "Gore" if sex == "female" else "Gabo"
        return self.rand.pick(["F", "Female", "f"] if sex == "female" else ["M", "Male", "m"])

    def registration_cells(self, reg: dict, person: dict, layout: dict, fields: list, is_paper: bool) -> Cells:
        loc = self.loc[person["location_id"]]
        cells = Cells(values=dict(sector=loc["sector"], cell=loc["cell"], village=loc["village"],
                                  household_size=person["household_size"],
                                  project=self.project_label(person["project_id"], layout["project_label"])))
        cells.values.update(self.name_values(reg, layout["name_style"]))
        # Sex
        sex = self.sex_text(person["sex"], layout["sex_style"])
        if layout["allow_code_1"] and self.rand.chance(self.rates["sex_ambiguous_code_1"]):
            cells.values["sex"] = "1"
            cells.note("sex", "sex_ambiguous_code_1", person["sex"], "1", "flag")
        else:
            cells.values["sex"] = sex
            if sex not in ("Female", "Male"):
                cells.note("sex", "sex_variant", person["sex"], sex, "fix")
        # Age or date of birth
        if layout["uses_dob"]:
            dob = person["date_of_birth"]
            if self.rand.chance(self.rates["impossible_age"]):
                bad = dob.replace(year=self.rand.pick([1890, 1899, 2031]))
                self.date_cell(cells, "dob", bad, layout["date_styles"], fields)
                cells.issues = [i for i in cells.issues if i["field"] != "dob"]
                cells.note("dob", "impossible_age", dob.isoformat(), cells.values["dob"], "flag")
            else:
                self.date_cell(cells, "dob", dob, layout["date_styles"], fields)
        else:
            cells.values["age"] = reg["age"]
            if self.rand.chance(self.rates["impossible_age"]):
                cells.values["age"] = self.rand.pick([131, 150, 230, -4])
                cells.note("age", "impossible_age", reg["age"], cells.values["age"], "flag")
        if self.rand.chance(self.rates["impossible_household_size"]):
            cells.values["household_size"] = self.rand.pick([0, 40, 63, -2])
            cells.note("household_size", "impossible_household_size", person["household_size"],
                       cells.values["household_size"], "flag")
        # Phone and national ID
        phone = reg["phone"]
        if phone:
            if layout["phone_style"] == "intl":
                cells.values["phone"] = f"+250 {phone[1:4]} {phone[4:7]} {phone[7:]}"
            elif layout["phone_style"] == "spaced":
                cells.values["phone"] = f"{phone[:4]} {phone[4:7]} {phone[7:]}"
            elif layout["phone_style"] == "number":
                cells.values["phone"] = int(phone)
                cells.note("phone", "phone_leading_zero_lost", phone, int(phone), "fix")
            else:
                cells.values["phone"] = phone
        national_id = reg["national_id"]
        if national_id:
            if layout["id_style"] == "number":
                cells.values["national_id"] = int(national_id[:15] + "0")
                cells.note("national_id", "id_stored_as_number", national_id, cells.values["national_id"], "flag",
                           detail="Excel keeps only 15 digits, so the last digit became 0")
            else:
                cells.values["national_id"] = national_id
        # Disability
        yes, no = layout["disability_style"]
        cells.values["disability"] = yes if person["disability"] else no
        if cells.values["disability"] not in ("Yes", "No"):
            cells.note("disability", "disability_variant", "yes" if person["disability"] else "no",
                       cells.values["disability"], "fix")
        self.district_cell(cells, person["district"], layout.get("kinyarwanda", False), fields)
        self.date_cell(cells, "reg_date", reg["registration_date"], layout["date_styles"], fields)
        # Missing values
        rate = self.rates["missing_required_paper" if is_paper else "missing_required"]
        if self.rand.chance(rate):
            options = [f for f in ("sex", "district", "village", "age") if f in fields and cells.values.get(f) is not None]
            if options:
                fld = self.rand.pick(options)
                true = {"sex": person["sex"], "district": person["district"], "village": loc["village"],
                        "age": reg["age"]}[fld]
                cells.blank(fld, true)
        if is_paper and self.rand.chance(self.rates["paper_transcription_error"]):
            kind = self.rand.pick(["name_letter", "phone_digit_dropped", "village_spelling"])
            if kind == "phone_digit_dropped" and isinstance(cells.values.get("phone"), str):
                old = cells.values["phone"]
                i = self.rand.int(3, len(old) - 1)
                cells.values["phone"] = old[:i] + old[i + 1:]
                cells.note("phone", "transcription_error", old, cells.values["phone"], "flag",
                           detail="digit dropped (phone too short)")
            elif kind == "village_spelling" and cells.values.get("village"):
                cells.values["village"] = letter_slip(self.rand, loc["village"])
                cells.note("village", "transcription_error", loc["village"], cells.values["village"],
                           "fix_or_flag", detail="village spelling")
            elif cells.values.get("name"):
                old = cells.values["name"]
                parts = old.split(" ")
                parts[0] = letter_slip(self.rand, parts[0])
                cells.values["name"] = " ".join(parts)
                cells.note("name", "transcription_error", old, cells.values["name"], "undetectable",
                           detail="name spelling (cannot be detected without the paper form)")
        return cells

    def write_register_file(self, key: str, regs: list[dict], services: list[dict]):
        is_paper = key == "PAPER"
        spec = self.w.cfg["field_officers"].get(key, {})
        layout = REGISTER_LAYOUTS["PAPER" if is_paper else spec["layout"]]
        file_name = (f"{PREFIX}digitised_paper_records.xlsx" if is_paper
                     else f"{PREFIX}beneficiary_register_{key}.xlsx")
        people = {p["person_id"]: p for p in self.w.people}
        fields = [f for f, _ in layout["columns"]]
        wb = self.book(file_name)
        ws = wb.create_sheet("Registrations" if is_paper else "Register")
        writer = SheetWriter(ws, file_name, "registration", self.log)
        if layout["title"]:
            writer.title(layout["title"].format(officer=key), len(fields))
            writer.blank_row(logged=False)
        writer.header(layout["columns"])
        prefix = "P" if is_paper else key
        reg_no = {}
        for n, reg in enumerate(sorted(regs, key=lambda r: (r["registration_date"], r["registration_id"])), 1):
            if not is_paper and self.rand.chance(self.rates["blank_row"]):
                writer.blank_row()
            person = people[reg["person_id"]]
            cells = self.registration_cells(reg, person, layout, fields, is_paper)
            reg_no[reg["registration_id"]] = f"{prefix}-{n:04d}"
            cells.values["reg_no"] = reg_no[reg["registration_id"]]
            if person.get("namesake_of") and not reg["is_duplicate"]:
                cells.note("name" if "name" in fields else "surname", "namesake_different_person",
                           person["namesake_of"], cells.values.get("name") or cells.values.get("surname"),
                           "keep_apart", detail="a different person with the same name, district and sex")
            if reg["is_duplicate"]:
                cells.note("name" if "name" in fields else "surname", "duplicate_beneficiary",
                           reg["person_id"], cells.values.get("name") or cells.values.get("surname"),
                           "flag", detail=reg["dup_change"])
            row = writer.record(cells, reg["registration_id"])
            self.registration_positions[reg["registration_id"]] = (file_name, ws.title, row, reg_no[reg["registration_id"]])
        writer.finish()
        # Services sheet: which registered person received which service, and when.
        ws2 = wb.create_sheet("Services")
        writer2 = SheetWriter(ws2, file_name, "service", self.log)
        columns = layout.get("service_columns") or [
            ("reg_no", "Reg No" if not is_paper else "Form No."), ("activity", "Activity"),
            ("date", "Date"), ("service_type", "Service received")]
        writer2.header(columns)
        for s in sorted(services, key=lambda s: (s["date"], s["service_id"])):
            cells = Cells(values=dict(reg_no=reg_no[s["registration_id"]],
                                      activity=self.w.activities[s["activity_id"]]["activity_name"],
                                      service_type=s["service_type"]))
            self.date_cell(cells, "date", s["date"], layout["date_styles"], [c for c, _ in columns])
            row = writer2.record(cells, s["service_id"])
            self.service_positions[s["service_id"]] = (file_name, ws2.title, row)
        writer2.finish()

    # ---------------------------------------------------------------- indicator reports
    def write_indicator_file(self, programme: dict):
        file_name = f"{PREFIX}indicator_reports_{programme['thematic_area'].upper()[:4]}.xlsx"
        rows = [r for r in self.w.indicator_rows if r["programme_id"] == programme["programme_id"]]
        # Decide which expected reports were never submitted.
        expected = [e for e in self.w.expected_reports
                    if self.w.projects[e["project_id"]]["programme_id"] == programme["programme_id"]]
        missing = set()
        for e in expected:
            if self.rand.chance(self.rates["missing_monthly_report"]):
                missing.add((e["project_id"], e["district"], e["period"]))
                self.missing_reports.append(e)
                self.log.add(source_file=file_name, entity="indicator_report", error_type="missing_monthly_report",
                             record_id=f"{e['project_id']}|{e['district']}|{e['period'].isoformat()}",
                             expected_handling="flag",
                             detail=f"no report for {e['project_id']} {e['district']} {e['period']:%b %Y}")
        # Percentage indicators have few rows, so their problems are injected at exact counts.
        pct_ids = [r["result_id"] for r in rows if self.w.indicators[r["indicator_code"]]["unit"] == "percent"]
        self.over_100 = self.rand.sample(pct_ids, max(2, round(self.rates["percent_over_100"] * len(pct_ids))))
        self.as_fraction = self.rand.sample([i for i in pct_ids if i not in self.over_100],
                                            max(2, round(self.rates["percent_as_fraction"] * len(pct_ids))))
        wb = self.book(file_name)
        for fy in sorted({r["fiscal_year"] for r in rows}):
            ws = wb.create_sheet(fy.replace("/", "-"))
            writer = SheetWriter(ws, file_name, "indicator_result", self.log)
            if fy == "FY2023/24":
                columns = [("project", "Project"), ("district", "District"), ("period", "Month"),
                           ("indicator", "Indicator"), ("sex", "Sex"), ("target", "Target"),
                           ("actual", "Actual"), ("entered_by", "Reported by")]
            else:
                columns = [("project", "Project code"), ("district", "District"), ("period", "Reporting period"),
                           ("indicator", "Indicator code"), ("sex", "Sex"), ("target", "Monthly target"),
                           ("actual", "Achieved"), ("entered_by", "Entered by")]
                if fy >= "FY2025/26":
                    columns.append(("comment", "Comments"))
            fields = [c for c, _ in columns]
            writer.title(f"Monthly indicator report - {programme['name']} - {fy} (SYNTHETIC)", len(fields))
            writer.header(columns, formats={"period": "mmm-yy"})
            for r in [r for r in rows if r["fiscal_year"] == fy]:
                if (r["project_id"], r["district"], r["period"]) in missing:
                    continue
                cells = self.indicator_cells(r, fy, fields)
                row = writer.record(cells, r["result_id"])
                self.indicator_positions[r["result_id"]] = (file_name, ws.title, row)
            writer.finish()

    def indicator_cells(self, r: dict, fy: str, fields: list) -> Cells:
        ind = self.w.indicators[r["indicator_code"]]
        unit = ind["unit"]
        if fy == "FY2023/24":
            period = r["period"].strftime("%b-%y")         # text, e.g. Jul-23
            indicator = ind["name"]
            project = self.w.projects[r["project_id"]]["name"]
        elif fy == "FY2024/25":
            period = dt.datetime(r["period"].year, r["period"].month, 1)
            indicator, project = r["indicator_code"], r["project_id"]
        else:
            period = r["period"].strftime("%Y-%m")         # text, e.g. 2025-07
            indicator, project = r["indicator_code"], r["project_id"]
        cells = Cells(values=dict(project=project, period=period, indicator=indicator, sex=r["sex"],
                                  target=r["target_value"], actual=r["actual_value"], entered_by=r["entered_by"],
                                  comment=self.rand.pick(REPORT_COMMENTS) if self.rand.chance(0.1) else None))
        self.district_cell(cells, r["district"], False, fields)
        roll = self.rand.uniform(0, 1)
        thresholds = np.cumsum([self.rates["missing_target"], self.rates["accuracy_outlier"],
                                self.rates["negative_value"]])
        if r["result_id"] in self.over_100:
            cells.values["actual"] = round(self.rand.uniform(104, 160), 1)
            cells.note("actual", "percent_over_100", r["actual_value"], cells.values["actual"], "flag")
        elif r["result_id"] in self.as_fraction:
            cells.values["actual"] = round(r["actual_value"] / 100, 3)
            cells.note("actual", "percent_as_fraction", r["actual_value"], cells.values["actual"], "fix",
                       detail="typed as a fraction instead of a percentage")
        elif roll < thresholds[0]:
            cells.values["target"] = None
            cells.note("target", "missing_target", r["target_value"], None, "flag")
        elif roll < thresholds[1] and r["actual_value"] and unit != "percent":
            cells.values["actual"] = r["actual_value"] * 10
            cells.note("actual", "accuracy_outlier", r["actual_value"], cells.values["actual"], "flag",
                       detail="extra zero typed: more than 3 x target")
        elif roll < thresholds[2] and r["actual_value"]:
            cells.values["actual"] = -r["actual_value"]
            cells.note("actual", "negative_value", r["actual_value"], cells.values["actual"], "flag")
        kind = {"percent": "percent", "hectares": "hectares"}.get(unit, "count")
        if isinstance(cells.values["actual"], (int, float)) and cells.values["actual"] >= 0:
            self.number_as_text(cells, "actual", kind)
        if isinstance(cells.values["target"], (int, float)):
            self.number_as_text(cells, "target", kind)
        return cells

    # ---------------------------------------------------------------- donor tracker
    def write_donor_file(self):
        file_name = f"{PREFIX}donor_opportunity_tracker.xlsx"
        wb = self.book(file_name)
        ws = wb.create_sheet("Donors")
        writer = SheetWriter(ws, file_name, "donor", self.log)
        columns = [("donor_id", "Donor ID"), ("donor_name", "Donor name"), ("donor_type", "Type"),
                   ("country", "Country"), ("thematic_interests", "Thematic interests"),
                   ("contact_person", "Contact person"), ("email", "Email"), ("phone", "Phone"), ("notes", "Notes")]
        writer.header(columns)
        for d in self.w.donors:
            individual = d["donor_type"] == "individual"
            contact = d["donor_name"] if individual else self.rand.pick(CONTACT_NAMES)
            handle = contact.split(" ")[-1].lower().replace(".", "")
            dtype = d["donor_type"].title()
            cells = Cells(values=dict(donor_id=d["donor_id"], donor_name=d["donor_name"], country=d["country"],
                                      thematic_interests=d["thematic_interests"], contact_person=contact,
                                      email=f"{handle}@example.org",
                                      phone=f"+{self.rand.pick(DIAL_CODES)} {self.rand.int(100, 999)} {self.rand.int(100000, 999999)}",
                                      notes="Met at partner forum" if self.rand.chance(0.2) else None))
            if self.rand.chance(0.25):
                dtype = self.rand.pick(DONOR_TYPE_VARIANTS[d["donor_type"]])
                cells.note("donor_type", "donor_type_variant", d["donor_type"], dtype, "fix")
            cells.values["donor_type"] = dtype
            if individual:
                cells.note("donor_name", "personal_data_individual_donor", None, d["donor_name"], "fix",
                           detail="individual's name must be pseudonymised")
            cells.note("contact_person", "personal_data_contact", None, contact, "fix",
                       detail="contact name, email and phone must not reach the dashboard")
            row = writer.record(cells, d["donor_id"])
            self.donor_positions[d["donor_id"]] = (file_name, ws.title, row)
        writer.finish()

        ws2 = wb.create_sheet("Opportunities")
        writer2 = SheetWriter(ws2, file_name, "opportunity", self.log)
        columns = [("opportunity_id", "Opp ID"), ("title", "Opportunity title"), ("donor", "Donor"),
                   ("thematic_fit", "Thematic fit"), ("stage", "Stage"), ("deadline", "Deadline"),
                   ("requested", "Requested amount"), ("awarded", "Awarded amount"),
                   ("probability", "Probability"), ("owner", "Owner"), ("next_action", "Next action"),
                   ("next_action_date", "Next action date")]
        fields = [c for c, _ in columns]
        writer2.header(columns, formats={"probability": "0%", "requested": "#,##0", "awarded": "#,##0"})
        rate = self.w.settings["currency"]["usd_to_rwf"]
        in_usd = self.rand.sample([o["opportunity_id"] for o in self.w.opportunities], 3)
        for o in self.w.opportunities:
            cells = Cells(values=dict(opportunity_id=o["opportunity_id"], title=o["title"],
                                      donor=o["donor_name"], thematic_fit=o["thematic_fit"], owner=o["owner"],
                                      next_action=o["next_action"], requested=o["requested_amount_rwf"],
                                      awarded=o["awarded_amount_rwf"], probability=o["probability"]))
            stage = o["stage"]
            if self.rand.chance(0.35):
                stage = self.rand.pick(STAGE_VARIANTS[o["stage"]])
                cells.note("stage", "stage_variant", o["stage"], stage, "fix")
            cells.values["stage"] = stage
            if self.rand.chance(0.15):
                cells.values["donor"] = self.rand.pick([o["donor_name"].upper(), f" {o['donor_name']}"])
                cells.note("donor", "donor_name_variant", o["donor_name"], cells.values["donor"], "fix")
            styles = {"datetime": 0.5, "dmy": 0.3, "iso": 0.1, "month_year": 0.1}
            self.date_cell(cells, "deadline", o["deadline"], styles, fields)
            self.date_cell(cells, "next_action_date", o["next_action_date"], {"datetime": 0.7, "dmy": 0.3}, fields)
            if o["opportunity_id"] in in_usd:
                usd = int(round(o["requested_amount_rwf"] / rate, -3))
                cells.values["requested"] = f"USD {usd:,}"
                cells.note("requested", "amount_in_usd", o["requested_amount_rwf"], cells.values["requested"],
                           "fix", detail=f"convert at {rate} RWF per USD")
            elif self.rand.chance(0.25):
                self.number_as_text(cells, "requested", "money", force=True)
            if o["awarded_amount_rwf"] and self.rand.chance(0.3):
                self.number_as_text(cells, "awarded", "money", force=True)
            p = o["probability"]
            p_roll = self.rand.uniform(0, 1)
            if p_roll < 0.2:
                cells.values["probability"] = f"{p:.0%}"
                cells.note("probability", "probability_format", p, cells.values["probability"], "fix")
            elif p_roll < 0.35:
                cells.values["probability"] = int(round(p * 100))
                cells.note("probability", "probability_format", p, cells.values["probability"], "fix",
                           detail="typed as 60 instead of 60%")
            elif p_roll < 0.45 and 0 < p < 1:
                cells.values["probability"] = "High" if p >= 0.6 else "Medium" if p >= 0.3 else "Low"
                cells.note("probability", "probability_format", p, cells.values["probability"], "flag",
                           detail="word instead of a number")
            row = writer2.record(cells, o["opportunity_id"])
            self.opportunity_positions[o["opportunity_id"]] = (file_name, ws2.title, row)
        writer2.finish()

    # ---------------------------------------------------------------- reference workbook (clean)
    def write_programme_master(self):
        wb = self.book(f"{PREFIX}programme_master.xlsx")
        tables = {
            "Programmes": pd.DataFrame(self.w.programmes),
            "Projects": pd.DataFrame([{k: (", ".join(v) if isinstance(v, list) else v) for k, v in p.items()}
                                      for p in self.w.projects.values()]),
            "Activities": pd.DataFrame(self.w.cfg["activities"])[
                ["activity_id", "project_id", "activity_name", "activity_type"]],
            "Locations": pd.DataFrame(self.w.locations),
            "Field officers": pd.DataFrame([dict(officer_code=k, districts=", ".join(v["districts"]))
                                            for k, v in self.w.cfg["field_officers"].items()]),
        }
        for name, df in tables.items():
            ws = wb.create_sheet(name)
            ws.append(list(df.columns))
            for cell in ws[1]:
                cell.font, cell.fill = Font(bold=True), SheetWriter.HEADER_FILL
            for row in df.itertuples(index=False):
                ws.append([v.isoformat() if isinstance(v, dt.date) else v for v in row])
            for i, col in enumerate(df.columns, 1):
                ws.column_dimensions[get_column_letter(i)].width = min(40, max(12, len(col) + 4))


# ======================================================================================
# 3. Orchestration and answer key
# ======================================================================================
def load_synthetic_config(name: str = "synthetic.yaml") -> dict:
    """Read a generator config. A config with 'extends: other.yaml' starts from that file and
    replaces only the settings it lists (used by the held-out set)."""
    cfg = load_yaml(name)
    base_name = cfg.pop("extends", None)
    if not base_name:
        return cfg

    def merge(base, extra):
        for k, v in extra.items():
            base[k] = merge(base[k], v) if isinstance(v, dict) and isinstance(base.get(k), dict) else v
        return base

    return merge(load_synthetic_config(base_name), cfg)


def generate(raw_dir: Path | None = None, truth_dir: Path | None = None, quiet: bool = False,
             config_name: str = "synthetic.yaml") -> dict:
    """Create all synthetic files. Returns a small summary dict (used by tests)."""
    cfg, settings = load_synthetic_config(config_name), load_settings()
    raw_dir = Path(raw_dir or resolve_path("raw"))
    truth_dir = Path(truth_dir or resolve_path("ground_truth"))
    raw_dir.mkdir(parents=True, exist_ok=True)
    truth_dir.mkdir(parents=True, exist_ok=True)
    for old in list(raw_dir.glob(f"{PREFIX}*.xlsx")) + list(truth_dir.glob(f"{PREFIX}*")):
        old.unlink()

    rand = Rand(cfg["seed"])
    world = World(cfg, settings, rand)
    mw = MessyWriter(world, raw_dir)
    mw.write_programme_master()
    officer_keys = list(cfg["field_officers"]) + ["PAPER"]
    for key in officer_keys:
        recs = [r for r in world.activity_records
                if (r["on_paper"] if key == "PAPER" else (not r["on_paper"] and r["field_officer"] == key))]
        mw.write_activity_file(key, recs)
    for key in officer_keys:
        regs = [r for r in world.registrations if r["file_key"] == key]
        srvs = [s for s in world.services if s["file_key"] == key]
        mw.write_register_file(key, regs, srvs)
    for programme in world.programmes:
        mw.write_indicator_file(programme)
    mw.write_donor_file()
    books = mw.books
    for name, wb in books.items():
        wb.save(raw_dir / name)

    write_answer_key(world, mw, truth_dir)
    summary = summarise(mw.log, truth_dir)
    fingerprint = content_fingerprint(raw_dir, truth_dir)
    if not quiet:
        print(f"Wrote {len(books)} synthetic Excel files to {raw_dir}")
        print(f"Wrote answer key and injected-error log to {truth_dir}")
        print(summary.to_string(index=False))
        print(f"Content fingerprint: {fingerprint}")
    return dict(files=sorted(books), log_rows=len(mw.log.rows), summary=summary, fingerprint=fingerprint,
                activity_records=len(world.activity_records), people=len(world.people),
                registrations=len(world.registrations), services=len(world.services),
                indicator_rows=len(world.indicator_rows), donors=len(world.donors),
                opportunities=len(world.opportunities))


def _with_position(rows: list[dict], positions: dict, key: str, extra: tuple = ()) -> pd.DataFrame:
    df = pd.DataFrame(rows)
    cols = ["source_file", "source_sheet", "source_row"] + list(extra)
    pos = {k: v for k, v in positions.items()}
    for i, col in enumerate(cols):
        df[col] = df[key].map(lambda k: pos[k][i] if k in pos else None)
    return df


def write_answer_key(world: World, mw: MessyWriter, truth_dir: Path):
    """Save the true values and where each record was written, for scoring later phases."""
    loc = pd.DataFrame(world.locations)
    act = _with_position(world.activity_records, mw.activity_positions, "record_id")
    act = act.merge(loc[["location_id", "province", "sector", "cell", "village"]], on="location_id", how="left")
    act.to_csv(truth_dir / f"{PREFIX}truth_activities.csv", index=False)

    people = pd.DataFrame(world.people)
    people.to_csv(truth_dir / f"{PREFIX}truth_people.csv", index=False)
    regs = _with_position(world.registrations, mw.registration_positions, "registration_id", extra=("reg_no",))
    regs.to_csv(truth_dir / f"{PREFIX}truth_registrations.csv", index=False)
    srv = _with_position(world.services, mw.service_positions, "service_id")
    srv.to_csv(truth_dir / f"{PREFIX}truth_services.csv", index=False)

    ind = _with_position(world.indicator_rows, mw.indicator_positions, "result_id")
    ind["submitted"] = ind["source_row"].notna()
    ind.to_csv(truth_dir / f"{PREFIX}truth_indicator_results.csv", index=False)
    exp = pd.DataFrame(world.expected_reports)
    missing = {(m["project_id"], m["district"], m["period"]) for m in mw.missing_reports}
    exp["submitted"] = [(r.project_id, r.district, r.period) not in missing for r in exp.itertuples()]
    exp.to_csv(truth_dir / f"{PREFIX}truth_expected_reports.csv", index=False)

    _with_position(world.donors, mw.donor_positions, "donor_id").to_csv(
        truth_dir / f"{PREFIX}truth_donors.csv", index=False)
    _with_position(world.opportunities, mw.opportunity_positions, "opportunity_id").to_csv(
        truth_dir / f"{PREFIX}truth_opportunities.csv", index=False)
    loc.to_csv(truth_dir / f"{PREFIX}truth_locations.csv", index=False)

    # Fill in, for each duplicate registration, where the first registration is.
    first = {r["person_id"]: mw.registration_positions[r["registration_id"]]
             for r in world.registrations if not r["is_duplicate"]}
    for row in mw.log.rows:
        if row["error_type"] == "duplicate_beneficiary":
            f, s, r, n = first[row["true_value"]]
            row["detail"] = f"{row['detail']}; first registered as {n} ({f}, {s} row {r})"
    log = pd.DataFrame(mw.log.rows)
    log.insert(0, "error_id", [f"ERR{i:05d}" for i in range(1, len(log) + 1)])
    log.to_csv(truth_dir / f"{PREFIX}injected_errors_log.csv", index=False)


def summarise(log: InjectionLog, truth_dir: Path) -> pd.DataFrame:
    df = pd.DataFrame(log.rows)
    summary = (df.groupby(["entity", "error_type", "expected_handling"], dropna=False)
               .size().reset_index(name="count").sort_values(["entity", "count"], ascending=[True, False]))
    summary.to_csv(truth_dir / f"{PREFIX}injected_errors_summary.csv", index=False)
    by_file = df.groupby(["source_file", "error_type"], dropna=False).size().reset_index(name="count")
    by_file.to_csv(truth_dir / f"{PREFIX}injected_errors_by_file.csv", index=False)
    return summary


def content_fingerprint(raw_dir: Path, truth_dir: Path) -> str:
    """One SHA-256 code summarising every cell of every synthetic file and the answer key.

    Excel files record the time they were saved, so their bytes differ on every run;
    this fingerprint looks only at the contents, which are identical for the same seed.
    """
    from openpyxl import load_workbook

    digest = hashlib.sha256()
    for path in sorted(Path(raw_dir).glob(f"{PREFIX}*.xlsx")):
        digest.update(path.name.encode())
        wb = load_workbook(path, read_only=True)
        for ws in wb.worksheets:
            digest.update(ws.title.encode())
            for row in ws.iter_rows(values_only=True):
                digest.update(repr(row).encode())
        wb.close()
    for path in sorted(Path(truth_dir).glob(f"{PREFIX}*.csv")):
        digest.update(path.name.encode())
        digest.update(path.read_bytes().replace(b"\r\n", b"\n"))
    return digest.hexdigest()


def verify(raw_dir: Path | None = None, truth_dir: Path | None = None, config_name: str = "synthetic.yaml") -> bool:
    """Check that the synthetic files on this computer match the expected fingerprint."""
    expected = load_synthetic_config(config_name).get("expected_fingerprint")
    actual = content_fingerprint(Path(raw_dir or resolve_path("raw")), Path(truth_dir or resolve_path("ground_truth")))
    if actual == expected:
        print(f"OK: your synthetic data matches the expected version ({actual[:16]}...).")
        return True
    print(f"MISMATCH: expected {expected}\n          found    {actual}\n"
          "Rerun 'python src/generate_synthetic.py' and check you installed requirements.txt exactly.")
    return False


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--raw-dir", type=Path, default=None, help="where to write the messy Excel files")
    parser.add_argument("--truth-dir", type=Path, default=None, help="where to write the answer key")
    parser.add_argument("--verify", action="store_true", help="check existing files instead of generating")
    parser.add_argument("--holdout", action="store_true",
                        help="the held-out test set (config/synthetic_holdout.yaml, written to data/holdout/)")
    args = parser.parse_args()
    config_name = "synthetic_holdout.yaml" if args.holdout else "synthetic.yaml"
    raw_dir = args.raw_dir or (resolve_path("holdout_raw") if args.holdout else None)
    truth_dir = args.truth_dir or (resolve_path("holdout_ground_truth") if args.holdout else None)
    if args.verify:
        sys.exit(0 if verify(raw_dir, truth_dir, config_name) else 1)
    generate(raw_dir, truth_dir, config_name=config_name)


def generate_holdout() -> dict:
    """Pipeline step: create the held-out test set in data/holdout/."""
    return generate(resolve_path("holdout_raw"), resolve_path("holdout_ground_truth"),
                    config_name="synthetic_holdout.yaml")


if __name__ == "__main__":
    main()
