"""Make the Excel file used in the demonstration's upload step (docs/demo_script.md). DEMONSTRATION ONLY.

It fills the blank monthly indicator report template with two rows:
  row 2  a report that is missing from the synthetic data (PRJ01, Nyagatare, February 2026, WASH-01): accepted;
  row 3  a report whose "Achieved" was typed as "about 30" (PRJ01, Kayonza, August 2024): held back, with the reason.
and saves it as data/exports/SYNTHETIC_june_report_PRJ01.xlsx (synthetic, like every demonstration file).

HOW TO RUN
    python tools/make_demo_upload.py
"""

from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook  # noqa: E402

from common import load_settings, resolve_path  # noqa: E402

ROWS = [["PRJ01", "Nyagatare", dt.datetime(2026, 2, 1), "WASH-01", None, 120, 104, "FO01", ""],
        ["PRJ01", "Kayonza", dt.datetime(2024, 8, 1), "WASH-01", None, 120, "about 30", "FO01", "estimate"]]


def make(out: Path | None = None) -> Path:
    if not load_settings().get("synthetic_mode"):
        raise SystemExit("STOPPED: synthetic_mode is off. This file is for the demonstration only.")
    wb = load_workbook(ROOT / "templates" / "indicator_report_template.xlsx")
    ws = wb["Report"]
    for i, row in enumerate(ROWS, start=2):
        for j, value in enumerate(row, start=1):
            ws.cell(i, j, value)
    out = Path(out or resolve_path("exports") / "SYNTHETIC_june_report_PRJ01.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    print(f"Written: {make()}")
