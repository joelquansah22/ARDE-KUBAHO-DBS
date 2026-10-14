"""Phase 4: indicator calculations, reconciliation, missing targets, no double counting, traceability."""

import datetime as dt

import duckdb
import pandas as pd
import pytest
from openpyxl import load_workbook

import clean
import indicators
import ingest
import model
import validate
from common import load_settings, load_yaml

S = load_settings()
CATALOGUE = {i["code"]: i for i in load_yaml("indicator_catalogue.yaml")["indicators"]}
SECTION_7_FIELDS = ["code", "name", "definition", "level", "unit", "numerator", "denominator", "calculation_rule",
                    "disaggregation", "target_setting_rule", "frequency", "data_source", "owner",
                    "higher_is_better", "notes"]
D = dt.date


# ---------------------------------------------------------------- small hand-made data
class FakeData:
    """The same shape as indicators.Data, filled by hand so every expected value can be worked out on paper."""

    def __init__(self, activities=(), services=(), results=(), completeness=()):
        lineage = lambda i: dict(source_file="f.xlsx", source_sheet="S", source_row=i + 2)
        self.activities = pd.DataFrame([{**lineage(i), **r} for i, r in enumerate(activities)],
                                       columns=["source_file", "source_sheet", "source_row", "project_id", "district",
                                                "date", "planned_date", "status", "on_time", "cost_rwf",
                                                "activity_id", "participants_total"])
        self.services = pd.DataFrame([{**lineage(i), **r} for i, r in enumerate(services)],
                                     columns=["source_file", "source_sheet", "source_row", "beneficiary_key",
                                              "activity_id", "project_id", "district", "sex", "date"])
        self.results = pd.DataFrame([{**lineage(i), **r} for i, r in enumerate(results)],
                                    columns=["source_file", "source_sheet", "source_row", "indicator_code",
                                             "project_id", "district", "period_date", "target_value", "actual_value",
                                             "disaggregation_key"])
        self.completeness = pd.DataFrame(list(completeness),
                                         columns=["project_id", "district", "period", "indicators_due", "submitted"])
        self.dim_indicator = pd.DataFrame(dict(code=list(CATALOGUE), indicator_id=range(1, len(CATALOGUE) + 1)))


def pick(values, **where):
    v = values
    for k, val in where.items():
        v = v[v[k] == val]
    assert len(v) == 1, (where, len(v))
    return v.iloc[0]


def run_method(fn, data, code, spec=None):
    """Run one calculation exactly as the pipeline does (monthly values, then the catalogue's rollup)."""
    spec = spec or CATALOGUE[code]
    builder = indicators.IndicatorBuilder.__new__(indicators.IndicatorBuilder)
    builder.settings, builder.data = S, data
    values, links = builder.compute(spec, "test", fn)
    return indicators.add_status(values, spec, S), links


def service(key, sex, date, activity="ACT10", project="PRJ04", district="Kirehe"):
    return dict(beneficiary_key=key, activity_id=activity, project_id=project, district=district, sex=sex, date=date)


# ---------------------------------------------------------------- the catalogue
@pytest.mark.parametrize("code", list(CATALOGUE))
def test_every_indicator_has_the_section_7_fields_and_is_a_draft(code):
    ind = CATALOGUE[code]
    for f in SECTION_7_FIELDS:
        assert f in ind and ind[f] not in (None, ""), (code, f)
    assert ind["status"] == "Draft, for co-design with ARDE/KUBAHO staff"
    assert "counting_rule" in ind


def test_survey_outcomes_are_marked_not_available():
    for code in ("WASH-03", "AGR-03"):
        assert CATALOGUE[code]["availability"] == "data source not yet available"
        assert "calculation" not in CATALOGUE[code]


# ---------------------------------------------------------------- no double counting
def test_a_person_with_many_services_counts_once():
    data = FakeData(services=[service("BEN-1", "female", D(2024, 3, 1)), service("BEN-1", "female", D(2024, 3, 9)),
                              service("BEN-1", "female", D(2024, 4, 2)), service("BEN-2", "male", D(2024, 3, 5))])
    values, _ = run_method(indicators.unique_people, data, "AGR-01")
    assert pick(values, period_type="month", period="2024-03", project_id="PRJ04", district="Kirehe",
                disaggregation="total")["actual"] == 2
    # Across a quarter the same woman still counts once (not once per month).
    assert pick(values, period_type="quarter", period="FY2023/24 Q3", project_id="All", district="All",
                disaggregation="total")["actual"] == 2
    assert pick(values, period_type="quarter", period="FY2023/24 Q4", project_id="All", district="All",
                disaggregation="sex=female")["actual"] == 1


def test_all_districts_is_counted_directly_not_added_up():
    # One person served in two districts' projects must count once for "All", not twice.
    data = FakeData(services=[service("BEN-1", "male", D(2024, 3, 1), district="Kirehe"),
                              service("BEN-1", "male", D(2024, 3, 2), district="Bugesera")])
    values, _ = run_method(indicators.unique_people, data, "AGR-01")
    assert pick(values, period_type="month", period="2024-03", project_id="All", district="All",
                disaggregation="total")["actual"] == 1


def test_share_of_women_uses_unique_people_not_attendance():
    data = FakeData(services=[service("BEN-W", "female", D(2024, 3, d)) for d in range(1, 6)]
                    + [service("BEN-M", "male", D(2024, 3, 1)), service("BEN-U", None, D(2024, 3, 1))])
    values, _ = run_method(indicators.share_women, data, "CC-01")
    v = pick(values, period_type="month", period="2024-03", project_id="All", district="All")
    assert v["actual"] == 50.0                     # 1 woman of 2 people with sex recorded (not 5 of 6 sessions)
    assert v["coverage_pct"] == pytest.approx(66.7, abs=0.1)   # 2 of 3 people have their sex recorded


# ---------------------------------------------------------------- the other derived indicators
def test_on_time_share_and_its_coverage():
    acts = [dict(project_id="PRJ01", district="Huye", date=D(2024, 3, 2), planned_date=D(2024, 3, 1),
                 status="completed", on_time=True, cost_rwf=10),
            dict(project_id="PRJ01", district="Huye", date=D(2024, 3, 25), planned_date=D(2024, 3, 2),
                 status="delayed", on_time=False, cost_rwf=10),
            dict(project_id="PRJ01", district="Huye", date=D(2024, 3, 3), planned_date=D(2024, 3, 3),
                 status="completed", on_time=None, cost_rwf=10)]          # timing cannot be checked
    values, _ = run_method(indicators.on_time, FakeData(activities=acts), "CC-02")
    v = pick(values, period_type="month", period="2024-03", project_id="PRJ01", district="Huye")
    assert v["actual"] == 50.0 and v["numerator"] == 1 and v["denominator"] == 2
    assert v["coverage_pct"] == pytest.approx(66.7, abs=0.1)


def test_cost_per_person_and_its_direction():
    acts = [dict(project_id="PRJ04", district="Kirehe", date=D(2024, 3, 2), planned_date=D(2024, 3, 1),
                 status="completed", on_time=True, cost_rwf=300_000),
            dict(project_id="PRJ04", district="Kirehe", date=D(2024, 3, 9), planned_date=D(2024, 3, 9),
                 status="completed", on_time=True, cost_rwf=None)]
    services = [service("BEN-1", "female", D(2024, 3, 2)), service("BEN-2", "male", D(2024, 3, 2)),
                service("BEN-1", "female", D(2024, 3, 9))]
    values, _ = run_method(indicators.cost_per_person, FakeData(activities=acts, services=services), "CC-03")
    v = pick(values, period_type="month", period="2024-03", project_id="PRJ04", district="Kirehe")
    assert v["actual"] == 150_000                  # 300,000 RWF / 2 different people
    assert v["coverage_pct"] == 50.0              # 1 of 2 activities has a cost
    assert CATALOGUE["CC-03"]["higher_is_better"] is False


def test_reporting_completeness():
    rc = [("PRJ01", "Huye", D(2024, 3, 1), "WASH-01", True), ("PRJ01", "Huye", D(2024, 4, 1), "WASH-01", False)]
    values, links = run_method(indicators.report_completeness, FakeData(completeness=rc), "CC-04")
    v = pick(values, period_type="quarter", period="FY2023/24 Q3", project_id="PRJ01", district="Huye")
    assert v["actual"] == 100.0
    v = pick(values, period_type="fiscal_year", period="FY2023/24", project_id="All", district="All")
    assert v["actual"] == 50.0
    assert set(links["record_key"].dropna()) == {"PRJ01|Huye|2024-03-01", "PRJ01|Huye|2024-04-01"}


# ---------------------------------------------------------------- reported values, targets and status
def result(month, target, actual, sex=None):
    return dict(indicator_code="AGR-01", project_id="PRJ04", district="Kirehe", period_date=D(2024, month, 1),
                target_value=target, actual_value=actual, disaggregation_key=f"sex={sex}" if sex else "total")


def test_reported_values_add_up_with_targets_and_traffic_lights():
    rows = [result(m, 10, a, sex) for m, a in ((1, 10), (2, 8), (3, 6)) for sex in ("female", "male")]
    rc = [("PRJ04", "Kirehe", D(2024, m, 1), "AGR-01", True) for m in (1, 2, 3)]
    values, _ = run_method(indicators.reported, FakeData(results=rows, completeness=rc), "AGR-01")
    q = pick(values, period_type="quarter", period="FY2023/24 Q3", project_id="PRJ04", district="Kirehe",
             disaggregation="total")
    assert q["actual"] == 48 and q["target"] == 60 and q["pct_achieved"] == 80.0 and q["status"] == "at risk"
    jan = pick(values, period_type="month", period="2024-01", project_id="PRJ04", district="Kirehe",
               disaggregation="sex=female")
    assert jan["status"] == "on track" and jan["coverage_pct"] == 100.0
    mar = pick(values, period_type="month", period="2024-03", project_id="PRJ04", district="Kirehe",
               disaggregation="total")
    assert mar["pct_achieved"] == 60.0 and mar["status"] == "off track"


def test_missing_target_shows_no_target_set_never_0_or_100_percent():
    rows = [result(1, None, 12, "female"), result(1, None, 9, "male"), result(2, 10, 10, "female"),
            result(2, 10, 10, "male")]
    values, _ = run_method(indicators.reported, FakeData(results=rows), "AGR-01")
    jan = pick(values, period_type="month", period="2024-01", project_id="PRJ04", district="Kirehe",
               disaggregation="total")
    assert jan["status"] == S["indicators"]["no_target_label"] and pd.isna(jan["pct_achieved"])
    q = pick(values, period_type="quarter", period="FY2023/24 Q3", project_id="PRJ04", district="Kirehe",
             disaggregation="total")
    assert q["status"] == S["indicators"]["target_incomplete_label"]
    assert pd.isna(q["pct_achieved"]) and pd.isna(q["target"])        # no RAG colour
    assert q["partial_pct_achieved"] == 100.0                           # February: 20 of 20
    assert q["target_note"] == "partial: 1 of 2 months"


def test_an_expected_report_that_never_arrived_shows_no_data():
    rc = [("PRJ04", "Kirehe", D(2024, 1, 1), "AGR-01", True), ("PRJ04", "Kirehe", D(2024, 2, 1), "AGR-01", False)]
    rows = [result(1, 10, 10, "female"), result(1, 10, 10, "male")]
    values, _ = run_method(indicators.reported, FakeData(results=rows, completeness=rc), "AGR-01")
    feb = pick(values, period_type="month", period="2024-02", project_id="PRJ04", district="Kirehe",
               disaggregation="total")
    assert feb["status"] == "no data" and feb["coverage_pct"] == 0.0


# ---------------------------------------------------------------- reconciliation
def test_reconciliation_flags_gaps_above_the_tolerance():
    tol = S["indicators"]["reconciliation_tolerance_pct"]
    builder = indicators.IndicatorBuilder.__new__(indicators.IndicatorBuilder)
    builder.settings, builder.catalogue = S, list(CATALOGUE.values())
    base = dict(indicator_code="WASH-01", indicator_id=1, project_id="PRJ01", district="Huye", period_type="month",
                period_start=D(2024, 1, 1), disaggregation="total")
    rows = []
    for i, (period, rep, der) in enumerate((("2024-01", 100, 100 + tol / 2), ("2024-02", 100, 100 + 3 * tol),
                                            ("2024-03", 100, None), ("2024-04", None, 40))):
        if rep is not None:
            rows.append(dict(base, period=period, source="reported", actual=rep, value_id=2 * i + 1))
        if der is not None:
            rows.append(dict(base, period=period, source="derived", actual=der, value_id=2 * i + 2))
    rec = builder.reconcile(pd.DataFrame(rows)).set_index("period")
    assert rec.loc["2024-01", "status"] == "within tolerance" and not rec.loc["2024-01", "flag"]
    assert rec.loc["2024-02", "status"] == "gap above tolerance" and rec.loc["2024-02", "flag"]
    assert rec.loc["2024-02", "pct_difference"] == 3 * tol
    assert rec.loc["2024-03", "status"] == "gap above tolerance" and rec.loc["2024-03", "derived"] == 0
    assert rec.loc["2024-04", "status"] == "derived only"


# ---------------------------------------------------------------- the real pipeline output
@pytest.fixture(scope="module")
def built(synthetic, tmp_path_factory):
    root = tmp_path_factory.mktemp("phase4")
    mp = pytest.MonkeyPatch()
    mp.setenv("ARDE_SALT_FILE", str(root.parent / "phase4_salt.txt"))
    staging, docs = root / "staging", root / "docs"
    docs.mkdir()
    ingest.run(synthetic["raw"], staging)
    clean.run(staging)
    validate.run(staging, docs)
    model.run(staging, root / "arde.duckdb", root / "exports", docs)
    tables = indicators.run(root / "arde.duckdb", root / "exports", docs)
    mp.undo()
    return dict(tables=tables, db=root / "arde.duckdb", docs=docs, exports=root / "exports")


def test_every_available_indicator_is_calculated_and_outcomes_are_not(built):
    codes = set(built["tables"]["indicator_values"]["indicator_code"])
    available = {c for c, i in CATALOGUE.items() if i["availability"] == "available"}
    assert codes == available and not {"WASH-03", "AGR-03"} & codes


def test_no_value_without_a_target_has_a_percentage(built):
    v = built["tables"]["indicator_values"]
    assert v[v["target"].isna()]["pct_achieved"].isna().all()
    assert set(v["status"]) <= {"on track", "at risk", "off track", "no target set", "target incomplete", "no data"}


def test_every_value_traces_back_to_real_fact_rows(built):
    v, links = built["tables"]["indicator_values"], built["tables"]["indicator_value_sources"]
    with_data = v[v["actual"].notna()]
    assert set(with_data["value_id"]) <= set(links["value_id"])
    con = duckdb.connect(str(built["db"]), read_only=True)
    facts = {t: {(f, s, int(r)) for f, s, r in con.execute(
        f"SELECT source_file, source_sheet, source_row FROM {t}").fetchall()}
        for t in ("fact_activity", "fact_beneficiary_service", "fact_indicator_result")}
    con.close()
    sample = links[links["source_table"] != "report_completeness"].sample(500, random_state=1)
    for r in sample.itertuples():
        assert (r.source_file, r.source_sheet, int(r.source_row)) in facts[r.source_table]


def test_a_derived_people_count_matches_its_source_rows(built):
    v, links = built["tables"]["indicator_values"], built["tables"]["indicator_value_sources"]
    month = v[(v["indicator_code"] == "AGR-01") & (v["source"] == "derived") & (v["project_id"] == "All")
              & (v["district"] == "All") & (v["period_type"] == "month") & (v["disaggregation"] == "total")]
    row = month.sort_values("actual").iloc[-1]
    mine = links[links["value_id"] == row["value_id"]]
    con = duckdb.connect(str(built["db"]), read_only=True)
    services = con.execute("SELECT source_file, source_sheet, source_row, beneficiary_key FROM fact_beneficiary_service").df()
    con.close()
    used = services.merge(mine, on=["source_file", "source_sheet", "source_row"])
    assert used["beneficiary_key"].nunique() == row["actual"]
    assert len(used) == row["n_rows"] >= row["actual"]


def test_coverage_of_on_time_uses_the_single_definition(built):
    # CC-02 coverage = the data-quality measure "on_time_checkable" (config/dq_metrics.yaml), applied to the
    # activities kept in the warehouse (the measure itself is over every activity received, quarantined ones too).
    v = built["tables"]["indicator_values"]
    cc02 = v[(v["indicator_code"] == "CC-02") & (v["project_id"] == "All") & (v["district"] == "All")
             & (v["period_type"] == "fiscal_year")]
    con = duckdb.connect(str(built["db"]), read_only=True)
    facts = con.execute("SELECT * FROM fact_activity").df()
    con.close()
    spec = load_yaml("dq_metrics.yaml")["metrics"]["on_time_checkable"]
    # Only activities with a planned date or an activity date can be placed in a period.
    dated = facts[facts["planned_date"].notna() | facts["date"].notna()]
    num, den = validate.measure(dated, spec, S)
    assert cc02["denominator"].sum() == num and cc02["n_rows"].sum() == den


def test_reconciliation_is_built_and_flags_large_gaps(built):
    rec = built["tables"]["indicator_reconciliation"]
    assert set(rec["indicator_code"]) == {"WASH-01", "AGR-01"}
    tol = S["indicators"]["reconciliation_tolerance_pct"]
    both = rec[rec["reported"].notna() & rec["derived"].notna()]
    assert (both["flag"] == (both["pct_difference"].abs() > tol)).all()


def test_reference_sheets_markdown_and_excel(built):
    md = (built["docs"] / "indicator_reference_sheets.md").read_text(encoding="utf-8")
    for code in CATALOGUE:
        assert f"## {code}:" in md
    assert "Draft, for co-design with ARDE/KUBAHO staff" in md
    ws = load_workbook(built["docs"] / "indicator_reference_sheets.xlsx")["Indicators"]
    headers = [c.value for c in ws[2]]
    assert headers[-1] == "Staff comments"
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    assert len(rows) == len(CATALOGUE) and all(r[-1] in (None, "") for r in rows)


def test_indicator_tables_are_exported_and_described(built):
    for name in indicators.OUTPUT_TABLES:
        assert (built["exports"] / f"{name}.csv").exists() and (built["exports"] / f"{name}.parquet").exists()
    text = (built["docs"] / "data_dictionary.md").read_text(encoding="utf-8")
    assert "## `indicator_reconciliation`" in text


# ---------------------------------------------------------------- aggregation methods (catalogue: aggregation)
def quarterly(code, values, project="PRJ01", district="Huye"):
    return [result_row(code, m, v, project, district) for m, v in values]


def result_row(code, month, value, project, district, target=10):
    return dict(indicator_code=code, project_id=project, district=district, period_date=D(2024, month, 1),
                target_value=target, actual_value=value, disaggregation_key="total")


@pytest.mark.parametrize("aggregation, expected_q3, expected_fy", [
    ("sum", 3 + 5 + 4, 3 + 5 + 4 + 6),
    ("latest", 4, 6),                         # a snapshot: the latest month in the period
    ("average", (3 + 5 + 4) / 3, (3 + 5 + 4 + 6) / 4),
])
def test_each_aggregation_method_rolls_months_up(aggregation, expected_q3, expected_fy):
    spec = {**CATALOGUE["WASH-02"], "aggregation": aggregation}
    rows = quarterly("WASH-02", [(1, 3), (2, 5), (3, 4), (4, 6)])
    values, _ = run_method(indicators.reported, FakeData(results=rows), "WASH-02", spec)
    q = pick(values, period_type="quarter", period="FY2023/24 Q3", project_id="PRJ01", district="Huye")
    fy = pick(values, period_type="fiscal_year", period="FY2023/24", project_id="PRJ01", district="Huye")
    assert q["actual"] == pytest.approx(expected_q3, abs=0.01)
    assert fy["actual"] == pytest.approx(expected_fy, abs=0.01)


def test_latest_uses_the_latest_target_too():
    spec = {**CATALOGUE["WASH-04"]}
    assert spec["aggregation"] == "latest"
    rows = [result_row("WASH-04", 3, 2, "PRJ01", "Huye", target=4), result_row("WASH-04", 6, 3, "PRJ01", "Huye", target=5)]
    values, links = run_method(indicators.reported, FakeData(results=rows), "WASH-04", spec)
    fy = pick(values, period_type="fiscal_year", period="FY2023/24", project_id="PRJ01", district="Huye")
    assert fy["actual"] == 3 and fy["target"] == 5 and fy["pct_achieved"] == 60.0     # not 5 of 9


def test_ratio_is_worked_out_again_for_the_whole_period():
    # CC-01: one woman in January and again in February, one man in February only.
    # Months: 100% then 50%. The quarter is 1 woman of 2 different people = 50%, not the average 75%.
    services = [service("BEN-W", "female", D(2024, 1, 5)), service("BEN-W", "female", D(2024, 2, 5)),
                service("BEN-M", "male", D(2024, 2, 6))]
    assert CATALOGUE["CC-01"]["aggregation"] == "ratio"
    values, _ = run_method(indicators.share_women, FakeData(services=services), "CC-01")
    q = pick(values, period_type="quarter", period="FY2023/24 Q3", project_id="All", district="All")
    assert q["actual"] == 50.0


def test_every_indicator_has_a_known_aggregation_method():
    for code, ind in CATALOGUE.items():
        assert ind["aggregation"] in {"sum", "latest", "average", "ratio"}, code
    assert CATALOGUE["WASH-04"]["aggregation"] == "latest"     # until the partner confirms


def test_like_for_like_people_served_from_activity_records():
    acts = [dict(project_id="PRJ01", district="Huye", date=D(2024, 3, 2), planned_date=D(2024, 3, 1),
                 status="completed", on_time=True, cost_rwf=1, activity_id="ACT01", participants_total=20),
            dict(project_id="PRJ01", district="Huye", date=D(2024, 3, 9), planned_date=D(2024, 3, 9),
                 status="delayed", on_time=False, cost_rwf=1, activity_id="ACT01", participants_total=15),
            dict(project_id="PRJ01", district="Huye", date=D(2024, 3, 9), planned_date=D(2024, 3, 9),
                 status="completed", on_time=True, cost_rwf=1, activity_id="ACT02", participants_total=99)]
    values, _ = run_method(indicators.activity_participants, FakeData(activities=acts), "WASH-01")
    assert pick(values, period_type="month", period="2024-03", project_id="PRJ01", district="Huye")["actual"] == 35
    assert CATALOGUE["WASH-01"]["calculation"]["derived"]["like_for_like"] is True


def test_reconciliation_finds_the_planted_discrepancies(built, synthetic):
    import score_cleaning
    r = score_cleaning.score_reconciliation(synthetic["truth"], built["db"])
    s = r["summary"]
    assert s["planted"] > 0 and s["catch_rate"] >= 95
    assert s["false_alarms"] == 0
    assert s["rounding_within"] == s["rounding_planted"]
