"""Phase 5: every dashboard page runs, shows the banner and labels statuses in words; the donor report."""

from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

import donor_report
from common import load_settings

PAGES = [None, "pages/1_Programme_Overview.py", "pages/2_Indicators.py", "pages/3_Geography.py",
         "pages/4_Data_Quality.py", "pages/5_Donor_Pipeline.py"]
BANNER = load_settings()["synthetic_banner"]
APP = Path(__file__).resolve().parent.parent / "dashboard" / "app.py"


def run_page(page):
    at = AppTest.from_file(str(APP), default_timeout=120)
    if page:
        at.switch_page(page)
    at.run()
    return at


@pytest.fixture(scope="module", autouse=True)
def warehouse_exists():
    from common import resolve_path
    if not resolve_path("warehouse").exists():
        pytest.skip("Build the warehouse first: python src/pipeline.py")


@pytest.mark.parametrize("page", PAGES)
def test_every_page_runs_and_shows_the_synthetic_banner(page):
    at = run_page(page)
    assert not at.exception, [e.value for e in at.exception]
    assert any(BANNER in m.value for m in at.markdown)


def test_executive_summary_answers_the_three_questions():
    at = run_page(None)
    text = " ".join(m.value for m in at.markdown)
    for question in ("Are we on track this quarter?", "Can we trust these numbers?", "Funding: the next 90 days"):
        assert question in text
    assert "✔ On track" in text and "✖ Off track" in text        # never colour alone


def test_indicator_statuses_are_in_words():
    at = run_page("pages/2_Indicators.py")
    import re
    table = next(m.value for m in at.markdown if "class='tbl'" in m.value)
    rows = re.findall(r"<tr>(.*?)</tr>", table.split("<tbody>")[1])
    assert rows
    for row in rows:      # every row's status has a symbol AND words, never colour alone
        assert re.search(r"[✔▲✖◔–…] (On track|At risk|Off track|No target set|Target incomplete|No data)", row)


def test_sex_filter_is_greyed_out_where_it_does_not_apply():
    at = run_page("pages/4_Data_Quality.py")
    sex = [r for r in at.sidebar.radio if r.label == "Sex"][0]
    assert sex.disabled
    at = run_page("pages/1_Programme_Overview.py")
    assert not [r for r in at.sidebar.radio if r.label == "Sex"][0].disabled


def test_people_counts_mention_pending_duplicate_review():
    at = run_page("pages/1_Programme_Overview.py")
    assert any("pending duplicate review" in m.value for m in at.markdown)


def test_donor_report_is_one_page_in_excel_and_pdf(tmp_path, monkeypatch):
    monkeypatch.setattr(donor_report, "LOG", tmp_path / "timings.csv")      # keep the real timing log clean
    result = donor_report.make_report("PRJ01", "FY2025/26 Q4", out_dir=tmp_path)
    assert result["pages"] == 1 and result["pdf"].stat().st_size > 1000
    ws = load_workbook(result["excel"])["Summary"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Quarterly progress report" in text and BANNER in text
    assert result["seconds"] < 30
    assert donor_report.LOG.exists()
