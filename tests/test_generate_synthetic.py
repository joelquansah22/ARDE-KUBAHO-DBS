"""Checks on the synthetic data generator (Phase 1)."""

import datetime as dt
import math

import pandas as pd
import pytest
from openpyxl import load_workbook

import generate_synthetic

EXPECTED_FILES = {
    "SYNTHETIC_programme_master.xlsx",
    "SYNTHETIC_activity_monitoring_FO01.xlsx",
    "SYNTHETIC_activity_monitoring_FO02.xlsx",
    "SYNTHETIC_activity_monitoring_FO03.xlsx",
    "SYNTHETIC_activity_monitoring_FO04.xlsx",
    "SYNTHETIC_beneficiary_register_FO01.xlsx",
    "SYNTHETIC_beneficiary_register_FO02.xlsx",
    "SYNTHETIC_beneficiary_register_FO03.xlsx",
    "SYNTHETIC_beneficiary_register_FO04.xlsx",
    "SYNTHETIC_digitised_paper_records.xlsx",
    "SYNTHETIC_indicator_reports_WASH.xlsx",
    "SYNTHETIC_indicator_reports_AGRI.xlsx",
    "SYNTHETIC_donor_opportunity_tracker.xlsx",
}

# Every kind of mess named in brief Section 6 (plus the timeliness additions) must be present.
REQUIRED_ERROR_TYPES = [
    "date_dmy_text", "date_month_year_text", "date_excel_serial",      # mixed date formats
    "district_variant", "district_typo",                               # district names
    "sex_variant", "sex_ambiguous_code_1",                             # sex coding
    "duplicate_beneficiary", "duplicate_activity_row",                 # duplicates
    "missing_target", "negative_value", "female_exceeds_total",        # impossible values
    "percent_over_100", "sum_mismatch",
    "number_as_text", "amount_in_usd",                                 # numbers as text
    "transcription_error", "missing_required",                         # paper records
    "late_entry", "status_completed_but_late",                         # timeliness
    "missing_monthly_report", "date_out_of_window", "district_wrong_province",
    "totals_row", "merged_header_row", "blank_row", "title_row",       # layout problems
    "personal_data_individual_donor",
]


@pytest.fixture(scope="module")
def log(synthetic):
    return pd.read_csv(synthetic["truth"] / "SYNTHETIC_injected_errors_log.csv", dtype=str, keep_default_na=False)


def test_all_files_created_with_synthetic_prefix(synthetic):
    created = {p.name for p in synthetic["raw"].glob("*.xlsx")}
    assert created == EXPECTED_FILES
    assert all(p.name.startswith("SYNTHETIC_") for p in synthetic["truth"].iterdir())


def test_answer_key_is_kept_out_of_the_raw_folder(synthetic):
    assert not list(synthetic["raw"].glob("*.csv"))
    assert (synthetic["truth"] / "SYNTHETIC_truth_activities.csv").exists()


def test_volumes_match_the_brief(synthetic):
    assert 2700 <= synthetic["activity_records"] <= 3300
    assert synthetic["people"] == 1500
    assert synthetic["registrations"] > synthetic["people"]          # duplicates were added
    assert synthetic["donors"] == 25
    assert synthetic["opportunities"] == 30
    acts = pd.read_csv(synthetic["truth"] / "SYNTHETIC_truth_activities.csv")
    assert 6 <= acts["district"].nunique() <= 8
    assert acts["project_id"].nunique() in range(4, 7)


@pytest.mark.parametrize("error_type", REQUIRED_ERROR_TYPES)
def test_every_mess_type_is_injected(log, error_type):
    assert (log["error_type"] == error_type).sum() > 0, f"no {error_type} injected"


def test_summary_counts_match_the_log(synthetic, log):
    summary = pd.read_csv(synthetic["truth"] / "SYNTHETIC_injected_errors_summary.csv")
    assert summary["count"].sum() == len(log)


def _same(cell_value, logged: str) -> bool:
    """Compare a cell with the value written in the CSV log."""
    if cell_value is None:
        return logged == ""
    if isinstance(cell_value, dt.datetime):
        return logged.startswith(cell_value.date().isoformat())
    if str(cell_value) == logged:
        return True
    try:
        return math.isclose(float(cell_value), float(logged))
    except ValueError:
        return False


def test_every_logged_cell_really_contains_the_messy_value(synthetic, log):
    cell_level = log[(log["column_letter"] != "") & (log["error_type"] != "duplicate_activity_row")]
    assert len(cell_level) > 1000
    books = {}
    mismatches = []
    for row in cell_level.itertuples():
        if row.source_file not in books:
            books[row.source_file] = load_workbook(synthetic["raw"] / row.source_file)
        ws = books[row.source_file][row.source_sheet]
        value = ws[f"{row.column_letter}{int(float(row.source_row))}"].value
        if not _same(value, row.messy_value):
            mismatches.append((row.error_id, row.error_type, value, row.messy_value))
    assert not mismatches, mismatches[:10]


def test_female_above_total_really_exceeds_total(synthetic, log):
    rows = log[log["error_type"] == "female_exceeds_total"]
    for row in rows.itertuples():
        assert "female participation" in row.detail
        assert int(row.detail.split("%")[0]) > 100


def test_duplicate_activity_rows_are_exact_copies(synthetic, log):
    rows = log[log["error_type"] == "duplicate_activity_row"].head(15)
    books = {}
    for row in rows.itertuples():
        wb = books.setdefault(row.source_file, load_workbook(synthetic["raw"] / row.source_file))
        ws = wb[row.source_sheet]
        original = int(row.detail.rsplit(" ", 1)[1])
        copy = int(float(row.source_row))
        assert [c.value for c in ws[original]] == [c.value for c in ws[copy]]


def test_answer_key_positions_are_unique_and_complete(synthetic):
    acts = pd.read_csv(synthetic["truth"] / "SYNTHETIC_truth_activities.csv")
    assert acts["source_row"].notna().all()
    assert not acts.duplicated(["source_file", "source_sheet", "source_row"]).any()
    regs = pd.read_csv(synthetic["truth"] / "SYNTHETIC_truth_registrations.csv")
    assert regs["source_row"].notna().all()
    assert regs.loc[regs["is_duplicate"], "person_id"].isin(regs.loc[~regs["is_duplicate"], "person_id"]).all()


def test_timeliness_fields_are_generated(synthetic):
    acts = pd.read_csv(synthetic["truth"] / "SYNTHETIC_truth_activities.csv", parse_dates=["date", "planned_date", "entry_date"])
    assert set(acts["status"]) == {"planned", "completed", "delayed", "cancelled"}
    assert acts["is_late_entry"].sum() > 0
    done_late = acts[(acts["status"] == "delayed") & acts["date"].notna()]
    assert ((done_late["date"] - done_late["planned_date"]).dt.days > 7).all()


def test_same_seed_gives_identical_output(synthetic, tmp_path):
    generate_synthetic.generate(tmp_path / "raw", tmp_path / "truth", quiet=True)
    for name in ["SYNTHETIC_injected_errors_log.csv", "SYNTHETIC_truth_activities.csv",
                 "SYNTHETIC_truth_registrations.csv"]:
        assert (tmp_path / "truth" / name).read_bytes() == (synthetic["truth"] / name).read_bytes()
