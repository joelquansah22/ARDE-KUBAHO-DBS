from openpyxl import load_workbook

import build_templates
from common import load_yaml, resolve_path


def test_templates_have_headers_and_dropdowns():
    build_templates.main()
    cfg = load_yaml("templates.yaml")["templates"]
    for file_name, spec in cfg.items():
        wb = load_workbook(resolve_path("templates") / file_name)
        assert wb["Lists"].sheet_state == "hidden"
        for sheet, columns in spec["sheets"].items():
            ws = wb[sheet]
            assert [c.value for c in ws[1]] == [c["header"] for c in columns]
            if any("list" in c for c in columns):
                assert ws.data_validations.dataValidation, f"{file_name}/{sheet} has no dropdowns"


def test_sex_dropdown_has_no_numeric_codes():
    values = build_templates.resolve_list(load_yaml("templates.yaml")["lists"]["sex"])
    assert values == ["Female", "Male"]


def test_templates_are_not_rebuilt_when_nothing_changed():
    build_templates.main()
    assert build_templates.main() is False
