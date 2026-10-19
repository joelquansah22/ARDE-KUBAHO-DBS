"""Scope extension: donor tracking entered on the dashboard (data/app/entries.db)."""

import datetime as dt
import sqlite3
from pathlib import Path

import pandas as pd
import pytest

import app_store
from app_store import EntryError

TODAY = app_store.today()
SOON = TODAY + dt.timedelta(days=40)


@pytest.fixture
def db(tmp_path):
    return tmp_path / "entries.db"


def opportunity(**kw):
    base = dict(title="Safe water for schools", donor_key="DON-abc", thematic_fit="WASH", stage="Concept note",
                deadline=SOON, owner="PM01", requested_amount_rwf=150_000_000, probability=0.3)
    return {**base, **kw}


# ---------------------------------------------------------------- validation on entry
@pytest.mark.parametrize("change, message", [
    (dict(title=""), "Title is required"),
    (dict(donor_key=None), "Donor is required"),
    (dict(probability=1.2), "between 0% and 100%"),
    (dict(probability=-0.1), "between 0% and 100%"),
    (dict(deadline=TODAY), "must be in the future"),
    (dict(deadline=None), "Deadline is required"),
    (dict(requested_amount_rwf=-5), "cannot be negative"),
    (dict(stage="Awarded"), "needs the amount awarded"),
    (dict(next_action="Send concept note"), "Give the date"),
    (dict(stage="Maybe"), "Stage must be one of"),
])
def test_opportunity_rules_are_checked_on_entry(db, change, message):
    with pytest.raises(EntryError, match=message):
        app_store.add_opportunity(opportunity(**change), "PM01", db)


def test_staff_must_be_chosen(db):
    with pytest.raises(EntryError, match="Choose your name"):
        app_store.add_opportunity(opportunity(), "Somebody", db)


def test_a_valid_opportunity_is_saved_with_its_author(db):
    opp_id = app_store.add_opportunity(opportunity(), "PM01", db)
    assert opp_id == "APP-O001"
    row = app_store.table("opportunities", db)[0]
    assert (row["created_by"], row["probability"], row["deadline"]) == ("PM01", 0.3, SOON.isoformat())
    assert app_store.table("change_history", db)[0]["action"] == "added"


# ---------------------------------------------------------------- individual donors and personal data
def test_individual_donors_are_never_named(db):
    problems = app_store.validate_donor({"donor_type": "individual", "donor_name": "Jane Doe"})
    assert any("Do not enter an individual donor's name" in p for p in problems)
    app_store.add_donor({"donor_type": "individual", "donor_name": "Jane Doe"}, "ED01", db)   # the name is dropped
    assert app_store.table("donors", db)[0]["donor_name"] is None
    donor_id = app_store.table("donors", db)[0]["donor_id"]
    # Even writing straight to the database cannot store a name for an individual.
    with sqlite3.connect(db) as con, pytest.raises(sqlite3.IntegrityError):
        con.execute("INSERT INTO donors (donor_id, donor_name, donor_type, created_by, created_at) "
                    "VALUES ('X', 'Jane Doe', 'individual', 'ED01', 'now')")
    assert donor_id == "APP-D001"


def test_an_organisation_needs_its_name(db):
    with pytest.raises(EntryError, match="Donor name is required"):
        app_store.add_donor({"donor_type": "foundation", "donor_name": " "}, "ED01", db)


def test_the_database_holds_no_personal_data_columns(db):
    from model import PERSONAL_COLUMNS
    app_store.add_donor({"donor_type": "foundation", "donor_name": "Example Foundation"}, "ED01", db)
    with sqlite3.connect(db) as con:
        columns = {c[1] for t in ("donors", "opportunities", "change_history")
                   for c in con.execute(f"PRAGMA table_info({t})")}
    banned = (PERSONAL_COLUMNS - {"name"}) | {"email", "phone", "contact_person", "national_id"}
    assert not columns & banned


# ---------------------------------------------------------------- updates and history
def test_updates_record_who_changed_what_and_when(db):
    opp_id = app_store.add_opportunity(opportunity(), "PM01", db)
    changed = app_store.update_opportunity(opp_id, dict(stage="Proposal submitted", probability=0.5), "MEAL01", db)
    assert set(changed) == {"stage", "probability"}
    history = pd.DataFrame(app_store.table("change_history", db))
    edits = history[history["action"] == "changed"].set_index("field")
    assert edits.loc["stage", "old_value"] == "Concept note" and edits.loc["stage", "new_value"] == "Proposal submitted"
    assert (edits["changed_by"] == "MEAL01").all() and edits["changed_at"].notna().all()
    assert app_store.update_opportunity(opp_id, dict(stage="Proposal submitted"), "MEAL01", db) == []   # no change


def test_updates_are_validated_too(db):
    opp_id = app_store.add_opportunity(opportunity(), "PM01", db)
    with pytest.raises(EntryError, match="needs the amount awarded"):
        app_store.update_opportunity(opp_id, dict(stage="Awarded"), "PM01", db)
    with pytest.raises(EntryError, match="cannot be changed here"):
        app_store.update_opportunity(opp_id, dict(title="New title"), "PM01", db)


def test_marking_a_next_action_done_keeps_it_in_the_history(db):
    opp_id = app_store.add_opportunity(opportunity(next_action="Send concept note", next_action_date=SOON), "PM01", db)
    app_store.mark_next_action_done(opp_id, "PM01", db)
    row = app_store.table("opportunities", db)[0]
    assert row["next_action"] is None and row["next_action_date"] is None
    done = [h for h in app_store.table("change_history", db) if h["action"] == "next action done"]
    assert "Send concept note" in done[0]["old_value"]
    with pytest.raises(EntryError, match="no next action"):
        app_store.mark_next_action_done(opp_id, "PM01", db)


# ---------------------------------------------------------------- the pipeline reads both sources
def test_pipeline_reads_entries_next_to_the_excel_tracker(synthetic, tmp_path):
    import ingest
    db = tmp_path / "entries.db"
    individual = app_store.add_donor({"donor_type": "individual"}, "ED01", db)
    foundation = app_store.add_donor({"donor_type": "foundation", "donor_name": "Example Foundation"}, "ED01", db)
    app_store.add_opportunity(opportunity(donor_key=None, donor_id=foundation), "PM01", db)
    without = ingest.run(synthetic["raw"], tmp_path / "staging_without")["tables"]
    both = ingest.run(synthetic["raw"], tmp_path / "staging_with", app_entries=db)["tables"]
    assert len(both["opportunity"]) == len(without["opportunity"]) + 1
    app_opp = both["opportunity"][both["opportunity"]["source_file"] == "entries.db"].iloc[0]
    assert app_opp["donor_key"] == ingest.pseudonym(ingest.load_salt(), "DON", foundation)
    # Excel individuals keep their numbers; the one added on the dashboard comes after them.
    excel = without["donor"].set_index("donor_key")["donor_name"]
    after = both["donor"].set_index("donor_key")["donor_name"]
    assert (after.loc[excel.index] == excel).all()
    app_label = both["donor"].loc[both["donor"]["donor_id"] == individual, "donor_name"].iloc[0]
    assert app_label.startswith("Individual donor #")
    assert app_label not in set(excel)


def test_entries_are_never_read_into_a_scoring_run(synthetic, tmp_path, monkeypatch):
    import ingest
    db = tmp_path / "entries.db"
    app_store.add_donor({"donor_type": "foundation", "donor_name": "Example Foundation"}, "ED01", db)
    monkeypatch.setattr(app_store, "db_path", lambda: db)
    tables = ingest.run(synthetic["raw"], tmp_path / "staging")["tables"]      # a test-set run: raw_dir given
    assert "entries.db" not in set(tables["donor"]["source_file"])


# ---------------------------------------------------------------- the page
def test_donor_entry_page_runs(tmp_path, monkeypatch):
    from common import resolve_path
    if not resolve_path("warehouse").exists():
        pytest.skip("Build the warehouse first")
    from streamlit.testing.v1 import AppTest
    db = tmp_path / "entries.db"
    app_store.add_opportunity(opportunity(next_action="Call donor", next_action_date=SOON), "PM01", db)
    monkeypatch.setattr(app_store, "db_path", lambda: db)
    at = AppTest.from_file(str(Path(__file__).resolve().parent.parent / "dashboard" / "app.py"), default_timeout=120)
    at.switch_page("pages/6_Donor_Entry.py")
    at.session_state["entry_staff"] = "PM01"
    at.run()
    assert not at.exception, [e.value for e in at.exception]
    assert any("APP-O001" in m.value for m in at.markdown)            # listed under next actions


# ---------------------------------------------------------------- backups, export and restore
def test_every_save_keeps_a_dated_backup_and_only_the_newest_are_kept(db, monkeypatch):
    settings = dict(app_store.load_settings(), app_backups_keep=3)
    monkeypatch.setattr(app_store, "load_settings", lambda: settings)
    for i in range(5):
        app_store.add_donor({"donor_type": "foundation", "donor_name": f"Foundation {i}"}, "ED01", db)
    backups = app_store.list_backups(db)
    assert len(backups) == 3 and all(b.name.endswith("_save.db") for b in backups)
    import sqlite3
    with sqlite3.connect(backups[0]) as con:                       # the newest copy has all five donors
        assert con.execute("SELECT COUNT(*) FROM donors").fetchone()[0] == 5


def test_restore_puts_a_backup_back_and_keeps_a_safety_copy(db):
    app_store.add_opportunity(opportunity(), "PM01", db)
    first = app_store.list_backups(db)[0]
    app_store.add_opportunity(opportunity(title="Second"), "PM01", db)
    assert len(app_store.table("opportunities", db)) == 2
    safety = app_store.restore(first, db)
    assert [o["title"] for o in app_store.table("opportunities", db)] == ["Safe water for schools"]
    assert safety.exists() and "before-restore" in safety.name
    app_store.restore(safety, db)                                  # a restore can itself be undone
    assert len(app_store.table("opportunities", db)) == 2


def test_restore_refuses_a_file_that_is_not_a_backup(db, tmp_path):
    other = tmp_path / "other.db"
    sqlite3.connect(other).execute("CREATE TABLE x (a)").connection.close()
    with pytest.raises(EntryError, match="not a donor entries backup"):
        app_store.restore(other, db)


def test_download_backup_holds_every_table_and_the_history(db):
    from io import BytesIO
    from openpyxl import load_workbook
    app_store.add_donor({"donor_type": "individual"}, "ED01", db)
    app_store.add_opportunity(opportunity(), "PM01", db)
    wb = load_workbook(BytesIO(app_store.export_workbook(db)))
    assert wb.sheetnames == ["About", "Donors", "Opportunities", "Activities", "Indicator reports", "Uploads",
                             "Change history"]
    assert wb["Opportunities"].max_row == 2 and wb["Change history"].max_row == 3
    assert wb["Donors"]["C2"].value is None                       # donor_name of the individual: empty


# ---------------------------------------------------------------- one home: importing the Excel tracker
@pytest.fixture(scope="module")
def tracker_staging(synthetic, tmp_path_factory):
    import clean
    import ingest
    import validate
    staging = tmp_path_factory.mktemp("tracker_staging")
    ingest.run(synthetic["raw"], staging)
    clean.run(staging)
    validate.run(staging, tmp_path_factory.mktemp("docs"))
    return staging


def test_import_preview_lists_what_moves_and_what_must_be_fixed_in_excel(tracker_staging, db):
    preview = app_store.import_preview(tracker_staging, db)
    clean_o = pd.read_parquet(tracker_staging / "clean_opportunity.parquet")
    assert len(preview["opportunities"]) + len(preview["problems"]) == len(clean_o)
    assert len(preview["opportunities"]) > 0 and len(preview["donors"]) > 0
    assert app_store.table("opportunities", db) == []              # a preview changes nothing


def test_import_keeps_ids_and_lineage_and_the_pipeline_stops_reading_those_excel_rows(synthetic, tracker_staging,
                                                                                         tmp_path):
    import ingest
    db = tmp_path / "entries.db"
    before = ingest.run(synthetic["raw"], tmp_path / "before", app_entries=db)["tables"]
    result = app_store.import_tracker("MEAL01", tracker_staging, db)
    opps = pd.DataFrame(app_store.table("opportunities", db))
    assert result["opportunities"] == len(opps) and opps["opportunity_id"].str.startswith("OPP").all()
    assert opps["origin_file"].eq("SYNTHETIC_donor_opportunity_tracker.xlsx").all() and opps["origin_row"].notna().all()
    donors = pd.DataFrame(app_store.table("donors", db))
    assert donors.loc[donors["donor_type"] == "individual", "donor_name"].isna().all()   # still never named
    assert app_store.list_backups(db)[0].name.endswith("_import.db")
    after = ingest.run(synthetic["raw"], tmp_path / "after", app_entries=db)
    t = after["tables"]
    # Not counted twice: every opportunity once, now from entries.db (except those left in Excel to correct).
    assert len(t["opportunity"]) == len(before["opportunity"])
    assert (t["opportunity"]["source_file"] == "entries.db").sum() == result["opportunities"]
    assert "moved to the dashboard" in " ".join(after["inventory"]["note"].dropna())
    # The same donors with the same keys and the same "Individual donor #NNN" labels as before the import.
    label = lambda tb: tb["donor"].set_index("donor_key")["donor_name"].sort_index()
    assert label(t).equals(label(before))
    assert app_store.import_preview(tracker_staging, db)["opportunities"].empty           # nothing twice


def test_rows_with_problems_stay_in_excel_with_the_reason(tracker_staging, tmp_path):
    import shutil
    staging = tmp_path / "staging"
    shutil.copytree(tracker_staging, staging)
    o = pd.read_parquet(staging / "clean_opportunity.parquet")
    o.loc[o.index[0], "probability"] = None                         # e.g. "likely" typed instead of a percentage
    o.to_parquet(staging / "clean_opportunity.parquet", index=False)
    preview = app_store.import_preview(staging, tmp_path / "entries.db")
    assert list(preview["problems"]["opportunity_id"]) == [o.loc[o.index[0], "opportunity_id"]]
    assert "Probability is required" in preview["problems"]["problem"].iloc[0]
    assert "row" in preview["problems"]["where"].iloc[0]


def test_the_whole_pipeline_runs_after_the_import(synthetic, tracker_staging, tmp_path):
    """Found on the demo data: after the import the model refused to save (a sheet listed twice). Now ingest, clean,
    validate, model and indicators all run, and every opportunity is still counted once."""
    import clean
    import indicators
    import ingest
    import model
    import validate
    db = tmp_path / "entries.db"
    app_store.import_tracker("MEAL01", tracker_staging, db)
    staging, wh, docs = tmp_path / "staging", tmp_path / "wh.duckdb", tmp_path / "docs"
    docs.mkdir()
    ingest.run(synthetic["raw"], staging, app_entries=db)
    clean.run(staging)
    validate.run(staging, docs)
    tables = model.run(staging, wh, tmp_path / "exports", docs)
    indicators.run(wh, tmp_path / "exports", docs)
    assert len(tables["fact_opportunity"]) == len(pd.read_parquet(tracker_staging / "clean_opportunity.parquet"))
    assert set(tables["fact_opportunity"]["source_file"]) == {"entries.db"}
    assert not tables["dq_by_source_file"].duplicated(["source_file", "source_sheet"]).any()
