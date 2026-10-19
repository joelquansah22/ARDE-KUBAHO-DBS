"""Extension v1.2: activity records and monthly indicator reports entered on the dashboard, Excel uploads, and the
Records page. Everything runs on synthetic data generated in a temporary folder (the real data/ folder is untouched),
except the page tests, which read the synthetic warehouse."""

import datetime as dt
import shutil
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

import app_store
from app_store import EntryError

TODAY = app_store.today()
ROOT = Path(__file__).resolve().parent.parent
APP = ROOT / "dashboard" / "app.py"


@pytest.fixture(scope="module")
def staging(synthetic, tmp_path_factory):
    """The synthetic set run through ingest, clean and validate: the lists and Excel records the checks compare with."""
    import clean
    import ingest
    import validate
    folder = tmp_path_factory.mktemp("programme_staging")
    ingest.run(synthetic["raw"], folder)
    clean.run(folder)
    validate.run(folder, tmp_path_factory.mktemp("docs"))
    return folder


@pytest.fixture
def db(tmp_path):
    return tmp_path / "entries.db"


def an_activity(**change) -> dict:
    """A valid activity: PRJ01's borehole work in a real Nyagatare village, done last month."""
    rec = dict(project_id="PRJ01", activity_id="ACT02", district="Nyagatare", sector="Rihonga", cell="Rusimbi",
               village="Giganza", status="completed", date=TODAY - dt.timedelta(days=20),
               planned_date=TODAY - dt.timedelta(days=21), participants_female=12, participants_male=8,
               participants_total=20, cost_rwf=450_000, officer="FO01")
    return {**rec, **change}


def missing_report(staging) -> dict:
    """A monthly report that is NOT in any Excel file (report_completeness says not submitted)."""
    rc = pd.read_parquet(staging / "clean_report_completeness.parquet")
    gap = rc[~rc["submitted"].astype(bool)].iloc[0]
    code = gap["indicators_due"].split(", ")[0]
    return dict(project_id=gap["project_id"], district=gap["district"], period=gap["period"], indicator_code=code,
                sex=None, target_value=40, actual_value=35)


# ---------------------------------------------------------------- activity records: the cleaner's rules on entry
def test_a_valid_activity_is_saved_with_its_author_history_and_backup(staging, db):
    record_id, warnings = app_store.add_activity(an_activity(), "FO01", db, staging)
    assert record_id == "APP-A001" and warnings == []
    row = app_store.table("activities", db)[0]
    assert row["created_by"] == "FO01" and row["entry_date"] == TODAY.isoformat()
    assert [h["action"] for h in app_store.table("change_history", db)] == ["added"]
    assert app_store.list_backups(db)[0].name.endswith("_save.db")


@pytest.mark.parametrize("change, message", [
    (dict(participants_total=25), "women + men"),                              # sum mismatch: refused on entry
    (dict(participants_female=-3, participants_total=5), "egative"),
    (dict(date=dt.date(2022, 1, 10), planned_date=dt.date(2022, 1, 9)), "project"),   # before PRJ01 started
    (dict(activity_id="ACT10"), "belongs to project"),
    (dict(date=TODAY + dt.timedelta(days=3)), "future"),
    (dict(officer="FO99"), "officer list"),
    (dict(participants_female=None), "number of women"),
])
def test_activity_rules_are_checked_on_entry_in_plain_words(staging, db, change, message):
    with pytest.raises(EntryError) as e:
        app_store.add_activity(an_activity(**change), "FO01", db, staging)
    assert message.lower() in str(e.value).lower(), str(e.value)
    assert app_store.table("activities", db) == []                              # nothing saved


def test_the_staff_code_must_be_chosen(staging, db):
    with pytest.raises(EntryError, match="Choose your name"):
        app_store.add_activity(an_activity(), "", db, staging)


def test_changes_to_an_app_entry_are_checked_and_recorded(staging, db):
    record_id, _ = app_store.add_activity(an_activity(), "FO01", db, staging)
    app_store.update_activity(record_id, dict(participants_female=14, participants_total=22), "MEAL01", db, staging)
    row = app_store.table("activities", db)[0]
    assert (row["participants_female"], row["participants_total"], row["updated_by"]) == (14, 22, "MEAL01")
    changed = {(h["field"], h["old_value"], h["new_value"]) for h in app_store.table("change_history", db)
               if h["action"] == "changed"}
    assert changed == {("participants_female", "12", "14"), ("participants_total", "20", "22")}
    with pytest.raises(EntryError, match="(?i)women"):                          # still validated
        app_store.update_activity(record_id, dict(participants_total=99), "MEAL01", db, staging)


def test_an_activity_like_an_excel_record_is_saved_with_a_warning(staging, db):
    excel = pd.read_parquet(staging / "clean_activity.parquet")
    r = excel[excel["date"].notna() & excel["village"].notna() & (excel["status"] == "completed")
              & excel["participants_female"].notna()].iloc[0]
    proj = pd.read_parquet(staging / "ref_project.parquet").set_index("project_id").loc[r["project_id"]]
    rec = dict(project_id=r["project_id"], activity_id=r["activity_id"], district=r["district"], sector=r["sector"],
               cell=r["cell"], village=r["village"], status="completed", date=r["date"], planned_date=r["date"],
               participants_female=int(r["participants_female"]), participants_male=int(r["participants_male"]),
               participants_total=int(r["participants_female"] + r["participants_male"]),
               cost_rwf=1000, officer=r["officer_code"] or "FO01")
    assert proj["start_date"] <= str(r["date"])
    record_id, warnings = app_store.add_activity(rec, "FO01", db, staging)
    assert record_id and any("Excel record looks the same" in w and r["source_file"] in w for w in warnings)


# ---------------------------------------------------------------- monthly indicator reports
def test_a_missing_monthly_report_can_be_entered(staging, db):
    rec = missing_report(staging)
    record_id, _ = app_store.add_indicator_report(rec, "MEAL01", db, staging)
    saved = app_store.table("indicator_reports", db)[0]
    assert record_id == "APP-R001" and saved["period"][:10] == str(rec["period"])[:10]


def test_a_report_already_in_excel_or_already_entered_is_refused(staging, db):
    excel = pd.read_parquet(staging / "clean_indicator_result.parquet")
    r = excel[excel["sex"].isna()].iloc[0]
    with pytest.raises(EntryError, match="already in an Excel file") as e:
        app_store.add_indicator_report(dict(project_id=r["project_id"], district=r["district"], period=r["period"],
                                            indicator_code=r["indicator_code"], actual_value=3), "MEAL01", db, staging)
    assert f"row {int(r['source_row'])}" in str(e.value)
    rec = missing_report(staging)
    app_store.add_indicator_report(rec, "MEAL01", db, staging)
    with pytest.raises(EntryError, match="already entered as APP-R001"):
        app_store.add_indicator_report(rec, "PM01", db, staging)


@pytest.mark.parametrize("change, message", [
    (dict(indicator_code="AGR-01", project_id="PRJ04", district="Bugesera", sex=None), "choose the sex"),
    (dict(sex="female"), "leave the sex empty"),
    (dict(indicator_code="CC-01"), "not reported monthly"),
    (dict(actual_value=-5), "egative"),
])
def test_indicator_report_rules(staging, db, change, message):
    with pytest.raises(EntryError) as e:
        app_store.add_indicator_report({**missing_report(staging), **change}, "MEAL01", db, staging)
    assert message.lower() in str(e.value).lower(), str(e.value)


# ---------------------------------------------------------------- the pipeline reads app entries, with lineage
def test_the_pipeline_reads_app_entries_with_lineage_and_flags_repeats(synthetic, staging, tmp_path):
    import clean
    import indicators
    import ingest
    import model
    import validate
    db = tmp_path / "entries.db"
    act_id, _ = app_store.add_activity(an_activity(), "FO01", db, staging)
    rep_id, _ = app_store.add_indicator_report(missing_report(staging), "MEAL01", db, staging)
    st, wh, docs = tmp_path / "staging", tmp_path / "wh.duckdb", tmp_path / "docs"
    docs.mkdir()
    ingest.run(synthetic["raw"], st, app_entries=db)
    tables = clean.run(st)
    acts = tables["activity"][tables["activity"]["source_file"] == "entries.db"]
    reps = tables["indicator_result"][tables["indicator_result"]["source_file"] == "entries.db"]
    assert list(zip(acts["source_sheet"], acts["source_row"])) == [("activities", 1)]
    assert list(zip(reps["source_sheet"], reps["source_row"])) == [("indicator_reports", 1)]
    assert reps["entered_by"].iloc[0] == "MEAL01"
    validate.run(st, docs)
    facts = model.run(st, wh, tmp_path / "exports", docs)
    indicators.run(wh, tmp_path / "exports", docs)
    assert (facts["fact_activity"]["source_file"] == "entries.db").sum() == 1
    assert (facts["fact_indicator_result"]["source_file"] == "entries.db").sum() == 1


def test_a_report_entered_in_the_app_and_later_in_excel_is_held_back(staging):
    """The page refuses a report already in Excel; if the Excel file comes later, the pipeline holds the app entry
    back, so the indicator is never counted twice."""
    import clean
    excel = pd.read_parquet(staging / "clean_indicator_result.parquet").head(1)
    app = excel.assign(source_file="entries.db", source_sheet="indicator_reports", source_row=1)
    cleaner = clean.Cleaner(staging)
    clean.flag_app_duplicates(cleaner, dict(activity=pd.DataFrame(columns=["source_file"]),
                                            indicator_result=pd.concat([excel, app], ignore_index=True)))
    issues = cleaner.log.frame()
    assert list(issues["issue_code"]) == ["app_report_repeats_excel"] and list(issues["source_file"]) == ["entries.db"]
    from common import load_settings
    assert "app_report_repeats_excel" in load_settings()["flag_severity"]["critical"]


# ---------------------------------------------------------------- Excel uploads
@pytest.fixture
def upload_env(synthetic, staging, tmp_path):
    raw = tmp_path / "raw"
    shutil.copytree(synthetic["raw"], raw)
    return dict(raw=raw, staging=staging, entries=tmp_path / "app" / "entries.db")


def _fo_file(synthetic) -> Path:
    return sorted(Path(synthetic["raw"]).glob("SYNTHETIC_activity_monitoring_FO*.xlsx"))[0]


def test_upload_check_summarises_rows_and_warns_about_records_already_kept(synthetic, upload_env):
    import uploads
    path = _fo_file(synthetic)
    data = path.read_bytes()
    assert uploads.guess_kind(data) == "activity_monitoring"
    s = uploads.check(data, "june.xlsx", "activity_monitoring", **upload_env)
    assert s["error"] is None and s["rows_read"] > 0
    assert s["accepted"] + s["held_back"] + s["removed"] == s["rows_read"]
    assert s["held_back"] == len(s["held_list"]) and s["to_review"] == len(s["review_list"])
    assert {"Sheet", "Row", "Problem", "Details"} <= set(s["held_list"].columns) and s["held_list"]["Row"].min() >= 2
    assert max(s["overlaps"], key=s["overlaps"].get) == path.name                # the same records, another file
    assert not (upload_env["entries"].parent / "upload_check").exists()          # the scratch copy is removed


def test_accept_keeps_the_file_with_who_and_when_and_can_replace_an_earlier_one(synthetic, upload_env):
    import uploads
    path = _fo_file(synthetic)
    data = path.read_bytes()
    s = uploads.check(data, "june.xlsx", "activity_monitoring", **upload_env)
    r = uploads.accept(data, "june.xlsx", "activity_monitoring", "FO01", s, [path.name], raw=upload_env["raw"],
                       entries=upload_env["entries"], run_pipeline=False)
    stored = upload_env["raw"] / r["stored_file"]
    assert stored.exists() and "_upload_" in stored.name and "_FO01_june" in stored.name
    assert stored.name.startswith("SYNTHETIC_")                                  # synthetic mode keeps the label
    assert not (upload_env["raw"] / path.name).exists()
    assert (upload_env["entries"].parent / "replaced_files" / path.name).exists()   # kept, never deleted
    u = app_store.table("uploads", upload_env["entries"])[0]
    assert (u["uploaded_by"], u["original_name"], u["rows_read"]) == ("FO01", "june.xlsx", s["rows_read"])
    # The same file again: warned.
    again = uploads.check(data, "june (1).xlsx", "activity_monitoring", **upload_env)
    assert [x["upload_id"] for x in again["same_file"]] == [r["upload_id"]]


def _filled_indicator_template(rows: list[list], drop: str | None = None) -> bytes:
    wb = load_workbook(ROOT / "templates" / "indicator_report_template.xlsx")
    ws = wb["Report"]
    headers = [c.value for c in ws[1]]
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j, v)
    if drop:
        ws.delete_cols(headers.index(drop) + 1)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_a_filled_indicator_report_template_is_read_and_a_missing_column_is_skipped(staging, upload_env):
    import uploads
    rec = missing_report(staging)
    row = [rec["project_id"], rec["district"], pd.Timestamp(rec["period"]).to_pydatetime(), rec["indicator_code"],
           None, 40, 35, "FO01", ""]
    data = _filled_indicator_template([row])
    assert uploads.guess_kind(data) == "indicator_report"
    s = uploads.check(data, "report.xlsx", "indicator_report", **upload_env)
    assert (s["rows_read"], s["accepted"], s["held_back"], s["skipped"]) == (1, 1, 0, [])
    s = uploads.check(_filled_indicator_template([row], drop="Achieved"), "report.xlsx", "indicator_report",
                      **upload_env)
    assert s["skipped"] and s["rows_read"] == 0                                # the whole sheet, with the reason


def test_a_synthetic_file_cannot_be_added_to_real_data(monkeypatch, upload_env):
    import uploads
    monkeypatch.setattr(uploads, "load_settings", lambda: dict(synthetic_mode=False))
    assert "synthetic" in uploads.check(b"", "SYNTHETIC_x.xlsx", "activity_monitoring", **upload_env)["error"]


def test_reset_demo_also_clears_replaced_upload_files(tmp_path):
    import reset_demo
    db = tmp_path / "entries.db"
    with app_store.connect(db):
        pass
    (tmp_path / "replaced_files").mkdir()
    (tmp_path / "replaced_files" / "SYNTHETIC_x.xlsx").write_bytes(b"x")
    reset_demo.clear_entries(db)
    assert not (tmp_path / "replaced_files").exists()


# ---------------------------------------------------------------- the pages (synthetic warehouse)
@pytest.fixture
def page(tmp_path, monkeypatch):
    from common import resolve_path
    if not resolve_path("warehouse").exists():
        pytest.skip("Build the warehouse first")
    from streamlit.testing.v1 import AppTest
    monkeypatch.setattr(app_store, "db_path", lambda: tmp_path / "entries.db")

    def run(name, **state):
        at = AppTest.from_file(str(APP), default_timeout=180)
        at.switch_page(name)
        for k, v in state.items():
            at.session_state[k] = v
        at.run()
        assert not at.exception, [e.value for e in at.exception]
        return at
    return run


def test_data_entry_page_runs_and_says_excel_records_are_read_only(page):
    at = page("pages/7_Data_Entry.py", entry_staff="FO01")
    text = " ".join(m.value for m in at.caption) + " ".join(m.value for m in at.markdown)
    assert "Records from Excel files are read-only here" in text and "upload it again" in text
    assert [s.label for s in at.selectbox][:7] == ["You are", "Project *", "Activity *", "District *", "Sector *",
                                                    "Cell *", "Village *"]


def test_upload_page_offers_every_blank_template(page):
    at = page("pages/8_Upload_Excel.py")
    labels = " ".join(b.label for b in at.get("download_button"))
    for kind in ("Activity monitoring", "Beneficiary register", "Monthly indicator report", "Donor and opportunity"):
        assert kind in labels


PERSONAL = ("beneficiary", "national", "phone", "first_name", "last_name", "full_name", "id_number")


@pytest.mark.parametrize("table", ["Activities", "Indicator results", "Donors", "Opportunities",
                                   "Held back for correction", "Review queue"])
def test_records_page_shows_origin_and_no_personal_data(page, table):
    at = page("pages/9_Records.py", records_table=table)
    frame = at.dataframe[0].value
    assert {"Source", "Origin"} <= set(frame.columns) and frame["Origin"].str.len().gt(0).all()
    assert not [c for c in frame.columns if any(p in c.lower() for p in PERSONAL)]
    if table == "Donors":
        assert not frame.loc[frame["Status"] == "Individual", "Donor"].str.match(r"^Individual donor #\d+$").eq(False).any()


def test_report_form_explains_a_project_with_no_indicator_reported_monthly(page):
    """Found on the demo data: PRJ03's only indicator (WASH-03) needs a survey, so there is nothing to report monthly;
    the form says so instead of failing."""
    at = page("pages/7_Data_Entry.py", entry_staff="MEAL01", new_rep_project="PRJ03")
    assert any("no indicator reported monthly" in m.value for m in at.info)
    assert [b for b in at.button if b.label == "Save monthly report"][0].disabled


def test_data_entry_page_runs_with_saved_entries_and_opens_them_for_change(page, tmp_path):
    """Found in the screenshots: with an entry saved (cost stored as 380000.0) the 'Change an entry' tab failed."""
    from common import resolve_path
    real_staging = resolve_path("staging")                    # app_store.db_path() is the fixture's temporary file
    app_store.add_activity(an_activity(cost_rwf=380000), "FO01", app_store.db_path(), real_staging)
    app_store.add_indicator_report(missing_report(real_staging), "MEAL01", app_store.db_path(), real_staging)
    at = page("pages/7_Data_Entry.py", entry_staff="FO01")
    assert any("APP-A001" in str(s.value) for s in at.selectbox if s.label == "Entry")
    at.radio(key="edit_kind").set_value("Monthly indicator report").run()
    assert not at.exception, [e.value for e in at.exception]
    assert any("APP-R001" in str(s.value) for s in at.selectbox if s.label == "Entry")


def test_records_search_and_source_filter(page):
    """Found while checking the page: the search failed on empty cells."""
    at = page("pages/9_Records.py", records_table="Activities")
    total = len(at.dataframe[0].value)
    at.text_input(key="records_search").set_value("kayonza completed").run()
    assert not at.exception, [e.value for e in at.exception]
    found = at.dataframe[0].value
    assert 0 < len(found) < total and found["District"].eq("Kayonza").all() and found["Status"].eq("Completed").all()
    at.text_input(key="records_search").set_value("").run()
    at.selectbox(key="records_source").set_value("Excel file").run()
    assert at.dataframe[0].value["Source"].eq("Excel file").all()
    at.text_input(key="records_search").set_value("no such thing anywhere").run()        # nothing left: no failure
    assert not at.exception and len(at.dataframe[0].value) == 0
    at = page("pages/9_Records.py", records_table="Held back for correction", records_source="Entered in the app",
              records_search="kayonza")                       # a search on an empty selection
    assert len(at.dataframe[0].value) == 0
