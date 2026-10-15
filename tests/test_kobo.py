"""Scope extension: KoboToolbox forms (XLSForm) and the Kobo export import adapter."""

import re

import pytest
from openpyxl import load_workbook

from common import load_yaml, resolve_path


@pytest.fixture(scope="module")
def forms(tmp_path_factory):
    if not (resolve_path("staging") / "ref_activity.parquet").exists():
        pytest.skip("Run python src/pipeline.py first (the forms' lists come from the programme master)")
    import build_kobo_forms
    out = tmp_path_factory.mktemp("forms")
    build_kobo_forms.main(out=out)
    sheets = {}
    for path in out.glob("*.xlsx"):
        wb = load_workbook(path)
        sheets[path.stem] = {name: [dict(zip([c.value for c in ws[1]], [c.value for c in row]))
                                    for row in ws.iter_rows(min_row=2)] for name, ws in
                             ((n, wb[n]) for n in wb.sheetnames)}
    return sheets


def test_forms_have_the_three_xlsform_sheets(forms):
    assert set(forms) == {"activity_monitoring_kobo", "indicator_report_kobo"}
    for form in forms.values():
        assert {"survey", "choices", "settings"} <= set(form)
        assert form["settings"][0]["form_id"] and form["settings"][0]["version"]


def test_every_dropdown_and_filter_has_its_choices(forms):
    for name, form in forms.items():
        lists = {c["list_name"] for c in form["choices"]}
        survey_names = {q["name"] for q in form["survey"]}
        assert len(survey_names) == len(form["survey"]), f"{name}: question names repeat"
        for q in form["survey"]:
            if q["type"].startswith("select_one"):
                assert q["type"].split()[1] in lists, (name, q["name"])
            for ref in re.findall(r"\$\{(\w+)\}", " ".join(str(q.get(k) or "") for k in
                                                            ("choice_filter", "constraint", "relevant"))):
                assert ref in survey_names, (name, q["name"], ref)
        for c in form["choices"]:
            assert c["name"] and " " not in str(c["name"]), c                      # choice names without spaces


def test_forms_check_the_rules_on_the_phone(forms):
    q = {x["name"]: x for x in forms["activity_monitoring_kobo"]["survey"]}
    assert "${female} + ${male}" in q["total"]["constraint"]                       # women + men = total
    for f in ("female", "male", "total", "cost"):
        assert ". >= 0" in q[f]["constraint"]                                        # no negative values
    assert q["date"]["constraint"] == ". <= today()"                                # no future dates
    assert q["activity"]["choice_filter"] == "project=${project}"
    assert "cell=${cell}" in q["village"]["choice_filter"]
    for f in ("district", "activity", "officer", "status"):
        assert q[f]["type"].startswith("select_one") and q[f]["required"] == "yes"
    r = {x["name"]: x for x in forms["indicator_report_kobo"]["survey"]}
    assert "AGR-01" in r["sex"]["relevant"] and r["period"]["appearance"] == "month-year"


def test_question_names_and_labels_are_columns_the_pipeline_reads(forms):
    from clean import norm_key
    mappings = load_yaml("source_mappings.yaml")
    for form, entity in (("activity_monitoring_kobo", "activity"), ("indicator_report_kobo", "indicator_result")):
        known = {norm_key(a) for aliases in mappings["fields"][entity].values() for a in aliases}
        ignored = {norm_key(c) for c in mappings["ignore_columns"]}
        for q in forms[form]["survey"]:
            assert norm_key(q["name"]) in known | ignored, (form, q["name"])
            if q.get("label"):
                assert norm_key(q["label"]) in known, (form, q["label"])


def test_kobo_exports_are_read_correctly():
    if not (resolve_path("raw") / "SYNTHETIC_programme_master.xlsx").exists():
        pytest.skip("Generate the synthetic data first")
    import score_cleaning
    k = score_cleaning.score_kobo()
    assert not k["crash"] and k["skipped"] == 0
    for name, f in k["files"].items():
        assert f["pct_correct"] == 100.0, f["wrong"].head()
        assert f["duplicates_removed"] == f["duplicates"]
        assert f["blank_costs_kept_blank"] == f["blank_costs"]
        assert f["quarantined"] == 0
    st = resolve_path("kobo_test") / "staging"
    import pandas as pd
    inv = pd.read_csv(st / "sheet_inventory.csv")
    kobo = inv[inv["source_file"].str.contains("ARDE_")]
    assert len(kobo) == 2 and kobo["unmapped_columns"].isna().all()           # Kobo's own columns are ignored
